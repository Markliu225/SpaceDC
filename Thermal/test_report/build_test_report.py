# -*- coding: utf-8 -*-
"""Build the SDTwin thermal module test report in Chinese or English.

    python build_test_report.py cn [output name]
    python build_test_report.py en [output name]

The cover, version record, contents, styles and page setup come from the thermal design report of the same language.
The body follows the Orbit module test report: 1 overview, 2 test content, 3 detailed test projects, 4 test content
sufficiency analysis, 5 test conditions and requirements, 6 test data, 7 test summary. Case results come from the test
evidence in tests/results and hash-bound bilingual prose in out/reviewed_results.json. Run reviewed_results.py
and figures_results.py first, then finalize.ps1 to
refresh the table of contents and export a PDF.
"""
import copy
import hashlib
import json
import re
import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

HERE = Path(__file__).resolve().parent
THERMAL = HERE.parent
TESTS = THERMAL / 'tests'
RESULTS_DIR = TESTS / 'results'
OUT = HERE / 'out'
EN_DIR = OUT / 'en'
sys.path.insert(0, str(HERE))
import content_cn  # noqa: E402
import case_evidence  # noqa: E402

try:
    import content_en  # noqa: E402
except ImportError:  # the English content is written separately
    content_en = None

TEMPLATE = {'cn': THERMAL / 'SDTwin_Thermal_Design_Report_CN.docx', 'en': THERMAL / 'SDTwin_Thermal_Design_Report_EN.docx'}
OUT_NAME = {'cn': 'SDTwin_Thermal_Test_Report_CN', 'en': 'SDTwin_Thermal_Test_Report_EN'}
DATE = {'cn': '2026年10月6日', 'en': '6 October 2026'}
DATE_ISO = '2026-10-06'
FIRST_ISO = '2026-10-05'
LABEL_FILL = 'E7E6E6'
CASE_W = [1560, 2450, 1500, 2799]          # portrait text width 8309 twips
ORDER = ['DM-001', 'PA-001', 'EN-001', 'EN-002', 'EN-003', 'HT-001', 'HT-002', 'HT-003', 'EC-001',
         'FE-001', 'FE-002', 'FE-003', 'FE-004', 'NI-001', 'NI-002', 'NI-003']

L = {
    'cn': dict(
        table='表', figure='图',
        case_labels=['用例编号', '用例名称', '测试目的', '设计依据', '前置条件', '输入数据', '测试步骤', '预期结果', '验收判据',
                     '用例设计', '执行人与日期', '实际结果', '记录人', '测试结论', '异常记录'],
        status={'pass': '通过', 'fail': '不通过', 'partial': '部分执行', 'not_run': '未执行'},
        executor=f'自动测试，{DATE["cn"]}',
        h=['概述', '标识', '软件与文档概述', '缩略语与约定', '测试内容', '详细测试项目', '测试内容充分性分析', '测试条件与要求', '测试数据',
           '测试总结'],
        project='测试项目 {k}：{name}', case_caption='{cid} {name}测试用例',
        cover=['软件模块测试报告', '参数装配 • 表面环境 • 组件传热 • 电热耦合 • 有限元对标 • 端到端运行'],
        keywords='集总热网络; 测试用例; 有限元对标; 验收判据',
        version1=['1.0', '初稿', FIRST_ISO, '', '测试框架与对比方案初稿'],
        version2=['2.0', '执行测试用例', DATE_ISO, '', '按已实现的热模块执行全部用例，记录实际结果、测试结论与证据'],
        version3=['2.1', '结果与图表说明', DATE_ISO, '', '明确环境与场景含义，补充输入与结果表，统一温差采样并增加放大图'],
        title='SDTwin 热模块软件模块测试报告',
    ),
    'en': dict(
        table='Table', figure='Figure',
        case_labels=['Case ID', 'Case Name', 'Test Objective', 'Design Basis', 'Preconditions', 'Input Data', 'Procedure',
                     'Expected Results', 'Acceptance Criteria', 'Case Designer', 'Executor / Date', 'Actual Results',
                     'Recorded By', 'Test Conclusion', 'Anomalies'],
        status={'pass': 'Pass', 'fail': 'Fail', 'partial': 'Partially Executed', 'not_run': 'Not Executed'},
        executor=f'Automated test, {DATE_ISO}',
        h=['Overview', 'Identification', 'Software and Document Overview', 'Abbreviations and Conventions', 'Test Content',
           'Detailed Test Projects', 'Test Content Sufficiency Analysis', 'Test Conditions and Requirements', 'Test Data',
           'Test Summary'],
        project='Test Project {k}: {name}', case_caption='{cid} - {name} Test Case',
        cover=['Software Module Test Report',
               'Parameter Assembly • Surface Environment • Heat Transfer • Electrothermal Coupling • FE Benchmarks • '
               'End-to-End Runs'],
        keywords='lumped thermal network; test cases; finite element benchmark; acceptance criteria',
        version1=['1.0', 'Initial issue', FIRST_ISO, '', 'Test framework and comparison plan'],
        version2=['2.0', 'Case execution', DATE_ISO, '', 'All cases executed on the implemented thermal module; actual '
                  'results, conclusions and evidence recorded'],
        version3=['2.1', 'Results and figure explanations', DATE_ISO, '', 'Environment definitions, scenarios, inputs and result '
                  'tables added; temperature-difference plot samples aligned and local views added'],
        title='Test Report for Thermal Module of SDTwin',
    ),
}


# =============================================================== text helpers
SUPER = str.maketrans('⁰¹²³⁴⁵⁶⁷⁸⁹⁻', '0123456789−')
SUPER_RUN = re.compile(r'[⁰¹²³⁴⁵⁶⁷⁸⁹⁻]+')


def norm(text):
    """Unify powers of ten and superscript units to the ^{...} markup with the U+2212 minus sign."""
    text = re.sub(r'×10 的 ?([−-]?\d+) ?次方', lambda m: '×10^{' + m.group(1).replace('-', '−') + '}', text)
    text = re.sub(r'×10\^\{-', '×10^{−', text)
    text = SUPER_RUN.sub(lambda m: '^{' + m.group(0).translate(SUPER) + '}', text)
    return text


def fmt(v, nd=1):
    return f'{v:.{nd}f}'.replace('-', '−')


def pct(v, nd=1):
    return fmt(100.0 * v, nd)


def scientific(v, nd=3):
    if v == 0:
        return '0'
    mantissa, exponent = f'{v:.{nd}e}'.split('e')
    return mantissa.replace('-', '−') + '×10^{' + str(int(exponent)).replace('-', '−') + '}'


def pct_wide(v):
    if abs(v * 100) < 1000:
        return pct(v)
    coefficient, exponent = f'{100 * v:.2e}'.split('e')
    return coefficient.replace('-', '−') + '×10^{' + str(int(exponent)).replace('-', '−') + '}'


# =============================================================== low-level XML helpers
def w_el(tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn('w:' + k), str(v))
    return el


def m_el(tag):
    return OxmlElement('m:' + tag)


def run_props(bold=False, italic=False, size=None, color=None, superscript=False):
    rpr = w_el('w:rPr')
    rpr.append(w_el('w:rFonts', eastAsia='SimSun'))
    if bold:
        rpr.append(w_el('w:b'))
    if italic:
        rpr.append(w_el('w:i'))
    if color:
        rpr.append(w_el('w:color', val=color))
    if size:
        rpr.append(w_el('w:sz', val=size))
    if superscript:
        rpr.append(w_el('w:vertAlign', val='superscript'))
    rpr.append(w_el('w:lang', eastAsia='zh-CN'))
    return rpr


def text_run(text, **kw):
    r = w_el('w:r')
    r.append(run_props(**kw))
    t = w_el('w:t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    return r


def math_run(text, upright=False):
    r = m_el('r')
    if upright:
        mrpr = m_el('rPr')
        sty = m_el('sty')
        sty.set(qn('m:val'), 'p')
        mrpr.append(sty)
        r.append(mrpr)
    rpr = w_el('w:rPr')
    rpr.append(w_el('w:rFonts', ascii='Cambria Math', hAnsi='Cambria Math'))
    r.append(rpr)
    t = m_el('t')
    t.text = text
    r.append(t)
    return r


def omath(token):
    """'S' -> italic S; 'R_JC' -> R with subscript JC; lower-case subscript words stay upright."""
    om = m_el('oMath')
    if '_' not in token:
        om.append(math_run(token))
        return om
    base, sub = token.split('_', 1)
    ssub = m_el('sSub')
    pr = m_el('sSubPr')
    ctrl = m_el('ctrlPr')
    ctrl.append(w_el('w:rPr'))
    ctrl[0].append(w_el('w:rFonts', ascii='Cambria Math', hAnsi='Cambria Math'))
    pr.append(ctrl)
    ssub.append(pr)
    e = m_el('e')
    e.append(math_run(base))
    ssub.append(e)
    s = m_el('sub')
    for piece in re.findall(r'[a-z]+|[A-Z0-9]+|[^A-Za-z0-9]+', sub):
        s.append(math_run(piece, upright=not piece[0].isupper()))
    ssub.append(s)
    om.append(ssub)
    return om


TOKEN = re.compile(r'(\$[^$]+\$|\^\{[^}]+\}|\*\*[^*]+\*\*)')


def fill(p_el, text, bold=False, size=None):
    """Append runs to a w:p, handling $math$, ^{superscript} and **bold** markup."""
    for part in TOKEN.split(norm(text)):
        if not part:
            continue
        if part.startswith('$'):
            p_el.append(omath(part[1:-1]))
        elif part.startswith('^{'):
            p_el.append(text_run(part[2:-1], bold=bold, size=size, superscript=True))
        elif part.startswith('**'):
            p_el.append(text_run(part[2:-2], bold=True, size=size))
        else:
            p_el.append(text_run(part, bold=bold, size=size))


def ppr(p_el):
    pr = p_el.find(qn('w:pPr'))
    if pr is None:
        pr = w_el('w:pPr')
        p_el.insert(0, pr)
    return pr


# =============================================================== document writer
class Writer:
    def __init__(self, doc, lang):
        self.d = doc
        self.lang = lang
        self.anchor = None
        self.tab = 0
        self.fig = 0

    def _place(self, el):
        self.anchor.addprevious(el)
        return el

    def heading(self, level, text):
        p = w_el('w:p')
        ppr(p).append(w_el('w:pStyle', val={1: '1', 2: '21', 3: '31'}[level]))
        fill(p, text)
        return self._place(p)

    def para(self, text, indent=True, jc=None, keep_next=False, before=None):
        p = w_el('w:p')
        if not indent or jc or keep_next or before:
            pr = ppr(p)
            if keep_next:
                pr.append(w_el('w:keepNext'))
            if before:
                pr.append(w_el('w:spacing', before=before))
            if not indent:
                pr.append(w_el('w:ind', firstLine=0))
            if jc:
                pr.append(w_el('w:jc', val=jc))
        fill(p, text)
        return self._place(p)

    def field(self, label, text):
        p = w_el('w:p')
        p.append(text_run(label, bold=True))
        fill(p, text)
        return self._place(p)

    def bullet(self, text):
        p = w_el('w:p')
        pr = ppr(p)
        pr.append(w_el('w:spacing', after=60))
        pr.append(w_el('w:ind', left=600, hanging=240))
        fill(p, '• ' + text)
        return self._place(p)

    def next_number(self, kind):
        return (self.tab if kind == 'table' else self.fig) + 1

    def caption(self, title, kind='table'):
        if kind == 'table':
            self.tab += 1
            n = self.tab
        else:
            self.fig += 1
            n = self.fig
        p = w_el('w:p')
        pr = ppr(p)
        pr.append(w_el('w:pStyle', val='EngineeringCaption'))
        if kind == 'table':
            pr.append(w_el('w:keepNext'))
        word = L[self.lang][kind]
        p.append(text_run(f'{word} {n}. ' if self.lang == 'en' else f'{word} {n} '))
        fill(p, title)
        return self._place(p)

    def picture(self, path, width_cm):
        p = self.d.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = 1
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        return self._place(p._p)

    # ---- tables
    def _new_table(self, n_rows, widths):
        t = self.d.add_table(rows=n_rows, cols=len(widths))
        t.style = self.d.styles['Table Grid']
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblpr = t._tbl.tblPr
        for old in tblpr.findall(qn('w:tblW')):
            tblpr.remove(old)
        style_el = tblpr.find(qn('w:tblStyle'))
        tblpr.insert(0 if style_el is None else list(tblpr).index(style_el) + 1, w_el('w:tblW', w=sum(widths), type='dxa'))
        grid = t._tbl.tblGrid
        for gc, wd in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(wd))
        for row in t.rows:
            trpr = row._tr.get_or_add_trPr()
            trpr.append(w_el('w:cantSplit'))
            trpr.append(w_el('w:jc', val='center'))
        return t

    @staticmethod
    def _cell_fmt(cell, width, fill_hex=None, padding_y=90):
        tcpr = cell._tc.get_or_add_tcPr()
        for old in tcpr.findall(qn('w:tcW')):
            tcpr.remove(old)
        tcpr.insert(0, w_el('w:tcW', w=width, type='dxa'))
        if fill_hex:
            tcpr.append(w_el('w:shd', val='clear', color='auto', fill=fill_hex))
        mar = w_el('w:tcMar')
        for side, v in (('top', padding_y), ('left', 120), ('bottom', padding_y), ('right', 120)):
            mar.append(w_el('w:' + side, w=v, type='dxa'))
        tcpr.append(mar)
        tcpr.append(w_el('w:vAlign', val='center'))

    @staticmethod
    def _cell_text(cell, lines, bold=False, center=False, keep_next=False, kind='plain'):
        """lines: str or list of str. kind 'bullet' prefixes •, 'steps' prefixes the step number."""
        for p in cell.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        if isinstance(lines, str):
            lines = [lines]
        first = True
        for i, line in enumerate(lines):
            p_el = cell.paragraphs[0]._p if first else cell.add_paragraph()._p
            first = False
            for r in p_el.findall(qn('w:r')):
                p_el.remove(r)
            pr = ppr(p_el)
            for old in list(pr):
                pr.remove(old)
            pr.append(w_el('w:pStyle', val='EngineeringTableText'))
            if keep_next:
                pr.append(w_el('w:keepNext'))
            if kind in ('bullet', 'steps', 'steps_cn'):
                pr.append(w_el('w:ind', left=230, hanging=230))
            if center:
                pr.append(w_el('w:jc', val='center'))
            prefix = '• ' if kind == 'bullet' else (f'{i + 1}. ' if kind == 'steps' else (f'{i + 1}．' if kind == 'steps_cn' else ''))
            fill(p_el, prefix + line, bold=bold)

    def table(self, headers, rows, widths, center_cols=(), group_rows=0, padding_y=90):
        t = self._new_table(1 + len(rows), widths)
        hdr = t.rows[0]
        trpr = hdr._tr.get_or_add_trPr()
        trpr.insert(list(trpr).index(trpr.find(qn('w:jc'))), w_el('w:tblHeader'))
        for j, h in enumerate(headers):
            c = hdr.cells[j]
            self._cell_fmt(c, widths[j], LABEL_FILL, padding_y)
            self._cell_text(c, h, bold=True, center=True, keep_next=True)
        compact = len(rows) <= 6 and sum(len(str(v)) for row in rows for v in row) <= 1500
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.rows[i + 1].cells[j]
                self._cell_fmt(c, widths[j], padding_y=padding_y)
                keep = i < len(rows) - 1 and (compact or (group_rows > 0 and (i + 1) % group_rows != 0))
                self._cell_text(c, val, center=j in center_cols, keep_next=keep)
        return self._place(t._tbl)

    def case_table(self, c):
        lab = L[self.lang]['case_labels']
        t = self._new_table(12, CASE_W)
        full = sum(CASE_W[1:])

        def label(r, col, text):
            cell = t.rows[r].cells[col]
            self._cell_fmt(cell, CASE_W[col], LABEL_FILL)
            self._cell_text(cell, text, bold=True, center=True)

        def wide(r, lines, kind='plain'):
            cell = t.cell(r, 1).merge(t.cell(r, 3))
            self._cell_fmt(cell, full)
            self._cell_text(cell, lines, kind=kind)

        def short(r, col, text, bold=False, center=False):
            cell = t.rows[r].cells[col]
            self._cell_fmt(cell, CASE_W[col])
            self._cell_text(cell, text, bold=bold, center=center)

        label(0, 0, lab[0]); short(0, 1, c['id'], bold=True, center=True)
        label(0, 2, lab[1]); short(0, 3, c['name'], bold=True)
        steps = 'steps_cn' if self.lang == 'cn' else 'steps'
        rows = ((2, 'goal', 'plain'), (3, 'basis', 'plain'), (4, 'pre', 'bullet'), (5, 'input', 'bullet'),
                (6, 'steps', steps), (7, 'expect', 'bullet'), (8, 'accept', 'plain'))
        for r, (k, key, kind) in enumerate(rows, start=1):
            label(r, 0, lab[k])
            wide(r, c[key], kind)
        label(8, 0, lab[9]); short(8, 1, ''); label(8, 2, lab[10]); short(8, 3, c.get('executor', ''))
        label(9, 0, lab[11]); wide(9, c['actual'])
        label(10, 0, lab[12]); short(10, 1, ''); label(10, 2, lab[13]); short(10, 3, conclusion_marks(self.lang, c['status']))
        label(11, 0, lab[14]); wide(11, c['anomaly'])
        # Keep the case identifier with its objective instead of leaving a lone
        # header row at a page bottom.
        for cell in t.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        return self._place(t._tbl)


# =============================================================== results
def load_results():
    out = {}
    for p in sorted(RESULTS_DIR.glob('*.json')):
        if p.stem.endswith(('_series', '_runs', '_archive', '_prerun')):
            continue
        d = json.loads(p.read_text(encoding='utf-8'))
        if 'case_id' in d and 'checks' in d:
            out[d['case_id']] = d
    return out


def load_en():
    out = {}
    if EN_DIR.exists():
        for p in sorted(EN_DIR.glob('*.json')):
            d = json.loads(p.read_text(encoding='utf-8'))
            out[d['case_id']] = d
    return out


RESULTS = load_results()
DRAFT = bool(__import__('os').environ.get('REPORT_DRAFT'))
RESULTS_EN = load_en()
REVIEWED = json.loads((OUT / 'reviewed_results.json').read_text(encoding='utf-8')) if (OUT / 'reviewed_results.json').exists() else {}


def series(cid):
    p = RESULTS_DIR / f'{cid}_series.json'
    return json.loads(p.read_text(encoding='utf-8')) if p.exists() else None


def metrics(cid):
    return RESULTS[cid]['metrics'] if cid in RESULTS else {}


def test_file(cid):
    hits = sorted(TESTS.glob(f'test_{cid.lower().replace("-", "_")}_*.py'))
    return hits[0].relative_to(THERMAL).as_posix() if hits else ''


def check_counts(cid):
    checks = RESULTS[cid]['checks']
    passed = sum(1 for c in checks if c['passed'])
    return len(checks), passed, len(checks) - passed


def conclusion_marks(lang, status):
    third = 'partial' if status == 'partial' else 'not_run'
    words = L[lang]['status']
    if lang == 'en':
        return '   '.join(('[X] ' if s == status else '[ ] ') + words[s] for s in ('pass', 'fail', third))
    return '  '.join(('☒ ' if s == status else '☐ ') + words[s] for s in ('pass', 'fail', third))


def status_word(lang, cid):
    return L[lang]['status'][RESULTS[cid]['status']] if cid in RESULTS else L[lang]['status']['not_run']


def case_result_fields(lang, cid):
    r = RESULTS[cid]
    reviewed = REVIEWED.get(cid)
    if reviewed:
        digest = hashlib.sha256((RESULTS_DIR / f'{cid}.json').read_bytes()).hexdigest()
        if reviewed['source_sha256'] != digest:
            raise RuntimeError(f'{cid}: report prose is stale; rerun reviewed_results.py')
        r = {**r, **{k: reviewed[k] for k in ('summary_cn', 'anomalies_cn')}}
    n, ok, bad = check_counts(cid)
    if lang == 'cn':
        summary, anomaly = r['summary_cn'].strip(), (r['anomalies_cn'].strip() or '无')
        tail = (f'测试程序 {test_file(cid)} 共记录 {n} 项检查，{ok} 项通过' + (f'，{bad} 项未通过' if bad else '')
                + f'。证据文件为 tests/results/{cid}.json。')
    else:
        en = reviewed or RESULTS_EN.get(cid)
        if en is None:
            if not DRAFT:
                raise SystemExit(f'missing English result text {EN_DIR / (cid + ".json")}')
            en = {'summary_en': 'English result text pending.', 'anomalies_en': 'Pending.'}
        summary, anomaly = en['summary_en'].strip(), ((en.get('anomalies_en') or '').strip() or 'None')
        tail = (f'Test program {test_file(cid)} recorded {n} checks, {ok} passed' + (f', {bad} failed' if bad else '')
                + f'. Evidence file: tests/results/{cid}.json.')
    return dict(actual=[summary, tail], anomaly=anomaly, status=r['status'], executor=L[lang]['executor'])


# =============================================================== template surgery
def set_runs(p, texts):
    runs = p.runs
    for r, txt in zip(runs, texts):
        r.text = txt
    for r in runs[len(texts):]:
        r._r.getparent().remove(r._r)


def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        set_runs(p, [text])
    else:
        p._p.append(text_run(text, color='000000', size=18))


def prepare(doc, lang, abstract):
    """Cover, version record and body anchors. The appendix section is dropped and the landscape summary section
    becomes the last section, as in the Orbit module test report; running header texts are cleared."""
    t = L[lang]
    # Keep the full contents on one readable page instead of a nearly empty overflow page.
    for style in doc.styles:
        if style.style_id in ('TOC1', 'TOC2', 'TOC3'):
            props = style.element.get_or_add_pPr()
            old_spacing = props.find(qn('w:spacing'))
            if old_spacing is not None:
                props.remove(old_spacing)
            props.append(w_el('w:spacing', before=0, after=40, line=280, lineRule='exact'))
    body = doc.element.body
    kids = list(body)
    set_runs(docx.text.paragraph.Paragraph(kids[6], doc), [t['cover'][0]])
    set_runs(docx.text.paragraph.Paragraph(kids[7], doc), [t['cover'][1]])
    cover_tbl = docx.table.Table(kids[8], doc)
    set_cell(cover_tbl.rows[3].cells[1], DATE[lang])
    vt = docx.table.Table(kids[12], doc)
    set_cell(vt.rows[0].cells[1], abstract)
    set_cell(vt.rows[1].cells[1], t['keywords'])
    for row, values in ((4, t['version1']), (5, t['version2']), (6, t['version3'])):
        for k, (cell, txt) in enumerate(zip(vt.rows[row].cells, values)):
            if row >= 5:
                # the empty template row carries no paragraph formatting: take the alignment of the row above
                src = vt.rows[4].cells[k].paragraphs[0]._p.find(qn('w:pPr'))
                dst = cell.paragraphs[0]._p
                old_pr = dst.find(qn('w:pPr'))
                if old_pr is not None:
                    dst.remove(old_pr)
                if src is not None:
                    dst.insert(0, copy.deepcopy(src))
            set_cell(cell, txt)
    first = next(i for i, el in enumerate(kids) if el.tag == qn('w:p') and el.find('.//' + qn('w:pStyle')) is not None
                 and el.find('.//' + qn('w:pStyle')).get(qn('w:val')) == '1')
    breaks = [i for i, el in enumerate(kids) if i > first and el.tag == qn('w:p') and el.find('.//' + qn('w:sectPr')) is not None]
    sec_body_end, sec_land_end = breaks[0], breaks[1]
    final = kids[-1]
    for el in kids[first:sec_body_end] + kids[sec_body_end + 1:sec_land_end] + kids[sec_land_end + 1:-1]:
        body.remove(el)
    # the landscape section becomes the last section: its properties replace the body-level section properties
    land_p = kids[sec_land_end]
    new_final = copy.deepcopy(land_p.find('.//' + qn('w:sectPr')))
    # The landscape summary needs less vertical margin than portrait case sheets.
    # Twenty-millimetre margins keep its conclusion and signature with the two-page table.
    margins = new_final.find(qn('w:pgMar'))
    margins.set(qn('w:top'), '1134')
    margins.set(qn('w:bottom'), '1134')
    body.replace(final, new_final)
    body.remove(land_p)
    for sec in doc.sections:
        for hdr in (sec.header, sec.first_page_header, sec.even_page_header):
            if hdr.is_linked_to_previous:
                continue
            for p in hdr.paragraphs:
                for r in p.runs:
                    r.text = ''
    doc.core_properties.title = t['title']
    return kids[sec_body_end], new_final


# =============================================================== static content per language
def content(lang):
    if lang == 'cn':
        c = content_cn
        return dict(groups=c.GROUPS_CN, cases=c.CASES_CN, abbr=c.ABBR_CN, coverage=c.COVERAGE_CN, compare=c.COMPARE_CN,
                    conditions=c.CONDITIONS_CN, data=c.DATA_CN)
    c = content_en
    if c is None:
        raise SystemExit('content_en.py is missing')
    return dict(groups=c.GROUPS_EN, cases=c.CASES_EN, abbr=c.ABBR_EN, coverage=c.COVERAGE_EN, compare=c.COMPARE_EN,
                conditions=c.CONDITIONS_EN, data=c.DATA_EN)


CASE_NAME = {'cn': {'cold0': '设计冷工况', 'nom0': '平均环境工况', 'hot75': '设计热工况'},
             'en': {'cold0': 'Design cold', 'nom0': 'Mean environment', 'hot75': 'Design hot'}}


# =============================================================== finite element result blocks after the case tables
def iss_temperature_figure(w, lang, key):
    captions = {
        'nom0_T_iso_noon': ('ISS 平均环境工况第三圈正午附近的有限元表面温度',
                          'ISS FE surface temperatures near third-orbit noon, mean environment'),
        'cold0_T_iso_ecl': ('ISS 设计冷工况第三圈地影中点附近的有限元表面温度',
                          'ISS FE surface temperatures near third-orbit mid-eclipse, design cold'),
        'hot75_T_hrs_noon': ('ISS 设计热工况第三圈正午附近的散热器局部温度场',
                           'ISS radiator FE temperature field near third-orbit noon, design hot'),
    }
    notes = {
        'nom0_T_iso_noon': (
            '本图给出对标对象的空间布局：两端长条为太阳翼，中央成组面板为散热器，其余为桁架、舱体及设备。工况为 nom0，β=0°，本时刻处于日照；具体辐射输入见第 2 章。它说明有限元保留各部件的空间温度分布，而热模块对每个节点只输出一个代表温度。',
            'This locates the benchmark objects: long strips at the ends are solar arrays, grouped central panels are radiators, and the remainder includes truss, modules and equipment. The case is nom0, β=0°, sunlit at this instant; Chapter 2 gives its radiation inputs. FE retains spatial temperature distributions, while the thermal module gives one representative temperature per node.'),
        'cold0_T_iso_ecl': (
            '本图对应 cold0、β=0° 的日食阶段，太阳直射被地球挡住，太阳翼仍继续辐射散热。它帮助定位 FE-001 的太阳翼对象，并把降温曲线与空间温度场对应起来；整站表面云图不能直接当作太阳翼平均温度。FE-001 的定量判定仍用前述同一时刻的太阳翼平均温度及差值。',
            'This shows eclipse in cold0 at β=0°: Earth blocks direct sunlight while the array continues radiating. It identifies the FE-001 arrays and connects the cooling history with a spatial field. The station-wide surface map is not an array-mean temperature. FE-001 acceptance still uses the preceding simultaneous array means and their differences.'),
        'hot75_T_hrs_noon': (
            '本图对应 hot75、β=75° 的连续日照环境，从右舷上方观察散热器。中央三排为右舷散热器面板，最上一排为后方左舷面板，前景横条为太阳翼。局部颜色变化显示面板间及面板内的温差，说明 FE-002 的回路平均温度不能代表最冷点；图中的局部颜色不作为单独的验收数值。',
            'This is the continuously sunlit hot75 case at β=75°, viewed from above the starboard side. The central three rows are starboard radiator panels, the uppermost row is the port wing behind them, and foreground strips are solar arrays. Spatial colour variation shows differences within and between panels, explaining why the FE-002 loop mean cannot represent the coldest point. Local colour is not a separate acceptance value.'),
    }
    index = 0 if lang == 'cn' else 1
    w.picture(OUT / f'fig_iss_{key}_{lang}.png', 14.0)
    w.caption(captions[key][index], kind='figure')
    w.para(notes[key][index])
    phase = '13884.1' if key == 'cold0_T_iso_ecl' else '11107.2'
    w.para((f'来源：已求解 COMSOL 模型的 {key}.png；绘图目标时刻 t={phase} s，导出程序取最近的已保存解。色标为表面温度 −80 至 80°C，超出范围时颜色饱和；它不是温差色标。图像保留原有限元温度场，仅重绘同范围色标。' if lang == 'cn' else
            f'Source: {key}.png from the solved COMSOL model. The requested time is t={phase} s; export selects the nearest saved solution. The scale is surface temperature from −80 to 80°C, with saturation outside that range, not temperature difference. The original FE field is retained; only the same-range colour bar is redrawn.'))


def block_en003(w, lang):
    m = metrics('EN-003')
    loads = [('太阳翼', '反照'), ('太阳翼', '红外'), ('散热器', '反照'), ('散热器', '红外')]
    names = {'cn': {('太阳翼', '反照'): '太阳翼反照', ('太阳翼', '红外'): '太阳翼红外', ('散热器', '反照'): '散热器反照',
                    ('散热器', '红外'): '散热器红外'},
             'en': {('太阳翼', '反照'): 'Solar array albedo', ('太阳翼', '红外'): 'Solar array infrared',
                    ('散热器', '反照'): 'Radiator albedo', ('散热器', '红外'): 'Radiator infrared'}}
    rows = []
    for case in ('cold0', 'nom0', 'hot75'):
        for obj, kind in loads:
            v = m.get(f'fe_{case}_{obj}_{kind}_load')
            if not v:
                continue
            rows.append([CASE_NAME[lang][case], names[lang][(obj, kind)], fmt(v['module_mean_W'] / 1000, 2),
                         fmt(v['fe_mean_W'] / 1000, 2), pct(v['rel_mean']), pct(v['rel_peak']),
                         pct_wide(v.get('pointwise_max_rel', v['inst_max_rel_above_10pct_peak']))])
    if not rows:
        return
    refined = m.get('fe_refined_comparison')
    if lang == 'cn':
        w.para(f'表 {w.next_number("table")} 给出步骤 3 的有限元热载荷对比，有限元值取自已求解模型，地球离散为 2 圈每圈 6 个点。'
               '相差为 earth_flux 减有限元再除以有限元，峰值相差按两者各自的峰值计算。逐点最大差对同一时刻的非零参考值计算，'
               '参考值极小时相对误差可很大，需结合后述微弱热流的绝对值判断适用范围。', keep_next=True)
        w.caption('EN-003 地球反照与红外热载荷对比，已求解模型')
        w.table(['工况', '热载荷', 'earth_flux 平均 kW', '有限元平均 kW', '平均相差 %', '峰值相差 %', '逐点最大差 %'], rows,
                [1200, 1300, 1250, 1250, 1000, 1000, 1309], center_cols=tuple(range(7)), group_rows=4)
    else:
        w.para(f'Table {w.next_number("table")} gives the step 3 comparison of thermal loads. The FE values come from the '
               'solved models, whose planet is discretised into 2 rings of 6 points. Differences are earth_flux minus FE '
               'divided by FE; peak differences compare the two peaks. Pointwise maxima compare nonzero FE values at '
               'the same instants. Tiny reference values can produce very large relative errors, so absolute weak-flux values must also be examined.', keep_next=True)
        w.caption('EN-003 Earth albedo and infrared thermal loads, solved models')
        w.table(['Case', 'Thermal load', 'earth_flux mean kW', 'FE mean kW', 'Mean diff %', 'Peak diff %', 'Pointwise max %'], rows,
                [1200, 1300, 1250, 1250, 1000, 1000, 1309], center_cols=tuple(range(7)), group_rows=4)
    diagnostic = [(case, obj, kind, m[f'fe_{case}_{obj}_{kind}_load'])
                  for case in ('cold0', 'nom0', 'hot75') for obj, kind in loads
                  if m.get(f'fe_{case}_{obj}_{kind}_load', {}).get('worst_relative_point')]
    if diagnostic:
        case, obj, kind, value = max(diagnostic, key=lambda item: item[3]['pointwise_max_rel'])
        point = value['worst_relative_point']
        if lang == 'cn':
            w.para(f'最大逐点相对差出现在{CASE_NAME[lang][case]}的{names[lang][(obj, kind)]}，时刻为 {point["time_s"]:.0f} s。'
                   f'该点 earth_flux 为 {fmt(point["module_W"], 2)} W，有限元为 {scientific(point["fe_W"])} W，绝对差为 {fmt(point["abs_diff_W"], 2)} W。'
                   '参考值接近零使相对差很大。本用例仍分别保留逐点、时间加权平均与峰值的 3% 检查。')
        else:
            w.para(f'The largest pointwise relative difference occurs for {names[lang][(obj, kind)].lower()} in the '
                   f'{CASE_NAME[lang][case].lower()} case at {point["time_s"]:.0f} s. The module gives {fmt(point["module_W"], 2)} W '
                   f'and FE gives {scientific(point["fe_W"])} W, an absolute difference of {fmt(point["abs_diff_W"], 2)} W. '
                   'The near-zero reference produces a large relative difference. Pointwise values, time-weighted means '
                   'and peaks retain separate 3% checks.')
    if refined:
        block_en003_refined(w, lang, refined)
    orientation = series('EN-003').get('orientation', {})
    weak_rows = []
    directions = {'cn': {'nadir': '对地', 'zenith': '背地', 'along_pos': '沿轨正向',
                          'cross_pos': '横向正向', 'sun': '朝向太阳', 'antisun': '背向太阳'},
                  'en': {'nadir': 'Nadir', 'zenith': 'Zenith', 'along_pos': 'Along track +',
                          'cross_pos': 'Cross track +', 'sun': 'Sun facing', 'antisun': 'Anti-Sun'}}
    def flux_value(value):
        if value == 0:
            return '0'
        if abs(value) < 1e-4:
            return scientific(value)
        return fmt(value, 8)
    for beta, block in orientation.items():
        candidates = []
        for key, vals in block.items():
            if not key.endswith('_albedo_module_W_m2'):
                continue
            refs = block[key.replace('_module_', '_reference_')]
            for index, (mod, reference) in enumerate(zip(vals, refs)):
                if reference > 1e-10 and abs(mod - reference) / reference > 0.01:
                    candidates.append((abs(mod - reference), abs(mod - reference) / reference,
                                       key.split('_albedo')[0], block['orbit_angle_deg'][index], mod, reference))
        if candidates:
            for row in [max(candidates, key=lambda v: v[0]), max(candidates, key=lambda v: v[1])]:
                weak_rows.append([beta.replace('beta_', ''), directions[lang].get(row[2], row[2]), fmt(row[3], 2),
                                  flux_value(row[4]), flux_value(row[5]), flux_value(row[0]), pct(row[1], 2)])
    if weak_rows:
        if lang == 'cn':
            w.para('对微弱反照，独立参照在地心角与方位角上分开积分，并在受光与可见边界处分段，'
                   '使用自适应积分核对均匀网格。下表给出平均环境参数下的代表性超限点，分别选取绝对差与相对差最大的点。'
                   '这说明当前 earth_flux 默认网格在微弱热流下不能保证统一的 1% 相对精度。')
            w.caption('EN-003 微弱反照的代表性误差')
            headers = ['β °', '法向', '轨道角 °', '模块 W/m²', '参照 W/m²', '绝对差 W/m²', '相对差 %']
        else:
            w.para('For weak albedo, the independent geocentric reference separates ring and azimuth integration, '
                   'splits at illumination and visibility boundaries and uses adaptive integration to check uniform grids. '
                   'The following mean-environment examples select the largest absolute and relative exceedances. '
                   'The default earth_flux grid therefore does not guarantee a uniform 1% relative accuracy at weak flux.')
            w.caption('EN-003 representative weak-albedo errors')
            headers = ['β °', 'Normal', 'Orbit angle °', 'Module W/m²', 'Reference W/m²', 'Abs diff W/m²', 'Rel diff %']
        w.table(headers, weak_rows, [650, 1050, 950, 1500, 1500, 1550, 1109], center_cols=tuple(range(7)))


def block_en003_refined(w, lang, refined):
    levels = refined['levels']
    rows = []
    for case in ('cold0', 'nom0', 'hot75'):
        for key, rec in refined['cases'].get(case, {}).items():
            row = [CASE_NAME[lang][case], rec['label_' + lang]]
            row += [fmt(rec['fe_mean_kW'][lv], 2) for lv in levels]
            row += [fmt(rec['module_mean_kW'], 2), pct(rec['rel_vs_finest']),
                    pct(rec['rel_vs_extrapolated']) if rec.get('extrapolation_monotone') else 'N/A']
            rows.append(row)
    lv_txt = [lv.replace('x', '×') for lv in levels]
    if lang == 'cn':
        w.para(f'为核对参照，用 fe_refine_ladder.py 在同一模型上逐级加密地球离散，只重算轨道热载荷，在 {refined["sample_start_s"]:.0f} 至 {refined["sample_end_s"]:.0f} s '
               f'按 {refined["step_s"]:.0f} s 间隔取 {refined["instants"]} 个时刻。这里使用样本算术平均，不是完整轨道周期的加权平均。表 {w.next_number("table")} 给出各级有限元热载荷在这些时刻的平均值、'
               'earth_flux 在同一时刻的平均值，以及相对最细一级与相对条件性外推估计的差。只有相邻级别变化同向且缩小时才计算外推值，'
               '否则标为 N/A。COMSOL 额外保存的日食边界点保留在原始文件中，但不纳入公共采样均值。'
               '三种离散级别不能单独证明已进入渐近收敛区。', keep_next=True)
        w.caption('EN-003 有限元地球离散加密后的热载荷对比')
        headers = ['工况', '热载荷'] + [f'有限元 {t} kW' for t in lv_txt] + ['模块 kW', '相对最细一级 %', '相对外推值 %']
    else:
        w.para('To verify the reference, fe_refine_ladder.py re-solved only the orbital thermal loads of the same models '
               f'with successively finer planet discretisations at {refined["instants"]} instants spaced '
               f'{refined["step_s"]:.0f} s from {refined["sample_start_s"]:.0f} to {refined["sample_end_s"]:.0f} s. These are arithmetic sample means, '
               f'not time-weighted means over an exact orbit. Table {w.next_number("table")} gives the mean FE loads of each '
               'level at these instants, the earth_flux mean at the same instants, and the difference of earth_flux from '
               'the finest level and from a conditional extrapolation. Extrapolation is shown only when successive changes '
               'have the same sign and shrink; otherwise it is N/A. Extra eclipse-boundary samples saved by COMSOL '
               'remain in the raw exports but are excluded from the common sample means. '
               'Three levels alone do not establish asymptotic convergence.', keep_next=True)
        w.caption('EN-003 thermal loads with refined FE planet discretisation')
        headers = ['Case', 'Thermal load'] + [f'FE {t} kW' for t in lv_txt] + ['Module kW', 'vs finest %',
                                                                              'vs extrap. %']
    n = len(headers)
    widths = [1300, 1500] + [int((8309 - 2800) / (n - 2))] * (n - 2)
    widths[-1] += 8309 - sum(widths)
    w.table(headers, rows, widths, center_cols=tuple(range(n)), group_rows=4)
    maxima = [(case, max(abs(rec['rel_vs_finest']) for rec in records.values()))
              for case, records in refined['cases'].items() if records]
    changes = [(abs(rec['fe_mean_kW'][levels[-1]] / rec['fe_mean_kW'][levels[0]] - 1), case, rec)
               for case, records in refined['cases'].items() for rec in records.values()
               if rec['fe_mean_kW'][levels[0]] != 0]
    if maxima and changes:
        change, case, rec = max(changes, key=lambda item: item[0])
        if lang == 'cn':
            detail = '、'.join(f'{CASE_NAME[lang][case]} {pct(value)}%' for case, value in maxima)
            w.para(f'在这些采样均值中，模块与最细一级的最大相对差分别为{detail}。'
                   f'地球离散从 {lv_txt[0]} 加密至 {lv_txt[-1]} 后，有限元均值的最大变化为 {pct(change)}%，'
                   f'出现在{CASE_NAME[lang][case]}的{rec["label_cn"]}。该变化说明原始参考值对地球离散敏感。'
                   '这里未验证完整轨道的逐点与峰值误差，也未重算温度场，因此不替代原验收。')
        else:
            detail = ', '.join(f'{CASE_NAME[lang][case]} {pct(value)}%' for case, value in maxima)
            w.para(f'Across these sample means, the largest module-to-finest relative differences are {detail}. '
                   f'Refining the planet from {lv_txt[0]} to {lv_txt[-1]} changes the FE mean by up to {pct(change)}%, '
                   f'for {rec["label_en"].lower()} in the {CASE_NAME[lang][case].lower()} case. This demonstrates '
                   'sensitivity of the original reference to planet discretisation. This comparison does not verify '
                   'full-orbit pointwise or peak errors and does not re-solve the temperature field, so it does not '
                   'replace the original acceptance checks.')


def block_fe001(w, lang):
    s = series('FE-001')
    m = metrics('FE-001')
    if not s:
        return
    rows = []
    for case in ('cold0', 'nom0', 'hot75'):
        c = s['cases'][case]
        fe, mo = c['stats_C']['fe'], c['stats_C']['module']
        d = c['module_minus_fe_K']
        rows.append([CASE_NAME[lang][case], ' / '.join(fmt(fe[k], 2) for k in ('min', 'mean', 'max')),
                     ' / '.join(fmt(mo[k], 2) for k in ('min', 'mean', 'max')),
                     fmt(max(abs(d['min']), abs(d['mean']), abs(d['max'])), 2)])
    point_rows = [[CASE_NAME[lang][c], fmt(m[c + '_comparison']['pointwise_max_abs_K'], 2),
                   fmt(m[c + '_comparison']['pointwise_rms_K'], 2), '3.00'] for c in ('cold0', 'nom0', 'hot75')]
    no = {case: m.get(f'{case}_no_earth', {}).get('with_minus_without_K', {}) for case in ('cold0', 'nom0', 'hot75')}
    peak = [no[c]['max'] for c in no if no[c]]
    low = [no[c]['min'] for c in ('cold0', 'nom0') if no[c]]
    if lang == 'cn':
        w.para(f'表 {w.next_number("table")} 与图 {w.next_number("figure")} 给出热模块与有限元太阳翼平均温度在最后一圈的对比。', keep_next=True)
        w.caption('FE-001 太阳能板最后一圈温度对比')
        w.table(['工况', '有限元 最低 / 平均 / 最高 °C', '热模块 最低 / 平均 / 最高 °C', '最大差 K'], rows,
                [1700, 2600, 2600, 1409], center_cols=(0, 1, 2, 3))
        w.caption('FE-001 同一时刻逐点温度差')
        w.table(['工况', '逐点最大差 K', '均方根差 K', '逐点限值 K'], point_rows,
                [2300, 2000, 2000, 2009], center_cols=(0, 1, 2, 3))
        w.picture(OUT / 'fig_fe001_solar_cn.png', 14.6)
        w.caption('FE-001 太阳能板最后一圈温度对比', kind='figure')
        w.para('这里的冷、平均工况均为 β=0°：每圈地球遮挡太阳约 36.1 min，直射吸热中断，而太阳翼仍向外辐射，因此出现明显降温。热工况为 β=75°，全圈日照，所以没有同样的日食降温段。三组辐射输入见第 2 章环境表；这些名称不是上排曲线的温度等级。')
        w.para('三列依次为冷、平均、热工况，蓝线为有限元，橙线为热模块。横轴从最后一圈起点计时，单位 min；上、中排纵轴为 °C。上排统一温度刻度，紫色区域在中排放大，灰色区域表示日食。两条曲线使用相同的有限元输出时刻，点间直线连接；下排为热模块减有限元，绿色带为 ±3 K。平均环境工况在 32.56 min 为 −7.0428 − 5.0425 = −12.0853 K，按两位小数显示为 −12.09 K。')
        if peak:
            w.para(f'关闭地球反照与红外后，最高温度低 {fmt(min(peak))} 至 {fmt(max(peak))} K，有日食工况的最低温度低 '
                   f'{fmt(min(low))} 至 {fmt(max(low))} K，说明设计报告第 5.3 节要求的 earth_flux 不可缺少。')
    else:
        w.para(f'Table {w.next_number("table")} and Figure {w.next_number("figure")} compare the thermal module with the FE '
               'mean temperature of the solar array wings over the last orbit.', keep_next=True)
        w.caption('FE-001 solar array temperatures over the last orbit')
        w.table(['Case', 'FE min / mean / max °C', 'Thermal module min / mean / max °C', 'Max difference K'], rows,
                [1700, 2600, 2600, 1409], center_cols=(0, 1, 2, 3))
        w.caption('FE-001 simultaneous pointwise temperature differences')
        w.table(['Case', 'Pointwise max K', 'RMS K', 'Pointwise limit K'], point_rows,
                [2300, 2000, 2000, 2009], center_cols=(0, 1, 2, 3))
        w.picture(OUT / 'fig_fe001_solar_en.png', 14.6)
        w.caption('FE-001 solar array temperatures over the last orbit', kind='figure')
        w.para('Cold and mean here use β=0°: Earth blocks sunlight for about 36.1 min per orbit. Direct solar heating stops while the array continues radiating, producing marked cooling. Hot uses β=75° and remains sunlit, so it has no corresponding eclipse-cooling interval. Chapter 2 lists the radiation inputs; the names do not rank temperatures in the upper curves.')
        w.para('Columns show cold, mean and hot cases; blue is FE and orange is the module. Time in min starts at the last-orbit boundary; top and middle temperatures are in °C. The top row shares one temperature scale, purple windows are enlarged in the middle row, and grey marks eclipse. Both curves use the same FE output instants, joined by straight lines. The bottom row is module minus FE; green denotes ±3 K. At 32.56 min in the mean case, −7.0428 − 5.0425 = −12.0853 K, displayed as −12.09 K to two decimal places.')
        if peak:
            w.para(f'With Earth albedo and infrared switched off, the maximum temperature is {fmt(min(peak))} to '
                   f'{fmt(max(peak))} K lower and the minimum temperature of the cases with eclipse is {fmt(min(low))} to '
                   f'{fmt(max(low))} K lower, so the earth_flux input required by Section 5.3 of the design report is '
                   'indispensable.')


def block_fe002(w, lang):
    m = metrics('FE-002')
    if not m:
        return
    rows = []
    for case in ('cold0', 'nom0', 'hot75'):
        for lp in ('A', 'B'):
            fe = m[f'{case}_{lp}_fe_orbit3_C']['mean']
            mo = m[f'{case}_{lp}_module_orbit3_C']['mean']
            rows.append([CASE_NAME[lang][case], lp, fmt(fe, 2), fmt(mo, 2), fmt(mo - fe, 2)])
    cold = m.get('coldest_panel', {})
    gaps = [v['mean_minus_coldest_K'] for v in cold.values() if isinstance(v, dict) and 'mean_minus_coldest_K' in v]
    if lang == 'cn':
        w.para(f'表 {w.next_number("table")} 与图 {w.next_number("figure")} 给出公共散热板第三圈平均温度的对比，有限元值为三个散热器单元面板平均温度之平均。', keep_next=True)
        w.caption('FE-002 公共散热板第三圈平均温度对比')
        w.table(['工况', '回路', '有限元 °C', '热模块 °C', '热模块减有限元 K'], rows, [2000, 1100, 1700, 1700, 1809],
                center_cols=(0, 1, 2, 3, 4))
        w.picture(OUT / 'fig_fe002_radiator_cn.png', 14.6)
        w.caption('FE-002 公共散热板第三圈温度对比', kind='figure')
        w.para('三列环境按第 2 章定义：前两列 β=0°、有日食，第三列 β=75°、无日食。A、B 回路各接收自己的有限元废热输入，面板在日照时侧对太阳、日食时朝地。各曲线是回路面板的平均温度；低于 0°C 指面板本身的温度，不是周围空气温度。')
        w.para('图中上、下排分别是 A、B 回路，三列依次为冷、平均、热环境，灰色区域表示日食。横轴从第三圈起点计时，单位 min；纵轴为回路面板平均温度，单位 °C，各子图刻度不同。蓝实线为有限元，橙虚线为给定同一进热时程后的热模块。曲线显示相位和波动，验收则比较上表的第三圈时间平均值；不能把图中的任一瞬时差直接当作 5 K 平均温度判据。')
        if gaps:
            w.para(f'同一时刻有限元最冷面板比散热器平均温度低 {fmt(min(gaps))} 至 {fmt(max(gaps))} K。设计报告第 10 章说明散热板只用集总温度，'
                   '不计算局部温差，防冻判断需要另算最冷点。有限元包含翼间遮挡，设计报告第 4.4 节说明本版不计组件间遮挡；其独立温度影响尚未分离。')
    else:
        w.para(f'Table {w.next_number("table")} and Figure {w.next_number("figure")} compare the third-orbit mean temperature '
               'of the common radiator; the FE value is the average of the panel mean temperatures of the three radiator '
               'units.', keep_next=True)
        w.caption('FE-002 common radiator third-orbit mean temperatures')
        w.table(['Case', 'Loop', 'FE °C', 'Thermal module °C', 'Module minus FE K'], rows, [2000, 1100, 1700, 1700, 1809],
                center_cols=(0, 1, 2, 3, 4))
        w.picture(OUT / 'fig_fe002_radiator_en.png', 14.6)
        w.caption('FE-002 common radiator temperatures over the third orbit', kind='figure')
        w.para('Chapter 2 defines the environments: the first two columns have β=0° and eclipse; the third has β=75° and no eclipse. Loops A and B receive their own prescribed FE waste-heat inputs. Panels face edge-on to sunlight and toward Earth in eclipse. Each curve is a loop panel-mean temperature; below 0°C describes the panels, not surrounding air.')
        w.para('Rows show loops A and B; columns show cold, mean and hot environments, with grey marking eclipse. Time in min is measured from the third-orbit start. The vertical axis is loop panel-mean temperature in °C, with different scales across panels. Blue solid curves are FE results; orange dashed curves are the module response to the same prescribed heat-input history. Curves show phase and variation, while acceptance compares the time means in the preceding table. An instantaneous gap is not the 5 K mean-temperature criterion.')
        if gaps:
            w.para(f'At the same instant the coldest FE panel is {fmt(min(gaps))} to {fmt(max(gaps))} K below the radiator '
                   'mean. Chapter 10 of the design report states that the radiator uses one lumped temperature without '
                   'local gradients, so freeze protection needs a separate coldest-point calculation. The FE model includes '
                   'wing obstruction, which Section 4.4 excludes from this version; its separate temperature contribution '
                   'has not been isolated.')


def block_fe003(w, lang):
    m = metrics('FE-003')
    s = series('FE-003')
    if not m or not s:
        return
    rows = []
    for dev in ('MBSU', 'DDCU', 'IEA'):
        r = m['step1_r_jc'][dev]
        items = [v for v in s['fe_items'].values() if v['device'] == dev]
        p = sum(v['P_load_W'] for v in items) / len(items)
        fe = sum(v['fe_Tmean_C'] for v in items) / len(items)
        mo = sum(v['module_mean_C'] for v in items) / len(items)
        worst = max(abs(v['diff_K']) for v in items)
        rows.append([dev, str(len(items)), f'{p:.0f}', f'{r["area_m2"]:.4g}', f'{r["contact_K_W"]:.5f}',
                     f'{r["conduction_K_W"]:.5f}', f'{r["R_JC_K_W"]:.5f}', fmt(mo, 2), fmt(fe, 2), fmt(worst, 2)])
    if lang == 'cn':
        w.para(f'表 {w.next_number("table")} 按设备类别给出 $R_JC$ 的组成与平均温度对比，温度为同类设备的平均值，最大差为单台设备差值绝对值的最大值。', keep_next=True)
        w.caption('FE-003 计算节点平均温度对比')
        w.table(['设备', '台数', '$P_load$ W', '冷板面积 m²', '接触部分 K/W', '导热部分 K/W', '$R_JC$ K/W', '热模块 °C', '有限元 °C', '最大差 K'],
                rows, [760, 560, 760, 860, 900, 900, 900, 880, 880, 909], center_cols=tuple(range(10)))
    else:
        w.para(f'Table {w.next_number("table")} gives the composition of $R_JC$ and the mean temperatures by device class; '
               'temperatures are averages over the units of a class and the maximum difference is the largest absolute '
               'difference of a single unit.', keep_next=True)
        w.caption('FE-003 computing node mean temperatures')
        w.table(['Device', 'No.', '$P_load$ W', 'Cold plate area m²', 'Contact K/W', 'Solid K/W', '$R_JC$ K/W',
                 'Module °C', 'FE °C', 'Max diff K'],
                rows, [760, 560, 760, 860, 900, 900, 900, 880, 880, 909], center_cols=tuple(range(10)))


def block_fe004(w, lang):
    m = metrics('FE-004')
    if not m:
        return
    om = m['orbit_means_K']
    bp = m['bus_and_heat_paths']
    gpu = {'caseA': '12 × V100', 'caseB': '12 × A100'}
    s = series('FE-004')
    input_rows = [[gpu[c], fmt(v['solar_irradiance_W_m2'], 2),
                   fmt(min(v['fe']['P_sources_W']), 2) + '–' + fmt(max(v['fe']['P_sources_W']), 2),
                   fmt(v['initial_temperature_K'], 2)] for c, v in s['cases'].items()]
    w.para((f'两组输入均重放周期 {fmt(s["period_s"], 2)} s 的有限元轨迹。下表的热源范围是完整时程的最小值至最大值，运行中使用原时程，并非以区间中值代替。' if lang == 'cn' else
            f'Both inputs replay the FE trajectory with period {fmt(s["period_s"], 2)} s. The heat-source range below is the minimum to maximum of the full history. The simulation uses that history, not its range midpoint.'), keep_next=True)
    w.caption('FE-004 整星场景输入' if lang == 'cn' else 'FE-004 Whole-Satellite Scene Inputs')
    w.table(['算例', '太阳常数 W/m²', '有限元总热源 W', '初始温度 K'] if lang == 'cn' else
            ['Case', 'Solar constant W/m²', 'FE total heat source W', 'Initial temperature K'],
            input_rows, [1900, 2100, 2300, 2009], center_cols=(0, 1, 2, 3))
    rows, rows2 = [], []
    for case in ('caseA', 'caseB'):
        for orb in ('4', '5'):
            v = om[case][orb]
            rows.append([gpu[case], orb, fmt(v['fe_baseplate_K'] - 273.15, 2), fmt(v['J_K'] - 273.15, 2),
                         fmt(v['J_K'] - v['fe_baseplate_K'], 2), fmt(v['fe_radiator_K'] - 273.15, 2),
                         fmt(v['R_K'] - 273.15, 2), fmt(v['R_K'] - v['fe_radiator_K'], 2)])
            b = bp[case][orb]
            fb = b['fe_budget']
            rows2.append([gpu[case], orb, f'{fb["P_sources_W"]:.0f}', f'{fb["heat_to_radiator_W"]:.0f}',
                          f'{fb["heat_other_surfaces_W"]:.0f}', fmt(100 * fb['heat_other_surfaces_W'] / fb['P_sources_W'])])
    if lang == 'cn':
        w.para(f'表 {w.next_number("table")} 与图 {w.next_number("figure")} 给出第四、五圈的平均温度对比，表 {w.next_number("table") + 1} 给出有限元能量收支剩余项。', keep_next=True)
        w.caption('FE-004 第四、五圈平均温度对比')
        w.table(['算例', '圈', 'GPU 基板 有限元 °C', '计算节点 热模块 °C', '差 K', '散热板 有限元 °C', '散热板 热模块 °C', '差 K'], rows,
                [1250, 560, 1100, 1100, 900, 1100, 1100, 1199], center_cols=tuple(range(8)))
        w.picture(OUT / 'fig_fe004_chain_cn.png', 14.6)
        w.caption('FE-004 GPU 基板、计算节点与散热板温度时程', kind='figure')
        w.para('上、下子图分别为 12 块 V100 与 12 块 A100，两者是不同 GPU 配置，采用本项目独立的轨道与功率时程。横轴为从仿真开始计的轨道周期数，纵轴为 °C。蓝色比较有限元 GPU 基板与热模块 J，绿色比较有限元散热板与热模块 R；实线为有限元，虚线为热模块。浅色区域 3–5 个周期对应第四、五圈验收窗口，不表示日食。判定采用窗口内各圈的平均温度。GPU 基板是芯片安装和传热的底板，曲线不代表芯片结温；两个模型也不包含完全相同的辐射表面。')
        w.caption('FE-004 有限元能量收支剩余项')
        w.table(['算例', '圈', '热源 W', '四个主表面及蓄热项 W', '剩余项 W', '占热源 %'], rows2,
                [1400, 600, 1300, 2100, 1500, 1409], center_cols=tuple(range(6)))
        w.para('剩余项包含四个散热板主表面以外的多种表面贡献，包括散热板侧面与其他组件表面。现有数据不能分离这些热量，'
               '因此不能把剩余项全部归入计算节点，不能用其修正温度，也不能据此反推并验证链路热阻。')
    else:
        w.para(f'Table {w.next_number("table")} and Figure {w.next_number("figure")} compare the mean temperatures of orbits 4 '
               f'and 5; Table {w.next_number("table") + 1} gives the FE energy-budget remainder.', keep_next=True)
        w.caption('FE-004 mean temperatures of orbits 4 and 5')
        w.table(['Case', 'Orb.', 'GPU baseplate FE °C', 'Node J module °C', 'Diff. K', 'Radiator FE °C',
                 'Radiator module °C', 'Diff. K'], rows,
                [1250, 560, 1100, 1100, 900, 1100, 1100, 1199], center_cols=tuple(range(8)))
        w.picture(OUT / 'fig_fe004_chain_en.png', 14.6)
        w.caption('FE-004 GPU baseplate, computing node and radiator temperature histories', kind='figure')
        w.para('The upper and lower panels contain 12 V100 and 12 A100 devices: different GPU configurations using this project’s separate orbital and power histories. The horizontal axis counts orbit periods from simulation start; temperature is in °C. Blue compares FE GPU baseplate with module J, and green compares FE radiator with module R. Solid curves are FE and dashed curves are the module. Shading from 3–5 periods marks the fourth- and fifth-orbit acceptance window, not eclipse. Acceptance uses each orbit mean. The baseplate mounts the chip and conducts its heat; its temperature is not GPU junction temperature. The models also do not include identical radiating surfaces.')
        w.caption('FE-004 FE energy-budget remainder')
        w.table(['Case', 'Orb.', 'Sources W', 'Four main faces and storage W', 'Remainder W', 'Share %'], rows2,
                [1400, 600, 1300, 2100, 1500, 1409], center_cols=tuple(range(6)))
        w.para('The remainder includes several surface contributions outside the four main radiator faces, including '
               'radiator edges and other components. Available exports cannot separate these contributions. The entire '
               'remainder cannot be assigned to the computing node, used to correct temperatures, or used to infer and '
               'validate the chain resistance.')


BLOCKS = {'EN-003': block_en003, 'FE-001': block_fe001, 'FE-002': block_fe002, 'FE-003': block_fe003,
          'FE-004': block_fe004}


# =============================================================== summary texts
def short_result(lang, cid):
    """Result and evidence column of the summary table."""
    if cid not in RESULTS:
        return '未执行' if lang == 'cn' else 'Not executed'
    n, ok, bad = check_counts(cid)
    key = SHORT.get(cid, {}).get(lang, '')
    if lang == 'cn':
        head = f'{n} 项检查全部通过' if not bad else f'{n} 项检查中 {bad} 项未通过'
        return f'{head}；{key}。证据 tests/results/{cid}.json' if key else f'{head}。证据 tests/results/{cid}.json'
    head = f'All {n} checks passed' if not bad else f'{bad} of {n} checks failed'
    return f'{head}; {key}. Evidence: tests/results/{cid}.json' if key else f'{head}. Evidence: tests/results/{cid}.json'


SHORT = {
    'DM-001': {'cn': '四个数据对象与表 6 一致，63 项异常构造全部被拒绝',
               'en': 'the four data objects match Table 6 and all 63 invalid constructions are rejected'},
    'PA-001': {'cn': '热容与热阻最大相对误差 1.0×10^{−16}，算例 1000 J/K、108.7 kJ/K 与 7.2 kJ·m^{−2}·K^{−1} 与验收值一致',
               'en': 'largest relative error of capacitances and resistances 1.0×10^{−16}; the examples give 1000 J/K, '
                     '108.7 kJ/K and 7.2 kJ·m^{−2}·K^{−1} as required'},
    'EN-001': {'cn': 'cos_incidence 最大误差 7.8×10^{−16}，G 与 Orbit 输出逐位相同',
               'en': 'largest cos_incidence error 7.8×10^{−16}; G identical to the Orbit output'},
    'EN-002': {'cn': '吸热与辐射最大相对误差 1.9×10^{−16}，第 4.4 节算例为 1600 W',
               'en': 'largest relative error 1.9×10^{−16}; the Section 4.4 example gives 1600 W'},
    'EN-003': {'cn': '微弱反照存在超过 1% 的相对误差，原始有限元 3% 对比未全部通过；三工况逐级加密结果另列',
               'en': 'weak-albedo relative errors exceed 1%; the original FE comparison does not fully meet 3%; '
                     'three-case refinement results are tabulated separately'},
    'HT-001': {'cn': '热流最大相对误差 9.7×10^{−17}，电池算例为 6 W 与 −6 W',
               'en': 'largest relative error of heat flows 9.7×10^{−17}; the battery example gives 6 W and −6 W'},
    'HT-002': {'cn': '电池算例温升速度 0.004 K/s，T1 残差与输入总功率之比最大 8.7×10^{−16}',
               'en': 'battery example 0.004 K/s; largest T1 residual ratio 8.7×10^{−16}'},
    'HT-003': {'cn': '阶跃响应最大误差 0.0000018 K，两个稳态算例误差低于 0.001 K',
               'en': 'largest step-response error 0.0000018 K; both steady cases within 0.001 K'},
    'EC-001': {'cn': '每个功率只进入对应组件，导数变化相对误差最大 6.0×10^{−16}，电池温度只在热模块中更新',
               'en': 'each power enters only its component; largest relative error of derivative changes 6.0×10^{−16}; '
                     'battery temperature updated only by the thermal module'},
    'FE-001': {'cn': '统计温度最大差 1.69 K；冷与平均工况逐点温差 9.18 K 与 12.09 K，超过 3 K',
               'en': 'temperature statistics within 1.69 K; cold and mean pointwise maxima 9.18 K and 12.09 K exceed 3 K'},
    'FE-002': {'cn': '第三圈平均温度差绝对值最大 3.39 K，门限 5 K',
               'en': 'largest absolute third-orbit mean difference 3.39 K against 5 K'},
    'FE-003': {'cn': '14 台设备平均温度差最大 0.25 K，门限 1 K',
               'en': 'largest mean temperature difference of 14 units 0.25 K against 1 K'},
    'FE-004': {'cn': '计算节点比 GPU 基板高 7.9 至 12.5 K，散热板比有限元高 5.9 至 9.2 K，超过 5 K，表面范围差异尚未分离验证',
               'en': 'computing node 7.9 to 12.5 K above the GPU baseplate and radiator 5.9 to 9.2 K above the FE '
                     'against 5 K; effects of different surface coverage have not been isolated'},
    'NI-001': {'cn': '容限收紧 100 倍后变化 5.5×10^{−4} K，RK45 与 Radau 相差 5.8×10^{−4} K，日食时刻误差 3.6×10^{−12} s',
               'en': 'change 5.5×10^{−4} K for 100 times tighter tolerances, RK45 and Radau within 5.8×10^{−4} K, '
                     'eclipse times within 3.6×10^{−12} s'},
    'NI-002': {'cn': '24 h 联合运行完成，两次运行逐位相同，事件与无效结果按第 7 章处理',
               'en': '24 h coupled run completed, two runs identical bit for bit, events and invalid results handled as '
                     'in Chapter 7'},
    'NI-003': {'cn': '面积与质量最大相对误差 1.4×10^{−16}，第 2 至 5 步检查全部通过',
               'en': 'largest relative error of areas and masses 1.4×10^{−16}; all checks of steps 2 to 5 passed'},
}


def counts():
    st = [RESULTS[c]['status'] for c in ORDER if c in RESULTS]
    return len(st), st.count('pass'), st.count('fail'), st.count('partial')


# =============================================================== build
def build(lang, out_name=None):
    C = content(lang)
    t = L[lang]
    n_exec, n_pass, n_fail, n_part = counts()
    failed = [c for c in ORDER if c in RESULTS and RESULTS[c]['status'] == 'fail']
    if lang == 'cn':
        abstract = (f'本报告按《SDTwin 热模块软件模块设计报告》规定热模块的测试范围、测试用例、参照数据与验收判据，并记录 {n_exec} 个用例的执行结果：'
                    f'{n_pass} 个通过，{n_fail} 个不通过。')
    else:
        abstract = (f'This report defines the verification scope, procedures, reference data and acceptance criteria of the '
                    f'SDTwin thermal module according to its design report, and records the execution results of {n_exec} '
                    f'cases: {n_pass} passed and {n_fail} failed.')
    doc = docx.Document(str(TEMPLATE[lang]))
    body_end, final = prepare(doc, lang, abstract)
    w = Writer(doc, lang)
    w.anchor = body_end

    # ------------------------------------------------------------ 1 overview
    w.heading(1, t['h'][0])
    w.heading(2, t['h'][1])
    if lang == 'cn':
        w.field('平台名称：', 'SDTwin。')
        w.field('被测模块：', '热模块 Thermal。')
        w.field('文档名称：', 'SDTwin 热模块软件模块测试报告。')
        w.field('文档用途：', '按《SDTwin 热模块软件模块设计报告》规定热模块各项计算功能的测试方法、参照数据和验收判据，并记录全部用例的执行结果。')
    else:
        w.field('Platform name: ', 'SDTwin.')
        w.field('Module under test: ', 'Thermal Module.')
        w.field('Document title: ', 'Test Report for Thermal Module of SDTwin.')
        w.field('Document purpose: ', 'Define repeatable functional and numerical verification of the thermal module '
                'against the SDTwin Thermal Module Design Report and record the execution results of all cases.')
    w.heading(2, t['h'][2])
    if lang == 'cn':
        w.para('本报告面向热设计、供电与计算模块开发人员，系统测试人员，以及负责资产配置和模型集成的人员。被测对象是设计报告规定的集总热模型，包括太阳能板、计算节点、冷板、电池、'
               '电源设备和公共散热板六个温度节点，SR、JC、CR、BR、DR 五条传热连接，第 5.1 节的四个公共数据对象，以及第 5 章的五个函数：')
        for line in ['assemble_thermal_parameters：从资产与场景装配热容、热阻与表面参数。',
                     'prepare_surface_environment：准备同一时刻各表面的辐照输入。',
                     'calculate_surface_heat：按式 T4 计算太阳能板与散热板的环境吸热和表面辐射。',
                     'calculate_heat_flows：按式 T2 计算五条连接的有符号热流。',
                     'thermal_derivative：按式 T3 计算六个组件的温度导数。']:
            w.bullet(line)
        w.para('测试用例按设计报告第 4 章的式 T1 至 T5、第 5 章的数据对象与函数、第 7 章的运行流程、第 8 章的数值设置和第 9 章的资产装配要求编写，'
               '每个用例在设计依据栏注明对应的章节、公式与函数。热模块已按设计报告实现为 thermal 软件包，四个数据对象与五个函数的名称、字段与调用顺序与设计报告一致。'
               '端到端用例另需积分流程、地球反照与红外计算、场景装配与供电结果，这些测试辅助程序位于 sdtwin_sim 目录。其中测试用供电程序按《SDTwin 供电模块软件模块设计报告》编写，'
               '电池参数为说明性测试值，不代表选定器件。')
        w.para(f'全部 {n_exec} 个用例已于{DATE["cn"]}执行，各用例的实际结果与测试结论见第 3 章，汇总见第 7 章。')
    else:
        w.para('This report is intended for thermal, power and computing developers, system testers, and personnel '
               'responsible for asset configuration and model integration. The object under test is the lumped thermal '
               'model specified by the design report: the six temperature nodes solar array, computing node, cold plate, '
               'battery, power equipment and common radiator, the five heat-transfer connections SR, JC, CR, BR and DR, '
               'the four common data objects of Section 5.1, and the five functions of Chapter 5:')
        for line in ['assemble_thermal_parameters: assembles thermal capacitances, thermal resistances and surface '
                     'parameters from the assets and the scene.',
                     'prepare_surface_environment: prepares the irradiation inputs of every surface at one instant.',
                     'calculate_surface_heat: calculates the environmental absorption and surface emission of the solar '
                     'array and the radiator by Equation T4.',
                     'calculate_heat_flows: calculates the signed heat flows of the five connections by Equation T2.',
                     'thermal_derivative: calculates the temperature derivatives of the six components by Equation T3.']:
            w.bullet(line)
        w.para('The test cases follow Equations T1 to T5 in Chapter 4, the data objects and functions in Chapter 5, the '
               'run procedure in Chapter 7, the numerical settings in Chapter 8 and the asset assembly requirements in '
               'Chapter 9 of the design report; the Design Basis row of each case names the sections, equations and '
               'functions it verifies. The thermal module is implemented as the thermal package of the design report, '
               'with the names, fields and call order of the four data objects and five functions unchanged. The '
               'end-to-end cases also need the coupled integration procedure, the Earth albedo and infrared calculation, '
               'the scene assembly and Power results; these test support programs are in the sdtwin_sim directory. The '
               'Power program used for testing follows the SDTwin Power Module Design Report, and its battery '
               'parameters are illustrative test values, not a selected device.')
        w.para(f'All {n_exec} cases were executed on {DATE["en"]}. The actual result and conclusion of each case are '
               'recorded in Chapter 3 and summarized in Chapter 7.')
    w.heading(2, t['h'][3])
    w.table(['术语', '说明'] if lang == 'cn' else ['Term', 'Definition'], C['abbr'], [2000, 6309], center_cols=(0,))
    if lang == 'cn':
        w.para('模块内部温度使用 K。有限元对比以 °C 给出温度、以 K 给出温差。最低与最高取比较窗口内极值；平均值按时间加权，另有注明时采用采样均值。')
    else:
        w.para('Internal temperatures are in K. FE comparisons use °C for temperatures and K for differences. '
               'Minima and maxima are window extrema; means are time-weighted unless a sample mean is explicitly stated.')

    # ------------------------------------------------------------ 2 test content
    w.heading(1, t['h'][4])
    if lang == 'cn':
        w.para('测试范围覆盖设计报告第 11 章的四项需求 TH-01 至 TH-04，共 16 个用例，按设计报告的章节分为六组。每个用例对应一个被测函数、数据对象或一项系统级对比，并在第 7 章汇总。')
    else:
        w.para('The verification scope covers the four requirements TH-01 to TH-04 in Chapter 11 of the design report '
               'with 16 cases organized into six groups that follow the design report chapters. Each case is traced to '
               'one function, data object or system-level comparison and to one final summary record.')
    for g, ids in C['groups']:
        w.bullet(f'{g}：{"、".join(ids)}。' if lang == 'cn' else f'{g}: {", ".join(ids)}.')
    w.caption('功能需求与测试用例追踪' if lang == 'cn' else 'Functional Requirement and Test-Case Traceability')
    rows = []
    k = 0
    for _, ids in C['groups']:
        for cid in ids:
            k += 1
            c = C['cases'][cid]
            rows.append([str(k), cid, c['func'], c['name'], c['req'], status_word(lang, cid)])
    w.table(['序号', '用例编号', '被测函数或对象', '测试名称', '对应需求', '状态'] if lang == 'cn' else
            ['No.', 'Case ID', 'Function', 'Test Name', 'Requirement', 'Status'], rows,
            [600, 950, 2600, 1800, 1400, 959], center_cols=(0, 1, 4, 5))
    trace_table = w.tab

    env = case_evidence.ENVIRONMENT[lang]
    w.para(env['intro'])
    w.caption(env['caption'])
    w.table(env['headers'], env['rows'], [1550, 650, 1300, 950, 1400, 2459], center_cols=tuple(range(6)))
    for paragraph in env['notes']:
        w.para(paragraph)

    # ------------------------------------------------------------ 3 detailed test projects
    w.heading(1, t['h'][5])
    if lang == 'cn':
        w.para('以下用例采用统一的执行记录格式，设计依据栏注明用例所验证的设计报告内容。验收判据为本版默认门限，模块需求给出更严格的数值时以需求为准。'
               '每个项目先说明场景与输入，在用例表后列出实测结果表；涉及温度时程的对比另给曲线及读图说明。')
    else:
        w.para('The following records use the same execution structure as the reference outline; the Design Basis row '
               'names the design report content each case verifies. The numerical thresholds are default verification '
               'gates and shall be tightened when an approved module requirement defines a stricter value. Each project '
               'explains its scene and inputs and provides a measured result table after the case record. Temperature-history '
               'comparisons also include curves with explicit reading instructions.')
    k = 0
    for g, ids in C['groups']:
        w.heading(2, g)
        if ids and ids[0].startswith('FE-'):
            if lang == 'cn':
                w.para('本组用例把热模块的计算结果与有限元结果比较。热模块用一个温度代表一个组件，比较对象取有限元同一部件的平均温度。有限元结果文件见第 6 章。')
            else:
                w.para('This group compares thermal module results with finite element results. The module represents each '
                       'component by one temperature, so the comparison uses the mean temperature of the same part in '
                       'the finite element model. The finite element result files are listed in Chapter 6.')
            iss_temperature_figure(w, lang, 'nom0_T_iso_noon')
        for cid in ids:
            k += 1
            c = dict(C['cases'][cid], id=cid)
            if cid in RESULTS:
                c.update(case_result_fields(lang, cid))
            else:
                c.update(actual='未执行。' if lang == 'cn' else 'Not executed.',
                         anomaly='无，用例未执行。' if lang == 'cn' else 'None recorded; case not executed.', status='not_run')
            w.heading(3, t['project'].format(k=k, name=c['name']))
            w.para(case_evidence.context(cid, lang))
            w.para(case_evidence.scenario(cid, lang))
            w.caption(t['case_caption'].format(cid=cid, name=c['name']))
            w.case_table(c)
            support = case_evidence.table_spec(cid, lang, metrics(cid))
            if support:
                w.para(('下表把代表性输入、独立参照和实际输出并列，便于核对本用例的判定。' if lang == 'cn' else
                        'The following table aligns representative inputs, independent references and measured outputs for checking the verdict.'), keep_next=True)
                w.caption(support['caption'])
                w.table(support['headers'], support['rows'], [2600, 2800, 2909])
            if cid in BLOCKS and cid in RESULTS:
                BLOCKS[cid](w, lang)
            if cid == 'FE-001':
                iss_temperature_figure(w, lang, 'cold0_T_iso_ecl')
            elif cid == 'FE-002':
                iss_temperature_figure(w, lang, 'hot75_T_hrs_noon')

    # ------------------------------------------------------------ 4 sufficiency
    w.heading(1, t['h'][6])
    t_cov = w.next_number('table')
    if lang == 'cn':
        w.para(f'16 个用例与设计报告逐章对应，对应关系见表 {t_cov}：式 T1 至 T5、第 5 章的四个数据对象与五个函数、第 6 章的接口约定、第 7 章的运行流程、'
               f'第 8 章的数值设置和第 9 章的资产装配与结果归档都有对应用例，四项需求 TH-01 至 TH-04 也都有对应用例，追踪关系见表 {trace_table}。')
        w.caption('设计报告章节与测试用例对应')
        w.table(['设计报告章节', '验证内容', '测试用例'], C['coverage'], [2600, 3100, 2609])
        w.para(f'表 {w.next_number("table")} 列出各项对比使用的参照与执行状态。参照分为设计报告算例、手算、解析解、独立数值计算和有限元结果五类，参照计算不使用热模块代码。')
        w.caption('对比项、参照与执行状态')
    else:
        w.para(f'The 16 cases map to the design report chapter by chapter, as listed in Table {t_cov}: Equations T1 to T5, '
               'the four data objects and five functions of Chapter 5, the interface conventions of Chapter 6, the run '
               'procedure of Chapter 7, the numerical settings of Chapter 8 and the asset assembly and result archiving '
               'of Chapter 9 all have cases, and so do the four requirements TH-01 to TH-04, traced in Table '
               f'{trace_table}.')
        w.caption('Design Report Sections and Test Cases')
        w.table(['Design report section', 'Verified content', 'Test cases'], C['coverage'], [2600, 3100, 2609])
        w.para(f'Table {w.next_number("table")} lists the reference and execution status of each comparison. References are '
               'design report examples, hand calculations, analytic solutions, independent numerical calculations and '
               'finite element results; no reference calculation uses thermal module code.')
        w.caption('Comparisons, References and Execution Status')
    rows = []
    for i, (item, ref, data_, ids) in enumerate(C['compare'], start=1):
        st = [RESULTS[x]['status'] for x in ids if x in RESULTS]
        if not st:
            word = t['status']['not_run']
        elif all(s == 'pass' for s in st):
            word = t['status']['pass']
        elif any(s == 'fail' for s in st):
            word = t['status']['fail']
        else:
            word = t['status']['partial']
        rows.append([str(i), item, ref, data_, f'{word}，{"、".join(ids)}' if lang == 'cn' else f'{word}, {", ".join(ids)}'])
    w.table(['序号', '对比项', '参照', '已有数据', '执行状态'] if lang == 'cn' else
            ['No.', 'Comparison', 'Reference', 'Available data', 'Execution status'], rows,
            [620, 1900, 2250, 2200, 1339], center_cols=(0,))
    fe2, fe4 = metrics('FE-002'), metrics('FE-004')
    gaps = [v['mean_minus_coldest_K'] for v in fe2.get('coldest_panel', {}).values() if isinstance(v, dict)]
    if lang == 'cn':
        w.para('有限元对比覆盖太阳能板、公共散热板、计算节点与冷板，以及 JC 与 CR 连接。电池与电源设备两个组件、SR、BR、DR 三条连接没有有限元参照，'
               '由设计报告算例、手算、解析解与整体热平衡检验。本报告的端到端用例使用测试用供电程序，电池电化学模型与器件参数仍需由 Power 模块另行验证。')
        if gaps:
            w.para(f'执行结果同时给出设计报告第 10 章所列限制的量级：散热器同一时刻最冷面板比平均温度低 {fmt(min(gaps))} 至 {fmt(max(gaps))} K；'
                   '设计热工况有限元回路 B 比回路 A 高约 4.7 K，模块为约 1.6 K。翼间遮挡与沿流向温差的贡献尚未分离。'
                   '整星能量收支在四个散热板主表面及蓄热项以外仍有约 9% 至 13% 的剩余量，现有导出不能将其逐节点分配，也不能据此确定 FE-004 超差的全部原因。')
        w.para(f'全部用例已有执行记录，{n_fail} 个用例未满足验收要求。对应偏差需在修正与复测后重新判定，当前报告保留实际结果。')
    else:
        w.para('The finite element comparisons cover the solar array, the common radiator, the computing node and the cold '
               'plate, and the JC and CR connections. The battery and power equipment and the SR, BR and DR connections '
               'have no finite element reference and are verified with design report examples, hand calculations, '
               'analytic solutions and the overall heat balance. The end-to-end cases use the Power stand-in; the '
               'battery electrochemical model and device parameters still require separate Power-module validation.')
        if gaps:
            w.para('The results also quantify the limitations listed in Chapter 10 of the design report: at the same '
                   f'instant the coldest radiator panel is {fmt(min(gaps))} to {fmt(max(gaps))} K below the radiator mean; '
                   'FE loop B is about 4.7 K warmer than loop A in the design hot case, against about 1.6 K in the module. '
                   'The contributions of wing obstruction and along-flow gradients are not isolated. The whole-satellite '
                   'energy budget has a roughly 9% to 13% remainder beyond four main radiator faces and storage. Available '
                   'exports cannot allocate it to individual nodes or establish the complete cause of FE-004 exceedances.')
        w.para(f'Every case has an execution record; {n_fail} cases do not meet acceptance. Their discrepancies require '
               'correction and repeat testing before reassessment. This report retains the measured results.')

    # ------------------------------------------------------------ 5 conditions
    conditions_heading = w.heading(1, t['h'][7])
    if lang == 'en':
        ppr(conditions_heading).append(w_el('w:pageBreakBefore'))
    for line in C['conditions']:
        w.bullet(line)
    if RESULTS:
        env = next(iter(RESULTS.values()))['environment']
        if lang == 'cn':
            w.para(f'本次执行环境为 Windows 11 家庭版，Python {env["python"]}，NumPy {env["numpy"]}，SciPy {env["scipy"]}，Astropy {env["astropy"]}，'
                   f'pytest {env["pytest"]}，Orbit 软件包 ntu_space_dynamics {env["ntu_space_dynamics"]}，OpenUSD {env["usd"]}；有限元热载荷用 COMSOL 6.3 与 MPh 导出。'
                   '在 Thermal 目录执行 .venv/Scripts/python.exe -m pytest tests -q 运行全部用例。最终整体验证输出保存在 tests/results/pytest_verified_run.txt，'
                   'COMSOL 加密计算完成后的 EN-003 复测输出保存在 tests/results/pytest_en003_final.txt。最终证据清单 verification_manifest.json 记录每个用例的执行时间、完整性与文件散列。'
                   '每个用例的检查项、关键数值、结果说明与异常记录保存在 tests/results 下以用例编号命名的结果文件中。')
        else:
            w.para(f'The execution environment was Windows 11 Home, Python {env["python"]}, NumPy {env["numpy"]}, SciPy '
                   f'{env["scipy"]}, Astropy {env["astropy"]}, pytest {env["pytest"]}, the Orbit package ntu_space_dynamics '
                   f'{env["ntu_space_dynamics"]} and OpenUSD {env["usd"]}; finite element thermal loads were exported with '
                   'COMSOL 6.3 and MPh. All cases are run from the Thermal directory with .venv/Scripts/python.exe -m '
                   'pytest tests -q. The final suite log is tests/results/pytest_verified_run.txt; the EN-003 rerun after '
                   'COMSOL refinement is tests/results/pytest_en003_final.txt. The final verification_manifest.json records '
                   'case execution times, completeness and file hashes. The checks, key '
                   'values, result text and anomalies of each case are retained in its named result file under '
                   'tests/results.')

    # ------------------------------------------------------------ 6 test data
    data_heading = w.heading(1, t['h'][8])
    ppr(data_heading).append(w_el('w:pageBreakBefore'))
    if lang == 'cn':
        w.para('测试输入与结果保存在项目目录，证据清单记录文件散列，数据记录注明来源、适用范围、时刻与坐标约定。')
        w.caption('测试数据清单')
        headers = ['序号', '测试数据', '数据类型', '形式', '来源']
    else:
        w.para('Test inputs and results are retained in the project directory. The evidence manifest records file hashes; '
               'data records state origin, validity range, time and frame conventions.')
        w.caption('Test Data Inventory')
        headers = ['No.', 'Test Data', 'Data Type', 'Form', 'Source']
    w.table(headers, [[str(i)] + list(r) for i, r in enumerate(C['data'], start=1)], [650, 2000, 1300, 1100, 3259],
            center_cols=(0, 2, 3), group_rows=len(C['data']))

    # ------------------------------------------------------------ 7 summary, landscape section
    w.anchor = final
    w.heading(1, t['h'][9])
    if lang == 'cn':
        w.para(f'全部 {n_exec} 个用例已执行，{n_pass} 个通过，{n_fail} 个不通过。状态按验收判据与保存的执行证据确定。')
        w.caption('热模块测试总结')
        headers = ['序号', '用例编号', '被测函数或对象', '验证重点', '状态', '结果与证据']
    else:
        w.para(f'All {n_exec} cases were executed: {n_pass} passed and {n_fail} failed. Each status was evaluated against the '
               'acceptance criteria with retained execution evidence.')
        w.caption('Thermal Module Test Summary')
        headers = ['No.', 'Case ID', 'Function', 'Verification Focus', 'Status', 'Result / Evidence']
    rows = []
    k = 0
    for _, ids in C['groups']:
        for cid in ids:
            k += 1
            c = C['cases'][cid]
            rows.append([str(k), cid, c['func'], c['focus'], status_word(lang, cid), short_result(lang, cid)])
    w.table(headers, rows, [650, 1000, 2600, 2700, 1000, 6004], center_cols=(0, 1, 4), padding_y=45)
    if failed:
        notes = ''.join(FAIL_NOTES[c][lang] for c in failed if c in FAIL_NOTES) + FAIL_CLOSE[lang]
        if lang == 'cn':
            w.para(f'不通过的用例为 {"、".join(failed)}。' + notes, before=120)
        else:
            w.para(f'The failed cases are {", ".join(failed)}. ' + notes, before=120)
    overall = 'pass' if n_exec == len(ORDER) and n_fail == 0 and n_part == 0 else ('fail' if n_fail else 'partial')
    if lang == 'cn':
        w.para('总体结论：' + conclusion_marks('cn', overall), indent=False, before=240)
        w.para('现场测试专家签字：____________________　日期：____________', indent=False)
    else:
        w.para('**Overall Conclusion:** ' + conclusion_marks('en', overall), indent=False, before=240)
        w.para('On-site Test Expert Signature: ______________________________    Date: __________________', indent=False)

    name = out_name or OUT_NAME[lang]
    path = Path(name).with_suffix('.docx') if ('/' in name or '\\' in name) else THERMAL / f'{name}.docx'
    doc.save(str(path))
    print('tables', w.tab, 'figures', w.fig, '->', path)


FAIL_NOTES = {
    'FE-001': {
        'cn': 'FE-001 的统计温度达到判据，但同一时刻的曲线对比存在超差，需在一致的日食处理与时间分辨率下继续核对。',
        'en': 'FE-001 meets the temperature-statistic criterion but fails simultaneous curve comparison; eclipse treatment and time resolution require further matched checks. ',
    },
    'FE-004': {
        'cn': 'FE-004 的两种模型使用的辐射表面范围不同，现有能量收支不能分离每个节点的贡献，不能据此认定热阻或温度已得到验证。',
        'en': 'FE-004 uses different radiating-surface coverage in the two models. The available energy budget cannot isolate each node contribution and does not establish validation of resistance or temperature. ',
    },
    'EN-003': {
        'cn': 'EN-003 保留原有限元对比的超限项，并分别记录微弱辐照下的积分精度和地球离散加密结果，原因分析不替代验收。',
        'en': 'EN-003 retains original FE exceedances and separately records weak-flux integration accuracy and planet-refinement results; diagnosis does not replace acceptance. ',
    },
}
FAIL_CLOSE = {
    'cn': '公式、接口或积分检查的通过不能替代有限元温度与环境热载荷验收。当前结果支持的范围以各用例记录为准。',
    'en': 'Passing formula, interface or integration checks does not replace FE temperature and environmental-load acceptance. Supported claims are limited to each case record.',
}


if __name__ == '__main__':
    lang = sys.argv[1] if len(sys.argv) > 1 else 'cn'
    build(lang, sys.argv[2] if len(sys.argv) > 2 else None)
