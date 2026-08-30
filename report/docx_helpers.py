"""Thin python-docx helpers for a Chinese engineering design report."""
from __future__ import annotations

import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_BODY = "宋体"
CN_HEAD = "黑体"
EN_BODY = "Times New Roman"
MATH = "Cambria Math"


def set_font(run, cn=CN_BODY, en=EN_BODY, size=12, bold=None, italic=None, color=None):
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:cs"), en)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _style_fonts(style, cn, en, size, bold=False):
    style.font.name = en
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:cs"), en)

_MARK = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\})")


def add_marked_runs(para, text, cn, size, bold=False):
    """Writes `text` into `para`, rendering _{x} as subscript and ^{x} as superscript."""
    for piece in _MARK.split(text):
        if not piece:
            continue
        if piece.startswith("_{"):
            run = para.add_run(piece[2:-1]); set_font(run, cn=cn, size=size, bold=bold)
            run.font.subscript = True
        elif piece.startswith("^{"):
            run = para.add_run(piece[2:-1]); set_font(run, cn=cn, size=size, bold=bold)
            run.font.superscript = True
        else:
            run = para.add_run(piece); set_font(run, cn=cn, size=size, bold=bold)


class Report:
    def __init__(self, lang="zh"):
        self.lang = lang
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        self._setup()

    # ------------------------------------------------------------------ setup
    def _setup(self):
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.6)
        sec.top_margin = Cm(2.8)
        sec.bottom_margin = Cm(2.5)

        normal = doc.styles["Normal"]
        _style_fonts(normal, CN_BODY, EN_BODY, 12)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5 if self.lang == "zh" else 1.3
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)

        for lvl, size, before, after in ((1, 16, 18, 10), (2, 14, 12, 6), (3, 12, 8, 4)):
            st = doc.styles[f"Heading {lvl}"]
            _style_fonts(st, CN_HEAD, EN_BODY, size, bold=True)
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.line_spacing = 1.3

        # footer page number
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(p, "PAGE", size=10.5)
        # ask Word to refresh fields (table of contents) on open
        settings = doc.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)

    @staticmethod
    def _field(paragraph, instr, size=12):
        run = paragraph.add_run()
        set_font(run, size=size)
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
        t = OxmlElement("w:t"); t.text = " "
        f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
        for el in (f1, it, f2, t, f3):
            run._element.append(el)

    # --------------------------------------------------------------- blocks
    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h(self, level, text):
        p = self.doc.add_heading(text, level=level)
        en = EN_BODY if self.lang == "zh" else "Arial"
        for r in p.runs:
            set_font(r, cn=CN_HEAD, en=en, size={1: 16, 2: 14, 3: 12}[level], bold=True, color="000000")
        return p

    def p(self, text, indent=True, align=None, size=12, bold=False, cn=CN_BODY, after=0):
        para = self.doc.add_paragraph()
        if indent and self.lang == "zh":
            para.paragraph_format.first_line_indent = Pt(size * 2)
        elif indent:
            para.paragraph_format.space_after = Pt(6)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if after:
            para.paragraph_format.space_after = Pt(after)
        add_marked_runs(para, text, cn, size, bold)
        return para

    def item(self, label, text, size=12):
        """Numbered item like '指标 1  ...' with a hanging indent."""
        para = self.doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(size * 2)
        para.paragraph_format.first_line_indent = Pt(-size * 2)
        r1 = para.add_run(label + "　")
        set_font(r1, size=size, bold=True)
        r2 = para.add_run(text)
        set_font(r2, size=size)
        return para

    def formula(self, text, size=12):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(3)
        para.paragraph_format.space_after = Pt(3)
        run = para.add_run(text)
        set_font(run, cn=MATH, en=MATH, size=size)
        return para

    def caption(self, text, kind):
        if kind == "fig":
            self.fig_no += 1
            label = (f"图 {self.fig_no}  {text}" if self.lang == "zh" else f"Figure {self.fig_no}  {text}")
        else:
            self.tab_no += 1
            label = (f"表 {self.tab_no}  {text}" if self.lang == "zh" else f"Table {self.tab_no}  {text}")
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.keep_with_next = (kind == "tab")
        run = para.add_run(label)
        set_font(run, cn=CN_HEAD, size=10.5, bold=True)
        return label

    def figure(self, path, caption, width_cm=15.0):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(path, width=Cm(width_cm))
        self.caption(caption, "fig")
        return self.fig_no

    # --------------------------------------------------------------- tables
    @staticmethod
    def _shade(cell, hex_fill):
        tcpr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcpr.append(shd)

    @staticmethod
    def _cell_text(cell, text, size=10.5, bold=False, align=None, cn=CN_BODY):
        cell.text = ""
        lines = str(text).split("\n")
        for i, line in enumerate(lines):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            para.paragraph_format.line_spacing = 1.15
            para.paragraph_format.space_after = Pt(1)
            if align == "center":
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_marked_runs(para, line, cn, size, bold)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def table(self, caption, header, rows, widths_cm=None, size=10.5, center_cols=(),
              header_fill="E7ECF2"):
        if caption:
            self.caption(caption, "tab")
        n_cols = len(header)
        t = self.doc.add_table(rows=1 + len(rows), cols=n_cols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for j, htxt in enumerate(header):
            c = t.rows[0].cells[j]
            self._cell_text(c, htxt, size=size, bold=True, align="center", cn=CN_HEAD)
            self._shade(c, header_fill)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                self._cell_text(t.rows[i + 1].cells[j], val, size=size,
                                align="center" if j in center_cols else None)
        if widths_cm:
            for row in t.rows:
                for j, w in enumerate(widths_cm):
                    row.cells[j].width = Cm(w)
        # repeat header row
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def spec(self, caption, rows, label_w=2.6, size=10.5):
        """Two-column model specification table: label | content."""
        if caption:
            self.caption(caption, "tab")
        t = self.doc.add_table(rows=len(rows), cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        total = 15.8
        for i, (lab, content) in enumerate(rows):
            c0, c1 = t.rows[i].cells
            self._cell_text(c0, lab, size=size, bold=True, align="center", cn=CN_HEAD)
            self._shade(c0, "E7ECF2")
            self._cell_text(c1, content, size=size)
            c0.width = Cm(label_w)
            c1.width = Cm(total - label_w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def kv_table(self, rows, widths=(5.0, 10.8), size=10.5):
        """Plain two-column table without caption (cover page use)."""
        t = self.doc.add_table(rows=len(rows), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (a, b) in enumerate(rows):
            self._cell_text(t.rows[i].cells[0], a, size=size, bold=True, align="center", cn=CN_HEAD)
            self._cell_text(t.rows[i].cells[1], b, size=size, align="center")
            t.rows[i].cells[0].width = Cm(widths[0])
            t.rows[i].cells[1].width = Cm(widths[1])
        return t

    def toc(self):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run("目　录" if self.lang == "zh" else "Contents")
        set_font(run, cn=CN_HEAD, size=16, bold=True)
        para.paragraph_format.space_after = Pt(12)
        p2 = self.doc.add_paragraph()
        self._field(p2, 'TOC \\o "1-3" \\h \\z \\u', size=12)

    def save(self, path):
        self.doc.save(path)
