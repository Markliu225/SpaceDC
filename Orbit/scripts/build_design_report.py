"""Build the English NTU Space Dynamics software module design report."""

from __future__ import annotations

import math
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "【1】面向群体协同的自主任务规划计算服务平台总体设计报告.docx"
FINAL = ROOT / "docs" / "NTU_Space_Dynamics_Software_Module_Design_Report.docx"
ASSET_DIR = ROOT / ".tmp" / "design_report" / "assets"
END_TO_END_FIGURE = ROOT / "outputs" / "end_to_end" / "end_to_end_timeseries.png"
REPORT_END_TO_END_FIGURE = ASSET_DIR / "end_to_end_timeseries_rgb.png"

BLACK = "000000"
DARK_GRAY = "404040"
MID_GRAY = "808080"
LIGHT_GRAY = "E7E6E6"
VERY_LIGHT_GRAY = "F4F4F4"
PALE_BLUE = "D9EAF2"
PALE_GREEN = "E2EAD4"
PALE_GOLD = "F5E7C6"
PALE_ROSE = "EED5D2"

PORTRAIT_WIDTH_IN = 8.27
PORTRAIT_HEIGHT_IN = 11.69
PORTRAIT_CONTENT_IN = 5.77
LANDSCAPE_CONTENT_IN = 9.69


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_in: float) -> None:
    width_dxa = int(round(width_in * 1440))
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_in)


def set_table_geometry(table, widths_in: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total_dxa = int(round(sum(widths_in) * 1440))
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(round(width * 1440))))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if not row._tr.xpath("./w:trPr/w:cantSplit"):
            tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_in[min(index, len(widths_in) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header_and_keep(row) -> None:
    repeat_table_header(row)
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True


def set_run_font(run, name="Times New Roman", size=None, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def clear_header_footer(container) -> None:
    for paragraph in container.paragraphs:
        clear_paragraph(paragraph)
    for table in list(container.tables):
        table._element.getparent().remove(table._element)


def clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_page_border(section, enabled: bool) -> None:
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:pgBorders"))
    if old is not None:
        sect_pr.remove(old)
    if not enabled:
        return
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "18")
        border.set(qn("w:space"), "28")
        border.set(qn("w:color"), BLACK)
        borders.append(border)
    sect_pr.append(borders)


def configure_portrait(section, *, cover=False) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(PORTRAIT_WIDTH_IN)
    section.page_height = Inches(PORTRAIT_HEIGHT_IN)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    set_page_border(section, cover)


def configure_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(PORTRAIT_HEIGHT_IN)
    section.page_height = Inches(PORTRAIT_WIDTH_IN)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.25)
    section.bottom_margin = Inches(1.25)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    set_page_border(section, False)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_page_number_start(section, value: int) -> None:
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        sect_pr.append(node)
    node.set(qn("w:start"), str(value))


def add_running_footer(section, *, restart=None) -> None:
    section.footer.is_linked_to_previous = False
    clear_header_footer(section.footer)
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    add_field(paragraph, "PAGE", "1")
    for run in paragraph.runs:
        set_run_font(run, size=9, color=DARK_GRAY)
    if restart is not None:
        set_page_number_start(section, restart)
    else:
        page_numbering = section._sectPr.find(qn("w:pgNumType"))
        if page_numbering is not None:
            section._sectPr.remove(page_numbering)


def add_running_header(section) -> None:
    section.header.is_linked_to_previous = False
    clear_header_footer(section.header)
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("NTU Space Dynamics — Software Module Design Report")
    set_run_font(run, size=8.5, color=MID_GRAY)


def ensure_style(document, name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Inches(0.25)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3

    sizes = (15, 13, 11.5)
    before = (14, 12, 10)
    after = (8, 6, 4)
    for index in range(1, 4):
        style = document.styles[f"Heading {index}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(sizes[index - 1])
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.left_indent = Inches(0)
        style.paragraph_format.space_before = Pt(before[index - 1])
        style.paragraph_format.space_after = Pt(after[index - 1])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = ensure_style(document, "Engineering Caption")
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Inches(0)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True

    table_text = ensure_style(document, "Engineering Table Text")
    table_text.font.name = "Times New Roman"
    table_text._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    table_text._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    table_text.font.size = Pt(9)
    table_text.paragraph_format.first_line_indent = Inches(0)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.12

    code = ensure_style(document, "Code Block")
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.paragraph_format.first_line_indent = Inches(0)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0

    equation = ensure_style(document, "Engineering Equation")
    equation.font.name = "Cambria Math"
    equation._element.rPr.rFonts.set(qn("w:ascii"), "Cambria Math")
    equation._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria Math")
    equation.font.size = Pt(10.5)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.first_line_indent = Inches(0)
    equation.paragraph_format.space_before = Pt(5)
    equation.paragraph_format.space_after = Pt(5)
    equation.paragraph_format.keep_together = True

    note = ensure_style(document, "Engineering Note")
    note.font.name = "Times New Roman"
    note._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    note._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    note.font.size = Pt(10)
    note.paragraph_format.first_line_indent = Inches(0)
    note.paragraph_format.left_indent = Inches(0.22)
    note.paragraph_format.right_indent = Inches(0.22)
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(6)
    note.paragraph_format.line_spacing = 1.18

    appendix = ensure_style(document, "Appendix Heading")
    appendix.base_style = document.styles["Normal"]
    appendix.font.name = "Times New Roman"
    appendix._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    appendix._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    appendix._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    appendix.font.size = Pt(15)
    appendix.font.bold = True
    appendix.font.color.rgb = RGBColor(0, 0, 0)
    appendix.paragraph_format.first_line_indent = Inches(0)
    appendix.paragraph_format.left_indent = Inches(0)
    appendix.paragraph_format.space_before = Pt(14)
    appendix.paragraph_format.space_after = Pt(8)
    appendix.paragraph_format.keep_with_next = True
    appendix.paragraph_format.keep_together = True
    appendix.paragraph_format.outline_level = 0


def add_numbering_definition(document, *, kind: str, levels: int = 1) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel" if levels > 1 else "singleLevel")
    abstract.append(multi)
    for level_index in range(levels):
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), str(level_index))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        if kind == "heading":
            paragraph_style = OxmlElement("w:pStyle")
            paragraph_style.set(qn("w:val"), f"Heading{level_index + 1}")
            level.append(paragraph_style)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        level.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        if kind == "bullet":
            lvl_text.set(qn("w:val"), "•")
        elif levels > 1:
            lvl_text.set(
                qn("w:val"),
                ".".join(f"%{i + 1}" for i in range(level_index + 1)),
            )
        else:
            lvl_text.set(qn("w:val"), "%1.")
        level.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "space")
        level.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(360 + level_index * 360))
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        left = 720 + level_index * 360
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        level.append(p_pr)
        if kind == "bullet":
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), "Arial")
            r_fonts.set(qn("w:hAnsi"), "Arial")
            r_pr.append(r_fonts)
            level.append(r_pr)
        abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def attach_num(paragraph, num_id: int, level: int = 0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    old = p_pr.find(qn("w:numPr"))
    if old is not None:
        p_pr.remove(old)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def configure_heading_numbering(document) -> tuple[int, int]:
    heading_num = add_numbering_definition(document, kind="heading", levels=3)
    for level in range(3):
        style = document.styles[f"Heading {level + 1}"]
        p_pr = style._element.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(heading_num))
        num_pr.extend([ilvl, num])
        old = p_pr.find(qn("w:numPr"))
        if old is not None:
            p_pr.remove(old)
        p_pr.append(num_pr)
    return add_numbering_definition(document, kind="bullet"), add_numbering_definition(
        document, kind="decimal"
    )


def add_heading(document, text: str, level: int = 1):
    return document.add_paragraph(text, style=f"Heading {level}")


def add_body(document, text: str, *, first_indent=True, keep=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Inches(0.25 if first_indent else 0)
    paragraph.paragraph_format.keep_together = keep
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    return paragraph


def add_note(document, label: str, text: str):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [PORTRAIT_CONTENT_IN])
    cell = table.cell(0, 0)
    shade_cell(cell, VERY_LIGHT_GRAY)
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["Engineering Note"]
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=10, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_bullets(document, items: list[str], bullet_num: int):
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.2
        attach_num(paragraph, bullet_num, 0)
        run = paragraph.add_run(item)
        set_run_font(run, size=10.5)


def add_numbered_steps(document, items: list[str], decimal_num: int):
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.2
        attach_num(paragraph, decimal_num, 0)
        run = paragraph.add_run(item)
        set_run_font(run, size=10.5)


def add_equation(document, equation: str):
    paragraph = document.add_paragraph(style="Engineering Equation")
    run = paragraph.add_run(equation)
    set_run_font(run, name="Cambria Math", size=10.5)
    return paragraph


def add_code(document, code: str):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [PORTRAIT_CONTENT_IN])
    cell = table.cell(0, 0)
    shade_cell(cell, VERY_LIGHT_GRAY)
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["Code Block"]
    run = paragraph.add_run(code.rstrip())
    set_run_font(run, name="Consolas", size=8.3)
    paragraph.paragraph_format.keep_together = False
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(
    document,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    widths_in: list[float],
    *,
    font_size=9,
    header_fill=LIGHT_GRAY,
):
    caption_p = document.add_paragraph(style="Engineering Caption")
    caption_p.paragraph_format.keep_with_next = True
    run = caption_p.add_run(caption)
    set_run_font(run, size=10)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_in)
    header = table.rows[0]
    set_repeat_table_header_and_keep(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        shade_cell(cell, header_fill)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.style = document.styles["Engineering Table Text"]
        run = paragraph.add_run(text)
        set_run_font(run, size=font_size, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Engineering Table Text"]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if index == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(str(value))
            set_run_font(run, size=font_size)
    set_table_geometry(table, widths_in)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_picture(document, path: Path, caption: str, width_in: float):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(width_in))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption.split("  ", 1)[0])
    caption_p = document.add_paragraph(style="Engineering Caption")
    cap_run = caption_p.add_run(caption)
    set_run_font(cap_run, size=10)


def font(size: int, *, bold=False):
    paths = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/timesbd.ttf" if bold else "C:/Windows/Fonts/times.ttf"),
    ]
    for candidate in paths:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def draw_centered_text(draw, box, text, font_obj, fill=(0, 0, 0), spacing=6):
    x0, y0, x1, y1 = box
    lines = text.split("\n")
    heights = [draw.textbbox((0, 0), line, font=font_obj)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = (y0 + y1 - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font_obj)
        width = bbox[2] - bbox[0]
        draw.text(((x0 + x1 - width) / 2, y), line, font=font_obj, fill=fill)
        y += height + spacing


def draw_box(draw, box, text, fill, *, title=False):
    draw.rounded_rectangle(box, radius=12, fill=fill, outline=(60, 60, 60), width=2)
    draw_centered_text(draw, box, text, font(26 if title else 22, bold=title))


def draw_arrow(draw, start, end, fill=(70, 70, 70), width=4):
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 16
    for offset in (2.55, -2.55):
        point = (
            end[0] + length * math.cos(angle + offset),
            end[1] + length * math.sin(angle + offset),
        )
        draw.line([end, point], fill=fill, width=width)


def create_architecture_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1050), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), "NTU SPACE DYNAMICS - LAYERED SOFTWARE ARCHITECTURE", font=font(34, bold=True), fill=(0, 0, 0))
    layer_specs = [
        (110, 270, PALE_BLUE, "Application and orchestration layer"),
        (300, 460, PALE_GREEN, "Analysis-service layer"),
        (490, 650, PALE_GOLD, "Dynamics and numerical core"),
        (680, 840, PALE_ROSE, "Data and external scientific libraries"),
    ]
    labels = [
        ["Python API / notebooks", "Batch and end-to-end runs", "Mission-planning integration"],
        ["Orbit propagation", "Solar, power and thermal", "Attitude dynamics", "Ground-target access"],
        ["States and elements", "Frame transformations", "Force-model composition", "Integration / interpolation"],
        ["EGM2008 / ICGEM / GRACE", "Astropy / ERFA / IERS", "SciPy / NumPy", "python-sgp4 / local ephemeris"],
    ]
    for (y0, y1, fill, title), items in zip(layer_specs, labels):
        draw.rounded_rectangle((45, y0, 1555, y1), radius=18, fill=tuple(int(fill[i:i+2], 16) for i in (0, 2, 4)), outline=(80, 80, 80), width=2)
        draw.text((70, y0 + 14), title, font=font(24, bold=True), fill=(0, 0, 0))
        gap = 20
        available = 1440
        box_w = (available - gap * (len(items) - 1)) / len(items)
        for index, item in enumerate(items):
            x0 = 80 + index * (box_w + gap)
            draw_box(draw, (x0, y0 + 64, x0 + box_w, y1 - 18), item, (255, 255, 255))
    for y0, y1 in ((270, 300), (460, 490), (650, 680)):
        draw_arrow(draw, (800, y0), (800, y1))
    image.save(path, dpi=(200, 200))


def create_workflow_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 760), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 35), "END-TO-END COMPUTATION FLOW", font=font(34, bold=True), fill=(0, 0, 0))
    boxes = [
        ((55, 230, 285, 485), "Inputs\nUTC epoch\nOE / state / TLE\nmodel configuration", PALE_BLUE),
        ((330, 230, 560, 485), "Propagation\nTwo-body / J2\nSGP4 / HPOP", PALE_GREEN),
        ((605, 230, 835, 485), "Environment\nframes and gravity\nSun / eclipse / SRP", PALE_GOLD),
        ((880, 230, 1110, 485), "Vehicle analysis\nattitude dynamics\npower and temperature", PALE_ROSE),
        ((1155, 230, 1545, 485), "Mission outputs\nEphemeris\naccess windows\nplots and traceable metadata", VERY_LIGHT_GRAY),
    ]
    for box, text, color in boxes:
        rgb = tuple(int(color[i:i+2], 16) for i in (0, 2, 4))
        draw_box(draw, box, text, rgb, title=False)
    for left, right in zip(boxes[:-1], boxes[1:]):
        draw_arrow(draw, (left[0][2] + 8, 358), (right[0][0] - 8, 358))
    draw.text((60, 600), "All public interfaces use SI units, timezone-aware UTC, and explicit GCRS / ITRS / TEME frame labels.", font=font(24, bold=True), fill=(50, 50, 50))
    image.save(path, dpi=(200, 200))


def prepare_end_to_end_figure(source: Path, destination: Path) -> None:
    with Image.open(source) as image:
        if image.mode == "RGBA":
            background = Image.new("RGB", image.size, "white")
            background.paste(image, mask=image.getchannel("A"))
            image = background
        else:
            image = image.convert("RGB")
        image.save(destination, format="PNG", optimize=True, dpi=(240, 240))


def add_cover(document: Document) -> None:
    section = document.sections[0]
    configure_portrait(section, cover=True)
    clear_header_footer(section.header)
    clear_header_footer(section.footer)

    top = document.add_paragraph()
    top.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    top.paragraph_format.space_after = Pt(110)
    run = top.add_run("CLASSIFICATION   PUBLIC")
    set_run_font(run, size=11, bold=True)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("NTU SPACE DATA CENTER")
    set_run_font(run, size=14, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("NTU SPACE DYNAMICS")
    set_run_font(run, size=27, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Software Module Design Report")
    set_run_font(run, size=21, bold=True)

    scope = document.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_after = Pt(105)
    run = scope.add_run("Orbit • Attitude • Solar Environment • Power and Thermal • Access Analysis")
    set_run_font(run, size=11, italic=True, color=DARK_GRAY)

    meta = document.add_table(rows=6, cols=2)
    set_table_geometry(meta, [1.65, 3.65])
    meta.style = "Table Grid"
    meta_rows = [
        ("Document ID", "NTU-SD-SDD-001"),
        ("Version", "1.0"),
        ("Status", "Design Baseline"),
        ("Prepared by", "NTU Space Data Center"),
        ("Review state", "Prepared for Technical Review"),
        ("Issue date", "17 August 2026"),
    ]
    for row, (label, value) in zip(meta.rows, meta_rows):
        shade_cell(row.cells[0], LIGHT_GRAY)
        for index, text in enumerate((label, value)):
            paragraph = row.cells[index].paragraphs[0]
            paragraph.style = document.styles["Engineering Table Text"]
            run = paragraph.add_run(text)
            set_run_font(run, size=10, bold=(index == 0))
    set_table_geometry(meta, [1.65, 3.65])

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(65)
    spacer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = spacer.add_run("Ningbo University / NTU Space Data Center")
    set_run_font(run, size=12, bold=True)


def add_front_matter(document: Document) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_portrait(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    clear_header_footer(section.header)
    clear_header_footer(section.footer)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(14)
    run = title.add_run("Document Summary")
    set_run_font(run, size=17, bold=True)

    add_body(
        document,
        "This report defines the architecture, algorithms, public interfaces, data contracts, "
        "verification approach, and deployment constraints of the NTU Space Dynamics Python module. "
        "The module provides auditable computation services for orbit propagation, high-order Earth "
        "gravity, reference-frame conversion, attitude dynamics, solar geometry, eclipse and radiation "
        "pressure, solar-array thermal power, and ground-target access analysis.",
    )
    add_body(
        document,
        "The structure and engineering-document conventions are derived from the Overall Design Report "
        "for the Group-Cooperative Autonomous Mission-Planning Computing Service Platform. The content, "
        "interfaces, model fidelity statements, and validation evidence in this report describe the "
        "current ntu-space-dynamics implementation.",
    )
    add_table(
        document,
        "Table 1  Document metadata",
        ["Item", "Value"],
        [
            ["Module", "ntu-space-dynamics 0.1.0"],
            ["Language", "Python 3.10 or later"],
            ["License", "MIT software license; gravity coefficients retain their source data terms"],
            ["Primary audience", "Software developers, simulation engineers, mission analysts, technical reviewers"],
            ["Primary input contract", "SI units, timezone-aware UTC, explicit reference-frame labels"],
            ["Primary outputs", "Orbit and attitude ephemerides, thermal-power histories, access windows, diagnostic plots"],
        ],
        [1.55, 4.22],
        font_size=9.5,
    )
    add_table(
        document,
        "Table 2  Revision history",
        ["Version", "Date", "Description", "Status"],
        [["1.0", "17 Aug 2026", "Initial English software module design baseline", "Issued for review"]],
        [0.75, 1.15, 3.15, 0.72],
        font_size=9,
    )
    keywords = document.add_paragraph()
    keywords.paragraph_format.first_line_indent = Inches(0)
    keywords.paragraph_format.space_before = Pt(8)
    label = keywords.add_run("Keywords: ")
    set_run_font(label, size=10.5, bold=True)
    run = keywords.add_run(
        "orbital dynamics, HPOP, EGM2008, SGP4, attitude, Sun geometry, thermal power, access analysis"
    )
    set_run_font(run, size=10.5)

    document.add_page_break()
    toc_title = document.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = toc_title.add_run("TABLE OF CONTENTS")
    set_run_font(run, size=17, bold=True)
    toc = document.add_paragraph()
    toc.paragraph_format.first_line_indent = Inches(0)
    add_field(
        toc,
        'TOC \\o "1-3" \\h \\z \\u',
        "Table of contents updates when opened in Microsoft Word.",
    )
    appendix_toc = document.add_paragraph()
    appendix_toc.paragraph_format.first_line_indent = Inches(0)
    add_field(
        appendix_toc,
        'TOC \\h \\z \\t "Appendix Heading,1"',
        "Appendix entries update when opened in Microsoft Word.",
    )


def add_main_body(document: Document, bullet_num: int, decimal_num: int) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_portrait(section)
    add_running_header(section)
    add_running_footer(section, restart=1)

    add_heading(document, "Document Overview", 1)
    add_heading(document, "Purpose", 2)
    add_body(
        document,
        "The purpose of this report is to provide an implementable and reviewable design baseline for the "
        "current NTU Space Dynamics software module. It explains what the module computes, how the major "
        "algorithms are composed, which frames and units are accepted, what each public interface returns, "
        "and where the present engineering models deliberately stop.",
    )
    add_heading(document, "Scope", 2)
    add_bullets(
        document,
        [
            "Orbit-state and classical-element conversion, analytical and numerical orbit propagation, and TLE-based SGP4 propagation.",
            "Static high-order Earth gravity from bundled EGM2008 or external fully normalized ICGEM/GRACE coefficient files.",
            "GCRS, ITRS, and TEME state transformation, plus WGS-84 geodetic/ECEF conversion.",
            "Euler-angle conversion, quaternion rigid-body dynamics, Sun position and beta-angle analysis.",
            "Conical eclipse, solar irradiance, cannonball solar-radiation pressure, and one-node solar-panel thermal power.",
            "Ground-target access from direct orbit propagation or Hermite interpolation of sampled states.",
            "End-to-end computation and diagnostic plotting of position, velocity, beta angle, illumination, power, temperature, quaternion, and body-rate histories.",
        ],
        bullet_num,
    )
    add_heading(document, "Intended readers", 2)
    add_body(
        document,
        "The report is written for spacecraft simulation developers, mission-analysis engineers, software "
        "integrators, verification personnel, and reviewers who need enough orbital-dynamics context to "
        "evaluate the design without requiring a full astrodynamics textbook treatment.",
    )
    add_heading(document, "Normative terminology", 2)
    add_table(
        document,
        "Table 3  Terminology and frame conventions",
        ["Term", "Meaning in this module"],
        [
            ["GCRS", "Geocentric Celestial Reference System; the numerical-propagation inertial-frame contract."],
            ["ITRS / ECEF", "Earth-fixed frame used for gravity synthesis and ground geometry."],
            ["TEME", "True Equator, Mean Equinox frame returned by the SGP4 implementation."],
            ["Osculating elements", "Instantaneous Keplerian elements derived from a Cartesian state."],
            ["Mean elements", "Averaged orbital parameters represented by a TLE and interpreted by SGP4."],
            ["HPOP", "Composable high-precision Cowell propagation architecture; not a single undeclared fidelity level."],
            ["Access", "Time interval satisfying Earth-clear line of sight, minimum elevation, sensor FOV, and optional range."],
        ],
        [1.25, 4.52],
        font_size=9.2,
    )

    add_heading(document, "Requirements and Design Objectives", 1)
    add_body(
        document,
        "The software was designed around explicit model semantics and auditability. A calculation result "
        "must identify its time system, units, frame, force-model selection, gravity data, and numerical "
        "tolerances. The module therefore favors a small set of composable Python interfaces over an opaque "
        "monolithic simulation service.",
    )
    add_table(
        document,
        "Table 4  Principal design objectives",
        ["ID", "Objective", "Design response"],
        [
            ["DO-01", "Complete requested functional coverage", "Public APIs cover all requested orbit, attitude, solar, power, pressure, gravity, and access functions."],
            ["DO-02", "Safe data contracts", "SI units, aware UTC datetimes, and explicit GCRS/ITRS/TEME labels are enforced at public boundaries."],
            ["DO-03", "Fidelity is selectable", "Two-body, J2, SGP4, and HPOP represent distinct model semantics instead of a single hidden propagator."],
            ["DO-04", "Reproducibility", "EGM2008 provenance and SHA-256 are documented; IERS, ephemeris, truncation, and tolerance settings are task inputs."],
            ["DO-05", "Offline operation", "Astropy built-in ephemeris and locally available IERS tables support deterministic disconnected runs."],
            ["DO-06", "Extensibility", "HPOP accepts custom force models and gravity coefficients; torque can be constant or callable."],
            ["DO-07", "Verification", "Twenty-three unit tests cover invariants, reference vectors, transforms, gravity synthesis, solar geometry, thermal trends, and access events."],
        ],
        [0.65, 1.7, 3.42],
        font_size=8.8,
    )
    add_note(
        document,
        "Design rule",
        "A frame label is part of the numerical value. A TEME vector must never be renamed GCRS without a real transformation, and a velocity transform must include transport terms caused by Earth rotation.",
    )

    add_heading(document, "Overall Design", 1)
    add_heading(document, "Layered architecture", 2)
    add_body(
        document,
        "The module is organized as a layered scientific library. User applications and batch workflows call "
        "analysis services. Those services reuse common types, frame transformations, gravity synthesis, "
        "force models, integration, interpolation, and event-search utilities. External scientific libraries "
        "are isolated behind explicit frame and model contracts.",
    )
    add_picture(document, ASSET_DIR / "architecture.png", "Figure 1  Layered software architecture", 5.55)
    add_heading(document, "Package decomposition", 2)
    add_table(
        document,
        "Table 5  Package-level decomposition",
        ["Package or module", "Responsibility", "Principal public API"],
        [
            ["types.py", "Validated immutable data contracts", "OrbitState, Ephemeris, ClassicalElements, AttitudeEphemeris"],
            ["orbits/", "Element conversion and propagation", "oe_to_rv, rv_to_oe, TwoBodyPropagator, J2Propagator, SGP4Propagator, HighPrecisionPropagator"],
            ["gravity.py", "ICGEM loading, Earth orientation, spherical harmonics", "GravityFieldCoefficients, EarthOrientationCache, SphericalHarmonicGravity"],
            ["coordinates.py", "Celestial/terrestrial transformations", "eci_to_ecef, ecef_to_eci, transform_ephemeris, geodetic/ECEF conversion"],
            ["attitude.py", "Euler conversion and rigid-body propagation", "ypr_to_euler, euler_to_dcm, EulerDynamics"],
            ["solar.py", "Sun, eclipse, beta angle, irradiance, SRP", "sun_position_analytical, sun_position_ephemeris, eclipse_fraction, sun_intensity, sun_radiation_pressure"],
            ["power.py", "One-node panel thermal and electrical model", "SolarPanel, power_considering_thermal"],
            ["access.py", "Ground target/sensor geometry and event refinement", "GroundTarget, NadirSensor, access_propagation, access_interpolation"],
        ],
        [1.22, 2.25, 2.30],
        font_size=8.3,
    )
    add_heading(document, "Shared data flow", 2)
    add_picture(document, ASSET_DIR / "workflow.png", "Figure 2  End-to-end computation flow", 5.55)
    add_body(
        document,
        "Orbit outputs are reused rather than re-derived by each subsystem. The same epoch sequence drives "
        "Sun vectors, eclipse, panel irradiance, attitude propagation, and access analysis. This alignment "
        "prevents hidden interpolation and frame mismatches between otherwise independent calculations.",
    )

    add_heading(document, "Concise Orbital-Dynamics Background", 1)
    add_heading(document, "State propagation and model hierarchy", 2)
    add_body(
        document,
        "An orbit is represented by position r and velocity v at an aware UTC epoch. Numerical propagation "
        "integrates the first-order state equations. The central term produces ideal Keplerian motion; each "
        "additional acceleration represents a declared perturbation.",
    )
    add_equation(document, "dr/dt = v,        dv/dt = −μ r / |r|³ + Σ aₚ")
    add_body(
        document,
        "Two-body propagation keeps only the central term. J2 adds the dominant oblateness effect. HPOP "
        "integrates a configurable sum of high-order gravity, third-body gravity, drag, solar pressure, "
        "relativity, and custom forces. SGP4 is different: it propagates the mean elements encoded in a TLE "
        "using the SGP4 analytical theory and returns TEME states.",
    )
    add_heading(document, "Classical orbital elements", 2)
    add_body(
        document,
        "For elliptic motion the six classical elements are semi-major axis a, eccentricity e, inclination i, "
        "right ascension of the ascending node Ω, argument of periapsis ω, and true anomaly ν. They provide "
        "a compact geometric description, but Ω and ω become singular for equatorial and circular orbits. "
        "The implementation uses deterministic conventions so that OE→RV→OE→RV reconstructs the original state.",
    )
    add_heading(document, "Non-spherical gravity", 2)
    add_body(
        document,
        "Earth's mass distribution is represented by normalized spherical-harmonic coefficients C̄nm and S̄nm. "
        "J2 is the dominant degree-two zonal contribution. Higher degrees and orders represent progressively "
        "finer latitude/longitude structure. Gravity synthesis is performed in ITRS; the acceleration is then "
        "rotated back to GCRS for integration.",
    )
    add_equation(document, "U(r, φ, λ) = μ/r · Σₙ (R/r)ⁿ Σₘ P̄ₙₘ(sin φ)[C̄ₙₘ cos(mλ)+S̄ₙₘ sin(mλ)]")
    add_body(
        document,
        "The packaged EGM2008 subset is fully normalized and tide-free through degree/order 120. The HPOP "
        "default truncation is 70×70. Increasing degree improves spatial detail but increases computation cost "
        "approximately with the number of retained coefficients.",
    )
    add_heading(document, "Frames, Sun geometry, and events", 2)
    add_body(
        document,
        "GCRS is used for numerical orbit states, ITRS for Earth-fixed gravity and targets, and TEME for raw "
        "SGP4 results. Sun beta angle is the signed elevation of the Sun direction above the orbital plane.",
    )
    add_equation(document, "β = asin(ĥ · ŝ),        ĥ = (r × v)/|r × v|")
    add_body(
        document,
        "Eclipse fraction is the visible fraction of the apparent solar disc after Earth occultation. Access "
        "is an event problem: a window is open only while the target elevation, sensor off-boresight, optional "
        "range, and Earth-clear line-of-sight margins are simultaneously non-negative. Sampled sign changes are "
        "refined with a scalar root solver.",
    )

    add_heading(document, "Detailed Functional Design", 1)
    add_heading(document, "Common data models", 2)
    add_table(
        document,
        "Table 6  Core data objects",
        ["Object", "Stored data", "Invariant or behavior"],
        [
            ["OrbitState", "epoch, position_m[3], velocity_m_s[3], frame", "Aware UTC; SI; frame is GCRS, ITRS, or TEME."],
            ["Ephemeris", "times[N], positions_m[N,3], velocities_m_s[N,3], frame", "Aligned lengths; monotonic time; state(i) reconstruction."],
            ["ClassicalElements", "a, e, i, Ω, ω, ν, optional epoch", "Elliptic osculating elements; radians and metres."],
            ["AttitudeEphemeris", "times_s, quaternion xyzw, body rates", "Body-to-inertial active quaternion; rates in body frame."],
            ["ThermalPowerResult", "temperature, electrical power, incident power", "Arrays aligned to input elapsed times."],
            ["AccessWindow", "start, end", "Positive duration property."],
            ["AccessResult", "windows plus sampled margins and geometry", "Provides traceable event-search diagnostics."],
        ],
        [1.35, 2.28, 2.14],
        font_size=8.7,
    )

    add_heading(document, "Orbit propagation", 2)
    add_table(
        document,
        "Table 7  Propagator selection and model semantics",
        ["Propagator", "Model", "Frame", "Primary use"],
        [
            ["TwoBodyPropagator", "Exact elliptic Kepler solution from the initial osculating state", "GCRS", "Baseline, initialization, unit tests, short arcs"],
            ["J2Propagator", "Central gravity + J2, integrated by DOP853", "GCRS", "Dominant node/perigee drift and a light numerical baseline"],
            ["SGP4Propagator", "Vallado SGP4/SDP4 applied to TLE mean elements", "TEME", "Catalog/TLE propagation and operational screening"],
            ["HighPrecisionPropagator", "Cowell integration with composable force stack", "GCRS", "Engineering high-fidelity prediction and sensitivity studies"],
        ],
        [1.35, 2.45, 0.70, 1.27],
        font_size=8.4,
    )
    add_heading(document, "Two-body propagator", 3)
    add_body(
        document,
        "TwoBodyPropagator converts the initial Cartesian state to elliptic elements, advances the mean anomaly "
        "with the Keplerian mean motion, solves Kepler's equation, and converts the result back to Cartesian "
        "position and velocity. It is analytical and does not accumulate numerical integration error over a "
        "single ideal orbit, but it deliberately omits all perturbations.",
    )
    add_heading(document, "J2 propagator", 3)
    add_body(
        document,
        "J2Propagator composes central gravity and the standard J2 acceleration and integrates the six-state "
        "Cartesian equations with SciPy DOP853. Users configure relative tolerance, position and velocity "
        "absolute tolerances, and maximum integration step. It excludes atmosphere, third bodies, SRP, and "
        "higher harmonics.",
    )
    add_heading(document, "SGP4 propagator", 3)
    add_body(
        document,
        "SGP4Propagator accepts two TLE lines, validates the epoch and satellite record through python-sgp4, "
        "and returns SI-unit TEME states. TLE age and element quality normally dominate the prediction error. "
        "The interface therefore keeps TEME explicit and requires transform_ephemeris for GCRS or ITRS use.",
    )
    add_heading(document, "High-precision propagator", 3)
    add_body(
        document,
        "HighPrecisionPropagator uses adaptive DOP853 Cowell integration. The default force stack is central "
        "gravity plus EGM2008 70×70, Sun/Moon third-body gravity, and Schwarzschild correction. When area-to-mass "
        "is positive, the stack also enables segmented exponential-atmosphere drag and cannonball SRP. The Sun "
        "and Moon ephemerides are shared across forces using sampled interpolation. Additional objects implementing "
        "acceleration(elapsed_s, state, epoch) can be appended through extra_forces.",
    )
    add_table(
        document,
        "Table 8  HPOP configuration",
        ["Parameter", "Default", "Design meaning"],
        [
            ["gravity_degree / order", "70 / same as degree", "Static harmonic truncation; None selects the J2 baseline path."],
            ["area_to_mass_m2_kg", "0.0", "Enables drag and SRP when positive."],
            ["drag_coefficient", "2.2", "Cannonball aerodynamic coefficient."],
            ["reflectivity_coefficient", "1.3", "Cannonball SRP coefficient."],
            ["include_third_body", "True", "Sun and Moon point-mass differential gravity."],
            ["include_relativity", "True", "Schwarzschild correction used by the engineering stack."],
            ["relative_tolerance", "1e-11", "DOP853 relative error control."],
            ["position / velocity tolerance", "1 mm / 1 μm·s⁻¹", "Component-wise absolute error scaling."],
            ["max_step_s", "120 s", "Upper bound on adaptive integration step."],
        ],
        [1.75, 1.22, 2.80],
        font_size=8.5,
    )

    add_heading(document, "High-order Earth gravity", 2)
    add_body(
        document,
        "GravityFieldCoefficients stores fully normalized static cosine/sine coefficient matrices, model GM, "
        "reference radius, tide system, and provenance. from_icgem reads static gfc records and supports explicit "
        "degree/order truncation. load_builtin_egm2008 verifies and loads the packaged 0–120 subset.",
    )
    add_body(
        document,
        "SphericalHarmonicGravity computes the non-central potential and acceleration in the fixed frame. It "
        "intentionally excludes the degree-zero central term because HPOP adds central gravity separately using "
        "the model's GM. EarthOrientationCache precomputes GCRS→ITRS rotations and interpolates them with SLERP; "
        "the default node interval is 300 s.",
    )
    add_table(
        document,
        "Table 9  Bundled EGM2008 data contract",
        ["Attribute", "Value"],
        [
            ["Packaged coverage", "Degree and order 0–120"],
            ["HPOP default", "70×70"],
            ["Normalization", "Fully normalized"],
            ["Tide system", "Tide-free"],
            ["Model GM", "3.986004415 × 10¹⁴ m³/s²"],
            ["Reference radius", "6,378,136.3 m"],
            ["Subset SHA-256", "6ae505fcd7acfed84c5f69ff3e729be9cbb450162fd6bcf9116584c5b00570f2"],
            ["External static models", "ICGEM gfc records; GRACE monthly/static files after epoch reduction if required"],
        ],
        [1.55, 4.22],
        font_size=8.6,
    )
    add_note(
        document,
        "Model limitation",
        "The loader does not directly evaluate time-variable gfct, trnd, asin, or acos records. A GRACE monthly solution or other time-variable field must first be reduced to static coefficients for the task epoch.",
    )

    add_heading(document, "Element and Cartesian-state conversion", 2)
    add_body(
        document,
        "oe_to_rv converts elliptic osculating classical elements to an inertial Cartesian state. rv_to_oe "
        "performs the inverse conversion and applies fixed conventions in singular cases: argument of periapsis "
        "is zero for circular inclined orbits; RAAN is zero for equatorial eccentric orbits; both are zero for "
        "circular equatorial orbits. The returned anomaly then represents argument of latitude or true longitude "
        "as appropriate.",
    )

    add_heading(document, "Coordinate and frame transformations", 2)
    add_body(
        document,
        "eci_to_ecef converts a GCRS OrbitState to ITRS; ecef_to_eci performs the inverse. Both transport "
        "velocity, including the rotational contribution. transform_ephemeris extends the same contract to "
        "sampled GCRS, ITRS, or TEME ephemerides through Astropy/ERFA. WGS-84 geodetic_to_ecef and "
        "ecef_to_geodetic support target definition and ground geometry.",
    )
    add_note(
        document,
        "Operational data",
        "Astropy transformations depend on UT1−UTC, polar motion, and leap seconds. Production runs should archive the controlled IERS Bulletin A/B table used for the calculation.",
    )

    add_heading(document, "Attitude conversion and dynamics", 2)
    add_body(
        document,
        "ypr_to_euler interprets yaw-pitch-roll as aerospace intrinsic Z-Y-X rotation and converts it to a "
        "requested Euler sequence. Upper-case sequences are intrinsic and lower-case sequences are extrinsic, "
        "matching SciPy Rotation. euler_to_dcm returns the active direction-cosine matrix.",
    )
    add_body(
        document,
        "EulerDynamics integrates the body-to-inertial quaternion in SciPy xyzw order and the rigid-body Euler "
        "equation. Inertia may be supplied as principal moments or a matrix. Torque is expressed in the body frame "
        "and may be constant or a callable of time, quaternion, and body rate.",
    )
    add_equation(document, "q̇ = 1/2 · q ⊗ [ωx, ωy, ωz, 0],        Iω̇ = τ − ω × (Iω)")

    add_heading(document, "Solar environment and radiation pressure", 2)
    add_table(
        document,
        "Table 10  Solar functions",
        ["Function", "Method", "Output"],
        [
            ["sun_position_analytical", "Low-order geocentric analytical Sun model; native MOD, optional GCRS conversion", "Earth-to-Sun vector [m]"],
            ["sun_position_ephemeris", "Astropy built-in ephemeris or controlled local JPL kernel", "GCRS-aligned Earth-to-Sun vector [m]"],
            ["sun_beta_angle", "Signed Sun elevation above the instantaneous orbit plane", "rad or deg"],
            ["eclipse_fraction", "Apparent-disc overlap with conical umbra/penumbra geometry", "Visible solar-disc fraction 0–1"],
            ["sun_intensity", "Solar constant scaled by actual Sun distance and optional eclipse fraction", "Direct irradiance [W/m²]"],
            ["sun_radiation_pressure", "Cannonball SRP, directed away from the Sun, optionally eclipsed", "Acceleration [m/s²]"],
        ],
        [1.55, 2.82, 1.40],
        font_size=8.4,
    )
    add_body(
        document,
        "The analytical and ephemeris Sun functions are separate because they serve different traceability and "
        "fidelity needs. The analytical model is small and deterministic. The ephemeris path can use Astropy's "
        "offline built-in solution or a task-controlled JPL binary kernel when jplephem is installed.",
    )

    add_heading(document, "Power considering panel temperature", 2)
    add_body(
        document,
        "SolarPanel contains area, reference conversion efficiency, temperature coefficient, absorptivity, "
        "emissivity, radiating area, heat capacity, and deep-space sink temperature. power_considering_thermal "
        "integrates a one-node lumped thermal balance. Incidence cosine is clipped to [0,1]; irradiance is expected "
        "to include eclipse before entering the model.",
    )
    add_equation(document, "Pincident = A · I · max(0, cos θ)")
    add_equation(document, "Pelectric = ηref[1 + kT(T − Tref)] · Pincident")
    add_equation(document, "C dT/dt = αPincident − Pelectric − εσArad(T⁴ − Tspace⁴)")
    add_body(
        document,
        "The panel is isothermal. Absorbed solar energy raises the node temperature after exported electrical "
        "power is removed; radiation to deep space cools the panel. Structural conduction, Earth infrared, "
        "albedo, battery behavior, and MPPT dynamics are outside the current model.",
    )

    add_heading(document, "Access propagation and interpolation", 2)
    add_body(
        document,
        "GroundTarget defines a fixed WGS-84 latitude, longitude, altitude, and name. NadirSensor defines a fixed "
        "VVLH boresight through pitch and roll, a half-cone field of view, minimum target elevation, and optional "
        "maximum range. Access is the intersection of the target local-horizon margin, sensor off-boresight margin, "
        "range margin, and Earth-clear geometry.",
    )
    add_body(
        document,
        "access_propagation first samples any compatible orbit propagator and then searches for events. "
        "access_interpolation receives an existing GCRS Ephemeris and uses cubic Hermite interpolation with the "
        "stored velocity derivatives. Both scan at scan_step_s and refine each sign-changing boundary with Brent's "
        "method to boundary_tolerance_s.",
    )
    add_table(
        document,
        "Table 11  Access outputs",
        ["Output", "Description"],
        [
            ["windows", "Tuple of AccessWindow(start, end) objects with duration_s."],
            ["sample_times", "Times used by the event scan."],
            ["combined_margin_rad", "Minimum of all active geometric margins; positive means accessible."],
            ["elevation_rad", "Target elevation in the local topocentric frame."],
            ["off_boresight_rad", "Angle between sensor boresight and target line of sight."],
            ["slant_range_m", "Spacecraft-to-target distance."],
        ],
        [1.70, 4.07],
        font_size=9,
    )

    add_heading(document, "Public Interface Design", 1)
    add_heading(document, "General contracts", 2)
    add_bullets(
        document,
        [
            "Lengths are metres, velocities are metres per second, elapsed times are seconds, torque is N·m, temperature is kelvin, and power is watts.",
            "Angles are radians unless degrees=True or a from_degrees constructor is explicitly used.",
            "datetime values must be timezone-aware; UTC is the operational convention.",
            "Numerical propagation initial states must be GCRS; SGP4 outputs TEME; Earth-fixed calculations use ITRS.",
            "Functions validate dimensions, finite values, monotonic time arrays, physical parameter ranges, and compatible frames.",
        ],
        bullet_num,
    )
    add_heading(document, "Primary calling interfaces", 2)
    add_table(
        document,
        "Table 12  Interface families",
        ["Family", "Primary call", "Input", "Output"],
        [
            ["Orbit", "propagator.propagate(initial, times) / propagate_grid(...)", "OrbitState + epochs", "Ephemeris"],
            ["Elements", "oe_to_rv(elements); rv_to_oe(r,v)", "ClassicalElements or Cartesian state", "Cartesian state or elements"],
            ["Gravity", "from_icgem(...); acceleration_fixed(r)", "Coefficient file + ITRS position", "Coefficients / acceleration"],
            ["Frames", "eci_to_ecef(state); transform_ephemeris(eph, target)", "State or ephemeris", "Transformed state/ephemeris"],
            ["Attitude", "EulerDynamics(...).propagate(q0, ω0, times)", "Inertia, torque, initial attitude", "AttitudeEphemeris"],
            ["Solar", "Sun/beta/eclipse/intensity/SRP functions", "Epoch, orbit state, area/mass", "Vectors, angles, fractions, acceleration"],
            ["Power", "power_considering_thermal(times,I,cosθ,panel)", "Time histories + SolarPanel", "ThermalPowerResult"],
            ["Access", "access_propagation(...) / access_interpolation(...)", "Orbit, target, sensor, event settings", "AccessResult"],
        ],
        [1.0, 2.15, 1.50, 1.12],
        font_size=7.8,
    )
    add_heading(document, "Minimal end-to-end call", 2)
    add_code(
        document,
        """from datetime import datetime, timedelta, timezone
from ntu_space_dynamics import (
    ClassicalElements, HighPrecisionPropagator, OrbitState,
    SolarPanel, oe_to_rv, sun_position_ephemeris,
    eclipse_fraction, sun_intensity, power_considering_thermal,
)

epoch = datetime(2026, 1, 1, tzinfo=timezone.utc)
oe = ClassicalElements.from_degrees(
    7_000_000.0, 0.001, 97.5, 20.0, 0.0, 0.0, epoch
)
r0, v0 = oe_to_rv(oe)
initial = OrbitState(epoch, r0, v0, "GCRS")
eph = HighPrecisionPropagator(gravity_degree=70).propagate_grid(
    initial, epoch + timedelta(hours=3), step_s=60.0
)""",
    )
    add_body(
        document,
        "The full example in examples/end_to_end.py adds Sun geometry, illumination, thermal power, rigid-body "
        "attitude, interpolated access, summary statistics, and multi-format plots.",
    )

    add_heading(document, "End-to-End Operational Workflow", 1)
    add_numbered_steps(
        document,
        [
            "Create an aware UTC epoch and define the initial Cartesian state, classical elements, or TLE record.",
            "Select the propagation semantics: two-body, J2, SGP4, or a declared HPOP force stack.",
            "Generate one aligned ephemeris and transform frames only through the public coordinate interfaces.",
            "Evaluate Sun vectors, beta angle, eclipse fraction, irradiance, SRP, thermal power, and attitude at the same epochs.",
            "Evaluate access from the generated ephemeris, refining each open/close event boundary.",
            "Export numerical results, model metadata, and diagnostic plots for engineering review.",
        ],
        decimal_num,
    )
    add_picture(
        document,
        REPORT_END_TO_END_FIGURE,
        "Figure 3  End-to-end propagation, solar, thermal-power, and attitude histories",
        5.15,
    )
    add_body(
        document,
        "The diagnostic figure is an engineering sanity-check surface rather than a formal orbit-determination "
        "validation product. It makes state discontinuities, eclipse transitions, unexpected power saturation, "
        "thermal divergence, quaternion normalization problems, and body-rate instability visible in one view.",
    )

    add_heading(document, "Numerical Design and Performance", 1)
    add_table(
        document,
        "Table 13  Numerical methods and controls",
        ["Area", "Method", "Primary controls"],
        [
            ["Two-body", "Analytical elliptic Kepler solution", "Epoch grid and gravitational parameter"],
            ["J2/HPOP", "SciPy solve_ivp DOP853", "rtol, position/velocity atol, max_step_s"],
            ["Attitude/thermal", "Adaptive non-stiff integration", "Input sampling and tolerance configuration"],
            ["Earth orientation", "Astropy/ERFA rotation nodes + SLERP", "IERS data and sample_step_s"],
            ["Sun/Moon", "Shared sampled ephemeris interpolation", "Ephemeris source and sampling interval"],
            ["Access state", "Cubic Hermite interpolation", "Orbit sample spacing and velocity quality"],
            ["Access events", "Fixed scan + Brent root refinement", "scan_step_s and boundary_tolerance_s"],
            ["Gravity", "Normalized harmonic summation", "Coefficient source, degree, order, tide system"],
        ],
        [1.25, 2.45, 2.07],
        font_size=8.7,
    )
    add_body(
        document,
        "High-order gravity is the main controllable cost in typical HPOP runs. Raising the degree/order "
        "increases both coefficient count and recurrence work. The recommended practice is to demonstrate "
        "convergence for the mission altitude and forecast duration, then select the lowest truncation that meets "
        "the error budget. Access scan_step_s must be smaller than the shortest window that must be detected.",
    )

    add_heading(document, "Verification and Validation", 1)
    add_body(
        document,
        "The current baseline passes 23 unit tests under Python's unittest runner. Tests use analytical invariants, "
        "round trips, Vallado reference vectors, published gravity metadata, finite-difference checks, model "
        "degeneration, and event-boundary refinement. This is implementation verification; mission-specific "
        "validation still requires comparison with authoritative ephemerides, tracking data, or a trusted tool.",
    )
    add_table(
        document,
        "Table 14  Automated verification summary",
        ["Test group", "Count", "Representative evidence"],
        [
            ["Elements and propagators", "7", "OE/RV round trips including singular cases; two-body period closure; J2 energy; HPOP/J2 equivalence; Vallado SGP4 vector"],
            ["Gravity", "8", "EGM metadata and zonals; C̄20→J2 reduction; finite-difference potential gradient; high-degree effect; ICGEM truncation; HPOP arc"],
            ["Coordinates and solar", "4", "GCRS/ITRS position-velocity round trip; analytical/ephemeris Sun direction; eclipse; beta and SRP direction"],
            ["Attitude, power, access", "4", "YPR rotation equivalence; torque-free invariants; thermal trend; refined access boundaries"],
        ],
        [1.65, 0.62, 3.50],
        font_size=8.5,
    )
    add_note(
        document,
        "Executed baseline",
        "23 tests passed on 17 August 2026. The end-to-end example also generated PNG, PDF, SVG, and TIFF diagnostic figures.",
    )

    add_heading(document, "Deployment and Configuration", 1)
    add_heading(document, "Runtime", 2)
    add_bullets(
        document,
        [
            "Python 3.10 or later.",
            "Required: NumPy ≥1.24, SciPy ≥1.10, Astropy ≥5.3, Matplotlib ≥3.7, python-sgp4 ≥2.23.",
            "Optional: jplephem ≥2.18 for a controlled local JPL binary ephemeris kernel.",
            "Packaged data: EGM2008 tide-free static coefficients through degree/order 120 and SHA-256 sidecar.",
        ],
        bullet_num,
    )
    add_heading(document, "Installation and VS Code interpreter", 2)
    add_code(
        document,
        """cd E:\\MyProject\\NTU_Space_Data_Center
python -m venv .venv
.\\.venv\\Scripts\\python.exe -m pip install -e .
.\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v""",
    )
    add_body(
        document,
        "In Visual Studio Code select Python: Select Interpreter and choose "
        "E:\\MyProject\\NTU_Space_Data_Center\\.venv\\Scripts\\python.exe. Production deployments should "
        "archive the exact coefficient file, IERS table, local ephemeris kernel if used, package versions, and "
        "model configuration alongside each result.",
    )

    add_heading(document, "Limitations and Planned Evolution", 1)
    add_table(
        document,
        "Table 15  Present limitations and recommended extensions",
        ["Subsystem", "Current boundary", "Recommended extension"],
        [
            ["Gravity", "Static fully normalized field; no tides or in-run time-variable coefficients", "Solid/ocean/pole tides; task-epoch GRACE evaluation; EOP/data version manager"],
            ["Atmosphere", "Segmented exponential density; no space weather response", "NRLMSISE-00/JB2008 with archived F10.7, F10.7a, and Ap/Kp"],
            ["SRP", "Cannonball area/mass and single reflectivity coefficient", "Faceted geometry, self-shadowing, Earth albedo and IR"],
            ["Ephemerides", "Astropy built-in by default", "Controlled DE440/DE441 or SPICE kernel bundle"],
            ["Thermal/power", "One isothermal panel node; no conduction or electrical storage/control", "Multi-node thermal network, Earth IR/albedo, battery, EPS, MPPT"],
            ["Attitude", "Rigid body with external torque; no actuator/sensor models", "Reaction wheels, magnetorquers, disturbances, guidance and control loops"],
            ["Access", "Fixed target and fixed VVLH sensor; no terrain/refraction/cloud/maneuver time", "Terrain horizon, atmosphere, keep-out constraints, slew and payload modes"],
            ["Validation", "Unit tests and engineering sanity plots", "Cross-tool campaigns and mission-tracking residual analysis"],
        ],
        [1.12, 2.30, 2.35],
        font_size=8.0,
    )


def add_traceability_landscape(document: Document) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_landscape(section)
    add_running_header(section)
    add_running_footer(section)
    add_heading(document, "Requirements Traceability Matrix", 1)
    add_body(
        document,
        "The matrix maps every requested capability to its implemented interface, computational model, primary output, and verification status.",
        first_indent=False,
    )
    rows = [
        ["F-01", "Twobody", "TwoBodyPropagator", "Elliptic Kepler analytical propagation", "GCRS Ephemeris", "Implemented / tested"],
        ["F-02", "J2", "J2Propagator", "Central gravity + J2, DOP853", "GCRS Ephemeris", "Implemented / tested"],
        ["F-03", "SGP4", "SGP4Propagator", "Vallado SGP4/SDP4, TLE mean elements", "TEME Ephemeris", "Implemented / Vallado vector tested"],
        ["F-04", "HPOP", "HighPrecisionPropagator", "Cowell + declared perturbation stack", "GCRS Ephemeris", "Implemented / baseline tested"],
        ["F-05", "High-order gravity", "GravityFieldCoefficients; SphericalHarmonicGravity", "EGM2008 70×70 default, packaged to 120; external static ICGEM/GRACE", "Potential / acceleration / HPOP", "Implemented / 8 gravity tests"],
        ["F-06", "OE2RV", "oe_to_rv", "Elliptic osculating elements to Cartesian state", "r [m], v [m/s]", "Implemented / round-trip tested"],
        ["F-07", "RV2OE", "rv_to_oe", "Cartesian state to elements with singular conventions", "ClassicalElements", "Implemented / singular cases tested"],
        ["F-08", "ECI2ECEF", "eci_to_ecef; ecef_to_eci; transform_ephemeris", "GCRS↔ITRS and TEME transforms with velocity", "OrbitState / Ephemeris", "Implemented / round-trip tested"],
        ["F-09", "YPR2EULER", "ypr_to_euler; euler_to_dcm", "Intrinsic Z-Y-X YPR and arbitrary Euler sequence", "Euler angles / DCM", "Implemented / rotation tested"],
        ["F-10", "EULERDYNAMICS", "EulerDynamics", "Quaternion kinematics + rigid-body Euler equation", "AttitudeEphemeris", "Implemented / invariants tested"],
        ["F-11", "SUN POSITION Analytical", "sun_position_analytical", "Low-order MOD model with optional GCRS conversion", "Sun vector [m]", "Implemented / direction tested"],
        ["F-12", "SUN POSITION ephemeris", "sun_position_ephemeris", "Astropy built-in or local JPL kernel", "Sun vector [m]", "Implemented / direction tested"],
        ["F-13", "SUN BETA ANGLE", "sun_beta_angle", "Signed Sun elevation above orbit plane", "Angle", "Implemented / sign tested"],
        ["F-14", "SUN INTENSITY", "eclipse_fraction; sun_intensity", "Conical eclipse and Sun-distance correction", "0–1 fraction; W/m²", "Implemented / eclipse tested"],
        ["F-15", "Power considering thermal", "SolarPanel; power_considering_thermal", "One-node energy balance and temperature-dependent efficiency", "ThermalPowerResult", "Implemented / thermal trend tested"],
        ["F-16", "Sun radiation pressure", "sun_radiation_pressure", "Eclipsed cannonball SRP", "Acceleration [m/s²]", "Implemented / direction tested"],
        ["F-17", "ACCESS PROPAGATION", "access_propagation", "Propagate then scan/refine target/sensor geometry", "AccessResult", "Implemented / event tested"],
        ["F-18", "ACCESS INTERPOLATION", "access_interpolation", "Cubic Hermite state interpolation + Brent refinement", "AccessResult", "Implemented / boundary tested"],
        ["F-19", "End-to-end diagnostics", "examples/end_to_end.py", "Aligned orbit, solar, power, thermal, attitude, access and plots", "Statistics + PNG/PDF/SVG/TIFF", "Implemented / executed"],
    ]
    add_table(
        document,
        "Table 16  Functional requirements response",
        ["ID", "Requested capability", "Implemented interface", "Design model", "Primary output", "Status"],
        rows,
        [0.58, 1.55, 2.30, 2.45, 1.45, 1.36],
        font_size=7.1,
    )


def add_appendices(document: Document) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_portrait(section)
    add_running_header(section)
    add_running_footer(section)

    paragraph = document.add_paragraph("Appendix A — Public API Inventory", style="Appendix Heading")
    paragraph.paragraph_format.keep_with_next = True
    api_rows = [
        ["AccessResult", "Type", "Access windows and sampled geometry/margins."],
        ["AccessWindow", "Type", "Start/end UTC and derived duration."],
        ["AttitudeEphemeris", "Type", "Quaternion and body-rate time history."],
        ["ClassicalElements", "Type", "Elliptic osculating elements."],
        ["Ephemeris", "Type", "Sampled orbit positions and velocities."],
        ["EulerDynamics", "Class", "Rigid-body attitude propagation."],
        ["EarthOrientationCache", "Class", "Cached GCRS→ITRS rotations with SLERP."],
        ["GroundTarget", "Type", "Fixed WGS-84 ground point."],
        ["GravityFieldCoefficients", "Type", "Static fully normalized harmonic coefficients."],
        ["HighPrecisionPropagator", "Class", "Composable Cowell HPOP."],
        ["J2Propagator", "Class", "Central + J2 numerical propagation."],
        ["NadirSensor", "Type", "Fixed VVLH sensor and access constraints."],
        ["OrbitState", "Type", "Single-epoch SI Cartesian state."],
        ["SGP4Propagator", "Class", "TLE propagation to TEME."],
        ["SolarPanel", "Type", "Thermal/electrical panel parameters."],
        ["SphericalHarmonicGravity", "Class", "Static non-central gravity synthesis."],
        ["ThermalPowerResult", "Type", "Temperature and power histories."],
        ["TwoBodyPropagator", "Class", "Exact elliptic Kepler propagation."],
        ["access_interpolation", "Function", "Access from a sampled GCRS ephemeris."],
        ["access_propagation", "Function", "Propagate then compute access."],
        ["ecef_to_eci", "Function", "ITRS/ECEF to GCRS state."],
        ["ecef_to_geodetic", "Function", "WGS-84 ECEF to latitude/longitude/altitude."],
        ["eci_to_ecef", "Function", "GCRS to ITRS/ECEF state."],
        ["eclipse_fraction", "Function", "Visible solar-disc fraction."],
        ["euler_to_dcm", "Function", "Euler angles to active DCM."],
        ["geodetic_to_ecef", "Function", "WGS-84 latitude/longitude/altitude to ECEF."],
        ["load_builtin_egm2008", "Function", "Load packaged EGM2008 to degree 120."],
        ["oe_to_rv", "Function", "Elements to Cartesian state."],
        ["power_considering_thermal", "Function", "Panel temperature and electrical power."],
        ["rv_to_oe", "Function", "Cartesian state to classical elements."],
        ["sun_beta_angle", "Function", "Signed Sun/orbit-plane angle."],
        ["sun_intensity", "Function", "Direct irradiance with optional eclipse."],
        ["sun_position_analytical", "Function", "Analytical Sun vector."],
        ["sun_position_ephemeris", "Function", "Ephemeris Sun vector."],
        ["sun_radiation_pressure", "Function", "Cannonball SRP acceleration."],
        ["transform_ephemeris", "Function", "GCRS/ITRS/TEME sampled state transform."],
        ["ypr_to_euler", "Function", "Aerospace YPR to arbitrary Euler sequence."],
    ]
    add_table(
        document,
        "Table A-1  Exported names in ntu_space_dynamics.__all__",
        ["Public name", "Kind", "Contract summary"],
        api_rows,
        [1.72, 0.70, 3.35],
        font_size=8.0,
    )

    paragraph = document.add_paragraph("Appendix B — Model-Selection Guidance", style="Appendix Heading")
    paragraph.paragraph_format.keep_with_next = True
    add_table(
        document,
        "Table B-1  Recommended model selection",
        ["Question", "Recommended starting point"],
        [
            ["Need a clean algorithm baseline or short initialization?", "TwoBodyPropagator."],
            ["Need the dominant oblateness drift with a small force stack?", "J2Propagator."],
            ["Starting from a TLE or catalog product?", "SGP4Propagator; keep TEME explicit."],
            ["Need perturbation composition and high-order gravity?", "HighPrecisionPropagator with declared coefficients, tolerances, and environmental inputs."],
            ["Need fast repeated access on an existing orbit?", "access_interpolation with sufficiently dense samples."],
            ["Need access while selecting or comparing propagators?", "access_propagation."],
            ["Need deterministic offline Sun calculations?", "Astropy built-in ephemeris or the analytical model, depending on the error budget."],
            ["Need higher solar-system accuracy?", "Local controlled JPL kernel with jplephem."],
        ],
        [2.75, 3.02],
        font_size=9.0,
    )

    paragraph = document.add_paragraph("Appendix C — References", style="Appendix Heading")
    paragraph.paragraph_format.keep_with_next = True
    references = [
        "Pavlis, N. K. et al. (2012). The development and evaluation of the Earth Gravitational Model 2008 (EGM2008). Journal of Geophysical Research, 117, B04406.",
        "Vallado, D. A. et al. (2006). Revisiting Spacetrack Report #3. AIAA/AAS Astrodynamics Specialist Conference.",
        "IERS Conventions (2010), IERS Technical Note 36.",
        "Montenbruck, O. and Gill, E. (2000). Satellite Orbits: Models, Methods and Applications.",
        "Astropy, ERFA, SciPy, NumPy, Matplotlib, and python-sgp4 project documentation.",
        "NTU Space Dynamics source documentation: README.md, docs/model-fidelity.md, docs/open-source-evaluation.md, and data/gravity/README.md.",
    ]
    for ref in references:
        add_body(document, ref, first_indent=False)


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build() -> Path:
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    if not END_TO_END_FIGURE.exists():
        raise FileNotFoundError(END_TO_END_FIGURE)
    FINAL.parent.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    create_architecture_diagram(ASSET_DIR / "architecture.png")
    create_workflow_diagram(ASSET_DIR / "workflow.png")
    prepare_end_to_end_figure(END_TO_END_FIGURE, REPORT_END_TO_END_FIGURE)

    document = Document()
    clear_document_body(document)
    configure_styles(document)
    bullet_num, decimal_num = configure_heading_numbering(document)
    document.core_properties.title = "NTU Space Dynamics Software Module Design Report"
    document.core_properties.subject = "Software architecture and detailed functional design"
    document.core_properties.author = "NTU Space Data Center"
    document.core_properties.keywords = "orbit, HPOP, EGM2008, attitude, solar, power, thermal, access"
    document.core_properties.comments = (
        "English design report derived from the structure of the group-cooperative autonomous "
        "mission-planning platform overall design report."
    )
    document.core_properties.created = datetime(2026, 8, 17)
    document.core_properties.modified = datetime(2026, 8, 17)

    add_cover(document)
    add_front_matter(document)
    add_main_body(document, bullet_num, decimal_num)
    add_traceability_landscape(document)
    add_appendices(document)
    for table in document.tables:
        if table.rows and not table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"):
            repeat_table_header(table.rows[0])
    set_update_fields(document)
    document.save(FINAL)
    return FINAL


if __name__ == "__main__":
    print(build())
