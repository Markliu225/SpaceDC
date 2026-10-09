"""Delivery checks and hashes. Does not convert an experiment failure to a pass."""
from pathlib import Path
import hashlib,json,platform
import numpy as np
from pypdf import PdfReader
from docx import Document
HERE=Path(__file__).resolve().parent;POWER=HERE.parent
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
s=json.loads((HERE/'results/summary.json').read_text())
assert s['simulink_all_pass']
assert s['checks']['pass'] and s['cycles']['pass_']
assert [m['id'] for m in s['cases'] if not m['pass']]==['NOISE']
assert not set(s['contract']['training_discharge_pulses_1based'])&set(s['contract']['validation_discharge_pulses_1based'])
for m in s['cases']:
 a=np.genfromtxt(HERE/'results'/f"{m['id']}_python.csv",delimiter=',',names=True)
 b=np.genfromtxt(HERE/'results'/f"{m['id']}_simulink.csv",delimiter=',',names=True)
 delta=np.genfromtxt(HERE/'results'/f"{m['id']}_difference.csv",delimiter=',',names=True)
 for col in a.dtype.names:assert np.max(abs(delta[col]-(a[col]-b[col])))<1e-12
 assert np.max(abs(np.diff(a['time_s'])-.1))<1e-9
for lang in ('CN','EN'):
 assert sha(POWER/f'SDTwin_Power_Design_Report_{lang}.docx')==sha(HERE/'design_before'/f'SDTwin_Power_Design_Report_{lang}.docx')
files=[]
for name in ('SDTwin_Power_Design_Report_CN_v19','SDTwin_Power_Design_Report_EN_v19','SDTwin_Battery_EKF_RLS_Test_Report_CN'):
 docx=POWER/(name+'.docx');pdf=POWER/(name+'.pdf');reader=PdfReader(pdf)
 assert all((p.extract_text() or '').strip() for p in reader.pages)
 if 'Test' in name:
  doc=Document(docx);forms=[t for t in doc.tables if len(t.rows)==11 and t.cell(0,0).text=='用例编号']
  assert len(forms)==13
  assert sum(t.cell(9,3).text=='不通过' for t in forms)==1
 files.append(dict(name=name,docx_sha256=sha(docx),pdf_sha256=sha(pdf),pages=len(reader.pages)))
paths=[POWER/'battery.py',HERE/'contract.json',HERE/'run.py',HERE/'analyze.py',HERE/'test_model.py',HERE/'build_reports.py']
paths+=list((HERE/'matlab').glob('*.m'))+list((HERE/'matlab').glob('*.slx'))+list((HERE/'fixtures').glob('*'))+list((HERE/'data').glob('*'))
manifest=dict(date='2026-10-08',execution_complete=True,all_acceptance_criteria_pass=False,
 case_records=13,passed_case_records=12,failed_case_records=1,failed_case='NOISE',
 independent_simulink_cases=9,compared_samples=sum(m['rows']for m in s['cases']),unit_test_functions_passed=8,
 validated_scope='25 C BAK pulse voltage and numerical/analytic/synthetic checks; real SOC and parameter truth unavailable',
 integration='Standalone rebuilt battery core; legacy Thermal SPM adapter not migrated',
 artifacts=files,hashes={str(p.relative_to(POWER)):sha(p) for p in paths},
 rendering='Word COM PDF export after packaged renderer reported missing LibreOffice; Poppler PNG review',
 visual_inspection=dict(date='2026-10-08',all_pages_reviewed=True,
  inspected_render_directories={
   'SDTwin_Power_Design_Report_CN_v19':'qa/final_v2/SDTwin_Power_Design_Report_CN_v19',
   'SDTwin_Power_Design_Report_EN_v19':'qa/final_v2/SDTwin_Power_Design_Report_EN_v19',
   'SDTwin_Battery_EKF_RLS_Test_Report_CN':'qa/final_v3'},
  findings='Final renders checked for clipping, overlap, detached captions and orphan paragraphs; none remained after revisions'),
 source_originals_preserved=True)
(HERE/'results/delivery_manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding='utf8')
print(json.dumps(dict(artifacts=files,case_records=13,passed=12,failed=1,simulink_samples=manifest['compared_samples']),indent=2))
