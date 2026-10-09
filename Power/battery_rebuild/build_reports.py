"""Update the existing design reports and retain Orbit-derived test forms."""
from pathlib import Path
from copy import deepcopy
import json,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION_START
HERE=Path(__file__).resolve().parent;POWER=HERE.parent;FIG=HERE/'figures'
S=json.loads((HERE/'results/summary.json').read_text());P=json.loads((HERE/'results/parameters.json').read_text())

def replace(p,text):
 rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 for e in list(p._p):
  if e.tag!=qn('w:pPr'):p._p.remove(e)
 r=p.add_run(text)
 if rpr is not None:r._r.insert(0,rpr)
 return p
def cell(c,text):
 replace(c.paragraphs[0],text)
 for p in c.paragraphs[1:]:p._p.getparent().remove(p._p)
def formula(c,text):
 cell(c,'');p=c.paragraphs[0]
 for j,line in enumerate(text.split('\n')):
  if j:p=c.add_paragraph()
  math=OxmlElement('m:oMath');run=OxmlElement('m:r');tx=OxmlElement('m:t');tx.text=line;run.append(tx);math.append(run);p._p.append(math)
def picture(p,file,width=14):
 replace(p,'');p.add_run().add_picture(str(FIG/(file+'.png')),width=Cm(width))

CN={
121:'整体分配仍先求太阳功率与负载需求，再在电池允许电流区间求解 P11。每次试探只重新求 OCV、欧姆压降与当前极化下的端电压，不推进状态，不调用估计器。配电损耗按 P3 计算。',
122:'限充时由允许的电池电流决定实际太阳输出。正常联合结果应提供端口功率、SOC 导数、极化电压导数与诊断。当前电池专项实现尚未完成这一整星接口的迁移，不引用旧总线对象作为新模型执行证据。',
7:'太阳能板  一阶等效电路电池  EKF 与 RLS  功率分配',
15:'本报告设计单星计算节点的供电模型。太阳能板、电池组、控制器和配电器根据日照与负载需求计算电力分配。电池部分改为一阶 Thevenin 等效电路，状态为 SOC 与 RC 极化电压；RLS 在线辨识电阻和时间常数，EKF 利用端电压测量估计状态。物理模型与估计器分别保存状态，温度仍由 Thermal 唯一维护。',
17:'本次修订覆盖电池方程、参数辨识、状态估计、热源与验证方法。可执行电池核心位于 Power/battery.py。太阳与配电设计沿用既有定义。新电池验证独立运行，尚未替换 Thermal 联合仿真的旧 SPM 适配器，不能将旧联合测试结果视为新模型验收。',
36:'电池核心已实现为 Power/battery.py，包含 Parameters、State、advance、response、safe_current 和 JointObserver。下列目录保留整个供电模块的目标接口结构；其中场景装配及联合求解接口仍需显式适配，不能直接把旧锂占比字段当作新状态。',
45:'各输入对应同一物理时刻。电池物理状态为 SOC 与极化电压 u₁；估计状态及其协方差另存。物理模型由电流推进，估计器由电流和电压测量推进。只有已接受的测量样本才更新 EKF 和 RLS，求根试算不得重复更新估计器。温度仅由 Thermal 保存。',
50:'电池端口功率已包含内部压降，P1 中不再重复扣除电池热源。热域接收 R₀ 与 R₁ 的耗散功率。RC 支路还存储能量，因此不能直接把电流乘以开路电压与端电压之差全都当作发热。',
51:'一阶 ECM 的两个物理状态是 SOC z 和 RC 极化电压 u₁。SOC 表示剩余电量比例，u₁ 表示有记忆的极化压降。端电压由开路电压、即时欧姆压降和极化压降计算。电流阶跃可使端电压立即跳变，但 SOC 与 u₁ 不跳变。',
52:'温度作为外部输入，参数表应注明支持的温度范围。本轮 BAK 数据标定仅覆盖 25°C；核心对其他温度明确拒绝计算。已保存的其他温度原始数据尚未完成容量和初始 SOC 校准，不能把这次结果称为多温度热电验证。',
53:'太阳与配电器仍采用既定准静态模型。电池按参考 PDF 的脉冲试验与静置回稳方法建立一阶等效电路，并在此基础上增加用户要求的 RLS 与 EKF。PDF 本身未规定该联合估计算法。一阶结构固定，不通过增加 RC 阶数改变本次任务。',
64:'物理模型输入单体电流、温度及状态，输出端电压、电功率、热源和下一步状态。估计器另读电压测量，输出校正前预测电压、估计 SOC、极化电压和在线参数。测量电压不得反向写入物理仿真状态，真实 SOC 不得作为估计器输入。',
66:'一个开路电压源与欧姆电阻 R₀ 串联，再串联一个 R₁ 与 C₁ 并联的支路。该支路保留停流后逐渐回稳的极化状态。参数均针对单体；电池组按相同单体、均温与均流假设换算。',
69:'SOC 与极化状态',
70:'采用放电电流为正、充电为负的约定。容量 Q 的单位为 Ah，时间单位为 s，因此 SOC 变化率需要除以 3600 Q。',
71:'SOC z 无量纲。给定电流后必须沿时间逐步积分，不允许用少量独立工作点代替动态过程。达到支持的 SOC 范围边界时，停止相应方向的充放电；越界输入直接报错，不能先积分越界再裁剪电量。',
72:'RC 极化电压的变化率为',
73:'u₁ 的单位为 V，R₁ 为 Ω，C₁ 为 F，时间常数 τ＝R₁C₁，单位为 s。静置时 i＝0，u₁ 按指数衰减。该状态是旧模型缺少秒级回稳响应时所需的动态记忆。',
74:'固定步长内保持电流与参数不变时，SOC 按电量积分更新，u₁ 使用指数精确更新。试验采用 0.1 s 步长；完整容量曲线采用不大于 1 s 的步长，切换时继承所有物理状态。',
76:'单体端电压由开路电压扣除欧姆和极化压降：',
77:'OCV 为当前 SOC 下的准开路电压插值，R₀ 描述电流变化带来的立即压降，u₁ 描述随时间积累和释放的极化。充电电流为负时，欧姆项抬高电压；停止充电后端电压立即回落并继续松弛，这是电路响应，不是电量减少。',
78:'OCV 查表来自训练脉冲前静置电压。SOC 端点外的少量线性延伸仅用于数值约束检查，未获实测验证。实测测试始终位于标定点覆盖范围，报告不得将插值、外推和实测范围混写。',
80:'电阻耗散作为不可逆热源，RC 储能另行核算：',
81:'单体热源为 i²R₀＋u₁²/R₁，整组乘 NₛNₚ。常参数时，i·OCV−i·V 等于热源与 RC 储能变化率之和。本次没有可逆熵热、老化和副反应的标定数据，不声称热模型已经获得真实电芯的热量验证。',
82:'在线参数属于估计器，更新参数不得把物理电池状态重置。若未来把随时间变化的 C₁ 直接用于物理模型，储能导数还包含 ½u₁²·dC₁/dt。此次能量验算固定物理参数，不能将估计参数跳变造成的储能变化计为真实发热。',
84:'P9 的热容和热阻由 Thermal 的电池节点资产给定。温度只在热域积分一次。当前新电池试验为恒温 25°C，尚未执行新 ECM 与六节点 Thermal 的联合闭环，不沿用旧 SPM 的热耦合通过结论。',
85:'标定参数与在线估计',
86:'参考数据为 MathWorks 提供的 BAK N18650CL-29 脉冲试验。采用示例的 2.84 Ah 容量及原始记录初始 SOC 0.055，通过电流积分形成 SOC 参考。该 SOC 是带假设的电量参考，不是独立测得的真实 SOC。',
87:'一阶离散模型与 RLS 参数映射为',
88:'EKF 每 0.1 s 先预测 SOC 和 u₁，再保存预测端电压，随后用电压残差校正状态。状态转移雅可比为 diag 1,a，测量雅可比为 OCV 对 SOC 的斜率与 −1。协方差采用 Joseph 形式更新。初始 SOC 可故意加入 10 个百分点偏差检验恢复。',
89:'RLS 每 1 s 更新一次。对电流和 OCV 基线减实测电压施加相同的 2 s 一阶滤波，使用四项回归量：上次滤波压差、当前及上次滤波电流、常数 1。常数项吸收局部 OCV 偏差，不增加物理 RC 状态。基线只累积预测产生的 OCV 增量，排除 EKF 校正跳变。遗忘因子为 0.9995。',
90:'把 a、b₀、b₁ 换算为 R₀、R₁ 与 τ，仅接受正值及预设范围内的参数。无电流激励时冻结参数，非法候选仅记录而不应用于电池。当前数据得到 R₀＝%.8f Ω，R₁＝%.8f Ω，τ＝%.6f s，C₁＝%.3f F，作为在线估计的离线初值。OCV 与容量不由本版 RLS 辨识。'%(P['r0'],P['r1'],P['tau'],P['tau']/P['r1']),
93:'给定当前 SOC 与 u₁，端电压满足 V＝OCV−u₁−iR₀。将其代入 P11 求电流，选择与零电流连续相接的可行分支。约束同时包括电流、电压、有效 SOC 区间和下一步电量边界。估计器不参与同一步电流试探的重复更新。',
95:'确定电流后推进 SOC 与 u₁，并向 Thermal 提供热源。到达电压或 SOC 边界时先处理保护，再继续仿真。当前 safe_current 检查瞬时电压及下一步 SOC；它不替代未来电压边界的事件定位，整星分配适配仍需实现该部分。',
100:'物理状态与估计状态分别保存。前者只有 SOC 与 u₁，后者还包含 EKF 协方差、RLS 系数、参数协方差和滤波历史。所有实例分别初始化。已有 PowerState 总线是联合求解目标接口，不能把旧 SPM 运行对象直接传入新核心。',
102:'电池参数包括 capacity_Ah、r0、r1、tau、soc_grid、ocv_grid、temperature_C 以及串并联配置。C₁＝τ/R₁。估计器配置包括两种更新周期、遗忘因子、测量噪声、过程噪声、初始协方差、参数范围及激励判据。完整数值保存在 battery_rebuild/contract.json。',
103:'有效 SOC 区间由 OCV 参数表给定，与器件电压、电流安全边界分开保存。本轮实测支持 SOC 约 0.103 至 0.987；表格延伸至 0.05 与 1 仅用于边界检查，不能解释为该范围均已标定。',
104:'输入与参数非有限、温度超出支持范围、SOC 越界均明确拒绝。估计状态触及边界时记录 boundary_hits，不能把裁剪后的数值当作无误差估计。RLS 另记录接受、拒绝和冻结次数。',
107:'读取器件、场景、参数版本与来源。检查容量、电阻和时间常数为有限正数，OCV 表的 SOC 结点严格递增，初始状态处于支持范围，RLS 周期为 EKF 周期的整数倍。不从电极几何推算本模型的容量。',
108:'缺少 OCV、容量或 RC 初值时拒绝启动，不能静默套用其他电芯。固定参数保存在参数对象；运行中的物理状态与估计状态分别保存，不回写标定资产。',
110:'物理响应 response 为纯求值，给定状态与单体电流计算端电压、功率、热源与 RC 储能。advance 推进物理状态。JointObserver.step 仅在测量时刻调用，先产生校正前预测，再更新估计状态与下一步使用的参数。',
111:'新的单体实现位于 Power/battery.py。response 的 ns、np_ 参数完成同质电池组换算。恒温假设必须显式满足，不能把目标卫星温度直接套入 25°C 的验证参数。',
112:'输入电流单位 A，放电为正。State.soc 无量纲，State.polarization_V 的单位为 V。温度单位 °C，接入 Thermal 的 K 温度时必须显式换算。Parameters 的标定温度独立于热域状态。',
113:'先检查输入和 OCV 支持范围，再按 P4 至 P8 计算端电压、整组功率、热源与储能。该路径不读取电压测量，也不调用 EKF。试探电流不会改变电池状态或估计器历史。',
114:'输出包含 voltage_V、pack_voltage_V、pack_current_A、power_W、heat_W 与 polarization_energy_J。状态推进使用 advance 返回的新 State。旧 BatteryResponse 是目标联合接口，迁移时须显式映射并补齐事件边界处理。',
126:'保护事件改变接通与锁存状态，保留边界时刻的 SOC 和极化状态。估计器只能在接受新的测量后推进一次，重复事件或求根试算不得重复计入电量或更新协方差。',
138:'联合连续状态的电池分量改为 SOC 与 u₁。现有调用示例保留目标接口结构，尚需从旧 SPM 适配器迁移。已执行的新电池验证直接调用 response、advance 与 JointObserver，不使用下述拟议联合接口冒充已集成代码。',
139:'接受物理步后保存 SOC、u₁ 和时刻。EKF/RLS 状态另存，包含所有滤波历史以便恢复运行；不能仅保存最后一个 SOC 或最后一组电阻作为完整检查点。',
141:'固定参数、模型版本、输入数据哈希与初始状态。物理电池、EKF 和 RLS 分别初始化；EKF 初值可以有误差，不能用仿真真值在每步重置。',
144:'每次物理推进读取当前 SOC、u₁ 和温度，由电流更新状态。随后在规定测量时刻先生成预测，再吸收电压观测。记录预测和校正的先后关系，避免测量泄漏。',
145:'已接受的保护边界保留物理电量与极化状态。断开负载后仍计算电池静置回稳；估计器在无激励时冻结参数辨识，但仍可用电压观测估计状态。',
148:'常参数、零阶保持电流下使用 P10 的精确离散更新。EKF 步长 0.1 s，RLS 步长 1 s。SOC、极化电压、测量电压分别有独立单位与误差标准。输入切换按实际时间处理，保存频率不能替代物理步长。',
151:'单支路 ECM 每步只需 OCV 插值与指数更新。RLS 处理四维系数协方差，EKF 处理二维状态协方差。性能优化不得省略状态推进或将已校正电压标成预测电压。',
158:'现有 Thermal/sdtwin_sim/power_stand_in.py 仍为历史 SPM 测试实现。新电池核心已完成独立验证，不与旧适配器共用锂状态。上线联合场景前必须迁移状态向量、参数装配及事件处理，并重跑联合回归。',
160:'SimReady 只保存几何、参数来源和模型绑定，不自动生成 OCV 或 RC 参数。新模型标识应与历史 SPM 分开，避免同一资产因代码更新而静默改变物理含义。',
162:'Battery01 可绑定新一阶 ECM 与 Thermal 电池节点。不同实例分别保存 SOC、极化状态及估计器协方差。热域接口沿用功率与温度定义，不能共用可变估计历史。',
165:'安装尺寸与显示缩放仍按工程单位处理。电芯容量、OCV 和 RC 参数来自电测试，不能从三维显示比例推导。热容量和热阻继续由 Thermal 资产提供。',
167:'选定目标电芯及试验温度后，先完成容量与 OCV 标定，再设置一阶 RC 初值、SOC 和极化初值。EKF 的初始偏差与噪声协方差、RLS 的约束与激励条件均必须进入运行记录。',
169:'本轮训练采用 25°C 的第 1、2、5、7、9 个放电脉冲的加载段。第 4、6、8 个脉冲及其充电和静置段留作验证。第 3 个脉冲保留为来源文档指出的异常诊断，不纳入训练。其他温度数据保留但不宣称已验证。',
170:'串并联换算只针对单体量执行一次。辨识参数表注明单位、训练数据和误差；容量与 OCV 固定，RLS 输出 R₀、R₁、τ，C₁ 由二者换算。在线参数是估计值，不是独立测量的电芯真值。',
172:'初始物理字段改为 soc 与 polarization_V；估计器保存 x、P、theta、S 和滤波历史。实例、时间和保护字段仍由联合场景维护。旧 x_n 与 x_p 不作兼容别名，以免掩盖状态含义改变。',
174:'本版提供一阶 ECM、EKF 与 RLS 的可复现研究实现。25°C 留出工况的电压通过预设标准，但含 3 mV 噪声的参数辨识未全部通过。真实 SOC、温度变化、老化、单体不一致和完整整星热耦合尚未验证。',
175:'优先解决含噪声时 R₁ 与时间常数的辨识偏差，并取得目标电芯独立 SOC 与多温度标定数据。本次固定一阶结构；若未来改阶数，应另立模型版本与验证合同。',
182:'整个供电总线接口仍属联合集成契约。已执行的电池 API 以 Power/battery.py 及 battery_rebuild 的测试入口为准。当前报告不宣称新算法已替换所有旧场景。',
187:'[2] 用户提供的《用 MATLAB Simulink 做锂电池建模》，第 3 至 7 页。脉冲、静置回稳、参数估计及独立工况验证。',
188:'[3] MathWorks. Estimate Battery Model Parameters from HPPC Data. BAK 实测数据与容量、初始 SOC 定义。https://www.mathworks.com/help/simscape-battery/ug/estimate-battery-model-parameters-from-hppc-data.html',
189:'[4] MathWorks. Recursive Least Squares Estimator. https://www.mathworks.com/help/ident/ref/recursiveleastsquaresestimator.html',
190:'[5] MathWorks. SOC Estimator Kalman Filter. https://www.mathworks.com/help/simscape-battery/ref/socestimatorkalmanfilter.html',
}

# English paragraphs are authored explicitly for the new battery chapter. Other
# battery-dependent paragraphs receive concise, index-specific contract updates.
EN={
7:'Solar array  First order ECM  EKF and RLS  Power allocation',
15:'The Power model supplies a satellite compute node. The battery is rebuilt as a first-order Thevenin ECM with SOC and RC polarization voltage as physical states. RLS identifies resistance and time constant online; EKF estimates states from measured current and terminal voltage. Plant states and estimated states are separate. Thermal remains the sole temperature owner.',
17:'This revision replaces the battery equations, calibration, observer and verification contract. The executable core is Power/battery.py. Solar and distribution definitions are retained. The new battery test bench runs independently; the existing Thermal SPM adapter has not been migrated, so its earlier test verdicts do not validate the new model.',
64:'The plant consumes current, temperature and physical state, and produces voltage, power, heat and the next state. The observer additionally consumes measured voltage and produces prior voltage, estimated SOC, polarization and parameters. Measured voltage never resets the plant. True SOC is not an observer input.',
66:'The cell comprises an OCV source, series resistance R0 and one parallel R1-C1 branch. Polarization retains memory and relaxes after current interruption. Pack scaling assumes identical cell states and temperature with equal parallel currents.',
69:'SOC and Polarization State',70:'Discharge current is positive and charging current is negative. Capacity Q is in Ah and time is in seconds; the SOC derivative therefore divides current by 3600 Q.',
71:'SOC is dimensionless. Integrate every time interval rather than evaluating isolated operating points. Stop the relevant current direction at the supported SOC boundary. An out-of-domain physical step is rejected rather than clipping inventory afterward.',
72:'The RC polarization voltage evolves according to',
73:'Polarization u1 is in V, R1 in ohm and C1 in F. The time constant is tau = R1 C1 in seconds. At zero current, polarization decays exponentially. This state restores the relaxation memory absent from the previous approximation.',
74:'For held current and parameters, use exact exponential propagation of polarization and Coulomb propagation of SOC. Pulse tests use 0.1 s steps. Capacity curves use steps no larger than 1 s and preserve state at each reversal.',
76:'Cell terminal voltage is OCV less the instantaneous ohmic and dynamic polarization drops:',
77:'OCV is interpolated from SOC. R0 produces the immediate voltage step and u1 the delayed response. Negative charging current raises voltage. Removing charging current causes a voltage drop and subsequent relaxation without an instantaneous change in SOC.',
78:'The OCV curve uses pre-pulse rest voltages from training records. Small endpoint extensions are unvalidated and used only for boundary checks. Measured-data validation remains inside the measured SOC support.',
80:'Resistor dissipation and RC energy storage are accounted for separately:',
81:'Cell irreversible heat is i squared R0 plus u1 squared divided by R1. Pack heat scales by Ns Np. With constant parameters, OCV work less terminal work equals dissipated heat plus RC energy change. Reversible, aging and reaction heat have not been calibrated.',
82:'Online parameter updates belong to the observer and do not reset plant state. A physically varying C1 introduces an additional one-half u1 squared dC1/dt storage term. Energy verification here uses fixed physical parameters; estimated parameter jumps are not counted as real heat.',
84:'P9 retains Thermal ownership of temperature. This test campaign is isothermal at 25 C; the new ECM has not been tested in the complete six-node Thermal loop. Earlier SPM thermal verdicts do not transfer.',
85:'Calibration and Online Estimation',
86:'The measured source is the MathWorks BAK N18650CL-29 HPPC data set. Capacity is 2.84 Ah and initial raw-record SOC is 0.055, following the example. Integrated-current SOC is a reference with assumptions, not independently measured true SOC.',
87:'The first-order discrete model and RLS coefficient mapping are',
88:'EKF runs at 0.1 s. It predicts SOC and polarization, records prior terminal voltage, then corrects state using the measured voltage innovation. The transition Jacobian is diag(1,a); the measurement Jacobian is [dOCV/dSOC, -1]. Covariance uses the Joseph update.',
89:'RLS runs at 1 s with a common 2 s first-order prefilter on current and OCV baseline minus measured voltage. Regressors are previous filtered voltage drop, present and previous filtered current, and a constant. The constant absorbs local OCV bias. The baseline accumulates predicted OCV increments without EKF correction jumps. The forgetting factor is 0.9995.',
90:f"Only positive, bounded R0, R1 and tau candidates are applied. Freeze identification without current excitation; log rejected candidates. Offline initialization is R0={P['r0']:.8f} ohm, R1={P['r1']:.8f} ohm, tau={P['tau']:.6f} s, C1={P['tau']/P['r1']:.3f} F. OCV and capacity are not RLS outputs.",
}
EN_GENERIC={
121:'The target allocator first evaluates solar power and load demand, then solves P11 on the admissible current interval. Each trial evaluates OCV, ohmic loss and voltage at fixed polarization without advancing state or invoking the observer. P3 determines distribution loss.',
122:'Charge limits determine permitted battery current and actual solar output. A normal shared result includes powers, SOC and polarization derivatives and diagnostics. The new cell test bench has not migrated this full-system interface; the legacy bus is not evidence of new-model execution.',
36:'The executable battery core is Power/battery.py. The package tree below remains the target full-system architecture. State, asset and event adapters require explicit migration from the legacy SPM implementation.',
45:'All inputs share physical time. Plant SOC and polarization are separate from observer estimates and covariances. Update EKF/RLS once per accepted measurement, never during repeated power-allocation trials. Thermal alone owns temperature.',
50:'Terminal power already includes internal voltage loss. Thermal receives resistor dissipation. RC storage is separate, so current times OCV minus terminal voltage must not all be counted as heat.',
51:'Physical states are SOC and polarization voltage. Current steps change the ohmic voltage immediately while both states remain continuous. The estimator additionally stores covariance and filter history.',
52:'Temperature is external. Only 25 C is calibrated and other temperatures are explicitly rejected. Other raw temperature files are retained without claiming calibrated capacity, SOC or thermal validation.',
53:'Solar and distribution models remain quasi-static. The supplied PDF provides pulse/rest calibration and independent validation; this revision adds the requested EKF/RLS observer and fixes ECM order at one.',
93:'At fixed SOC and polarization, V = OCV - u1 - i R0. Solve P11 on the feasible branch connected to zero current, enforcing current, voltage and supported SOC limits. Allocation trials do not update the observer.',
95:'Advance SOC and polarization after selecting current. Handle boundaries before further propagation. The current helper constrains instantaneous voltage and next-step inventory; future voltage event localization still belongs to the full-system adapter.',
100:'Keep physical state separate from observer state. The latter includes EKF covariance, RLS coefficients and covariance, and filter history. The legacy PowerState is a target integration contract, not an accepted input to the new core.',
102:'Battery parameters are capacity_Ah, r0, r1, tau, soc_grid, ocv_grid and temperature_C. C1=tau/r1. Estimator rates, noise, covariance, excitation and bounds are frozen in battery_rebuild/contract.json.',
103:'Measured SOC support is approximately 0.103 to 0.987. Extensions to 0.05 and 1 are unvalidated boundary-test support, distinct from voltage/current safety limits.',
104:'Reject nonfinite inputs, unsupported temperature and physical SOC violations. Log estimated-state boundary hits and all accepted, rejected and frozen RLS updates.',
107:'Validate positive finite capacity, resistance and time constant, strictly ordered OCV knots, valid initial states and an integer RLS-to-EKF period ratio. Capacity is not derived from electrode geometry.',
108:'Missing OCV, capacity or RC parameters prevents startup. Do not silently use another cell. Runtime plant and observer states never overwrite calibration assets.',
110:'response is a pure plant evaluation; advance propagates physical state. JointObserver.step first records the prior voltage, then updates state and the parameters used at the next sample.',
111:'The implementation is Power/battery.py. response scales identical cells using ns and np_. The 25 C parameters cannot be applied to arbitrary satellite temperatures.',
112:'Current is in A and positive for discharge. State.soc is dimensionless and State.polarization_V is in V. Convert Thermal kelvin input explicitly to degrees Celsius.',
113:'Validate input and OCV support before evaluating voltage, pack power, heat and storage. This pure path does not read measured voltage or execute EKF.',
114:'Outputs are voltage_V, pack_voltage_V, pack_current_A, power_W, heat_W and polarization_energy_J. advance returns State. The target BatteryResponse bus requires an explicit adapter and boundary handling.',
126:'Protection events preserve SOC and polarization. Repeated trial evaluations never integrate charge or update estimator covariance. Each accepted measurement is processed once.',
138:'The battery components of a future shared state vector are SOC and u1. The pseudocode retains target interfaces. Executed tests call response, advance and JointObserver directly, not an unimplemented full-system adapter.',
139:'Persist accepted SOC, polarization and time separately from estimator state. A restart requires covariance, coefficients and filter histories as well as estimated SOC.',
141:'Freeze model, parameters, input hashes and initial states. Initialize plant, EKF and RLS separately; never reset an observer with true SOC at every step.',
144:'Advance physical states with the prescribed current. At a measurement time, predict voltage first and only then consume the voltage observation. Record both stages to expose leakage.',
145:'Accepted protection boundaries preserve physical states. Disconnection still permits relaxation. Without excitation freeze parameter identification while allowing EKF state correction.',
148:'Use exact held-input propagation at 0.1 s; RLS updates at 1 s. Capacity curves use at most 1 s. Output density does not substitute for physical propagation.',
151:'Each plant step needs OCV interpolation and exponential propagation. RLS has four regression coefficients and EKF has two states. Do not label corrected voltage as a prior prediction.',
158:'Thermal/sdtwin_sim/power_stand_in.py remains the historical SPM adapter. The new battery core is independently tested. Shared-state, scene and event adapters must be migrated and retested before full-system deployment.',
160:'SimReady carries geometry, parameter provenance and model bindings. It does not generate OCV or RC parameters. The new ECM must have an identifier distinct from the historical SPM.',
162:'Bind each battery instance to the new ECM and its Thermal node. Allocate separate plant states and observer histories. Electrical/thermal ports keep their physical meaning.',
165:'Engineering geometry and display scaling do not determine capacity, OCV or RC coefficients. Thermal assets continue to supply heat capacity and thermal resistance.',
167:'Calibrate capacity and OCV for the selected cell and temperature, then initialize RC parameters and states. Record estimator initial bias, noise, covariance, bounds and excitation requirements.',
169:'Training uses the load portions of 25 C discharge pulses 1, 2, 5, 7 and 9. Pulses 4, 6 and 8 with subsequent rest/charge are held out. Pulse 3 is retained as an anomaly case and excluded from fitting. Other temperatures are not validated.',
170:'Apply pack scaling only once. RLS estimates R0, R1 and tau; C1 follows algebraically. Estimated parameters are not independently measured ground truth.',
172:'Physical fields are soc and polarization_V. Store x, P, theta, S and filter history for the observer. The old x_n/x_p fields are not compatibility aliases for different physical states.',
174:'This is a reproducible research implementation. Held-out voltage tests at 25 C pass their targets, but noisy parameter identification does not pass every criterion. Real SOC, aging, variable temperature, imbalance and full spacecraft coupling remain unvalidated.',
175:'Address the noisy R1/time-constant bias and acquire independent SOC and temperature calibration before extending claims. A change in RC order requires a new model and validation contract.',
182:'Full Power bus interfaces remain an integration contract. Executed battery APIs and tests are in Power/battery.py and battery_rebuild. This revision does not claim replacement of every legacy scene.',
187:'[2] User-supplied MathWorks Chinese battery modeling guide, pages 3 to 7. Pulse/rest calibration and independent operating-data validation.',
188:'[3] MathWorks. Estimate Battery Model Parameters from HPPC Data. https://www.mathworks.com/help/simscape-battery/ug/estimate-battery-model-parameters-from-hppc-data.html',
189:'[4] MathWorks. Recursive Least Squares Estimator. https://www.mathworks.com/help/ident/ref/recursiveleastsquaresestimator.html',
190:'[5] MathWorks. SOC Estimator Kalman Filter. https://www.mathworks.com/help/simscape-battery/ref/socestimatorkalmanfilter.html',
}

def design(lang):
 en=lang=='EN';src=HERE/'design_before'/f'SDTwin_Power_Design_Report_{lang}.docx';d=Document(src)
 for idx,txt in CN.items():
  if en:txt=EN.get(idx,EN_GENERIC.get(idx));idx=idx+1 if idx>=35 else idx
  if txt is not None:replace(d.paragraphs[idx],txt)
 # Replace both battery-dependent diagrams while retaining the electrical topology.
 picture(d.paragraphs[44 if en else 43],'FLOW');picture(d.paragraphs[143 if en else 142],'RUNTIME')
 equations={10:'dz/dt = −i / (3600 Q)',11:'du₁/dt = −u₁ / (R₁C₁) + i / C₁',12:'V = OCV(z) − iR₀ − u₁',13:'q = i²R₀ + u₁²/R₁ ;   Q_B = NₛNₚq\nE_RC = ½C₁u₁² ;   i·OCV − iV = q + dE_RC/dt',15:'a = exp(−Δt/τ) ;   zₖ₊₁ = zₖ − iₖΔt/(3600Q)\nu₁,ₖ₊₁ = a u₁,ₖ + R₁(1−a)iₖ\nR₀ = b₀ ;   R₁ = (b₁+ab₀)/(1−a) ;   τ = −Δt/ln(a)'}
 for k,t in equations.items():formula(d.tables[k].cell(0,0),t)
 formula(d.tables[16].cell(0,0),'NₛNₚ i [OCV(z) − u₁ − iR₀] = P_need − P_pv_max')
 # Battery rows of shared contracts are rewritten rather than relabeling old math.
 L=lambda a,b:b if en else a
 cell(d.tables[0].cell(2,1),'');cell(d.tables[0].cell(3,1),L('2026年10月8日','8 October 2026'))
 cell(d.tables[1].cell(0,1),EN[15] if en else CN[15]);cell(d.tables[1].cell(1,1),'ECM; EKF; RLS; Simulink')
 for j,t in enumerate(['19.0','Battery reconstruction' if en else '电池模型重构','2026-10-08','',L('一阶 ECM 与 EKF RLS，增加实测留出验证及失败记录','First-order ECM, EKF/RLS, held-out measured tests and retained failures')]):cell(d.tables[1].cell(5,j),t)
 cell(d.tables[2].cell(4,1),L('物理状态为 SOC 与 RC 极化电压；估计状态和热域温度另存','Physical SOC and RC polarization; observer state and Thermal temperature are separate'))
 cell(d.tables[5].cell(2,1),'P4–P10, 1RC ECM');cell(d.tables[5].cell(2,2),'response / advance / JointObserver.step')
 replacements={
 17:{4:['PowerState','soc; polarization_V; load_connected; trip_latched',L('SOC 无量纲，极化电压 V，保护状态离散','Dimensionless SOC, polarization in V, discrete protection')],6:['BatteryResponse','voltage_V; power_W; heat_W; polarization_energy_J',L('实际单体 API 见 Power/battery.py','Executable cell API: Power/battery.py')],7:['PowerResult','P_pv_W; P_load_W; Q_B_W; Q_D_W; dsoc_dt; du1_dt',L('目标联合接口，尚待新模型适配','Target integration bus; new-model adapter pending')]},
 27:{1:[L('电池物理状态','Battery states'),L('零阶保持精确更新','Exact held-input update'),'dt = 0.1 s; cycle dt ≤ 1 s'],2:[L('状态与参数估计','State and parameter estimation'),'EKF / RLS','dt_EKF = 0.1 s; dt_RLS = 1 s']},
 28:{5:[L('初值与连接','Initial state and links'),'initial_state; observer_state',L('SOC、极化电压、协方差与滤波历史','SOC, polarization, covariance and filter history')]},
 29:{2:['Battery01',L('容量、OCV、RC 参数、初始 SOC 与极化','Capacity, OCV, RC parameters, initial SOC and polarization'),L('电端口接控制器，热端口接 Thermal','Electrical port to controller; heat port to Thermal')]},
 30:{2:[L('电芯参数','Cell parameters'),'Q_Ah, N_s, N_p',L('容量试验及拓扑配置','Capacity characterization and topology')],3:[L('电池动态','Battery dynamics'),'R0, R1, tau, C1',L('脉冲与回稳数据，一阶正参数拟合','Pulse/rest data, positive 1RC fitting')],4:['OCV','soc_grid; ocv_grid',L('静置电压与电量参考，注明外推范围','Rest voltage and Coulomb reference; disclose extrapolation')],5:[L('初值与工作边界','Initial values and limits'),'SOC; u1; current; voltage',L('有效域与器件安全边界分别保存','Separate model domain and device safety limits')]},
 33:{1:[L('轨道电量与温升','Orbital energy and temperature'),L('一阶 ECM；恒温试验已执行，完整热耦合待迁移','First-order ECM; isothermal tests executed, full coupling pending')],2:[L('更换电芯','Cell replacement'),L('重新标定容量、OCV 与 RC 参数','Recalibrate capacity, OCV and RC parameters')],3:[L('秒级脉冲与回稳','Second-scale pulses and relaxation'),L('一个 RC 动态状态，显式推进','One dynamic RC state with explicit propagation')],4:[L('在线状态与参数','Online state and parameters'),L('EKF 与 RLS；保留噪声辨识失败','EKF/RLS with retained noisy-identification failure')]}}
 for ti,rs in replacements.items():
  for row,values in rs.items():
   for j,value in enumerate(values):cell(d.tables[ti].cell(row,j),value)
 # Remove remaining old state symbols from target-interface pseudocode and rows.
 swaps=[('dx_n_dt','dsoc_dt'),('dx_p_dt','du1_dt'),('x_n','soc'),('x_p','polarization_V'),('两个锂状态','SOC 与极化状态'),('锂状态','SOC 与极化状态'),('锂占比约束','SOC 约束'),('初始锂占比','初始 SOC 与极化'),('两极初值','SOC 与极化初值'),('材料参数','电芯参数'),('current lithium states','current SOC and polarization'),('Two lithium states','SOC and polarization states'),('two lithium states','SOC and polarization states'),('lithium states','SOC and polarization states'),('lithium state','SOC and polarization'),('lithium-state','state'),('lithium-fraction constraint','SOC constraint')]
 for table in d.tables:
  done=set()
  for row in table.rows:
   for c in row.cells:
    if c._tc in done:continue
    done.add(c._tc)
    for p in c.paragraphs:
     if p._p.xpath('.//m:oMath'):continue
     txt=p.text
     for a,b in swaps:txt=txt.replace(a,b)
     if txt!=p.text:replace(p,txt)
 # Explicit algorithm appendix makes update order and equations reviewable.
 ap=d.add_paragraph(L('附录 D 电池估计器执行契约','Appendix D Battery Estimator Execution Contract'),'Appendix Heading');ap.paragraph_format.page_break_before=True
 for text in [
  L('RLS 回归为 yₖ＝a yₖ₋₁＋b₀iₖ＋b₁iₖ₋₁＋d。y 与 i 均先经过相同滤波，d 是局部偏差项。增益的分子为 Sφ，分母为 λ＋φᵀSφ。系数更新 θ←θ＋g·误差，协方差更新 S←S−gφᵀS 后除以 λ。','RLS uses y[k]=a y[k-1]+b0 i[k]+b1 i[k-1]+d after a common prefilter. d is a nuisance offset. Gain g=S phi/(lambda+phi transpose S phi); theta=theta+g residual; S=(S-g phi transpose S)/lambda.'),
  L('EKF 先使用上一步已接受参数预测状态与协方差，按 OCV 斜率线性化测量方程，再用 Joseph 公式更新协方差。输出预测电压时尚未读取本次校正量；RLS 接受的新参数从下一样本使用。','EKF predicts using previously accepted parameters, linearizes voltage using OCV slope and applies the Joseph covariance update. Prior voltage is recorded before correction. Newly accepted RLS parameters apply to the next sample.'),
  L('过去 20 s 的电流标准差超过 0.15 A 且至少有 10 个采样时才尝试辨识。R₀ 范围 0.001 至 0.2 Ω，R₁ 范围 0.001 至 0.3 Ω，τ 范围 1 至 500 s。非法候选不应用于 ECM，但保留无约束回归系数与协方差以继续学习。','Identification requires at least ten samples and current standard deviation above 0.15 A over the previous 20 s. Bounds are 0.001–0.2 ohm for R0, 0.001–0.3 ohm for R1 and 1–500 s for tau. Invalid candidates are not applied, but unconstrained regression state continues learning.'),
  L('所有数据、参数、源文件哈希、Python 与 Simulink 结果及逐点差值保存在 battery_rebuild。实测数据许可随原始与导出数据一同保存。完整用例与未通过项目见电池专项测试报告。','battery_rebuild retains data, parameters, source hashes, Python/Simulink outputs and pointwise differences. The measured-data license is retained with exports. The battery test report records each case and failures.')]:d.add_paragraph(text)
 d.save(POWER/f'SDTwin_Power_Design_Report_{lang}_v19.docx')

def para(d,text,style=None):return d.add_paragraph(text,style)
def table(d,headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):cell(c,v)
 for values in rows:
  for c,v in zip(t.add_row().cells,values):cell(c,str(v))
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3)
    for r in p.runs:r.font.size=Pt(10)
 return t
def image(d,name,caption):
 d.add_picture(str(FIG/(name+'.png')),width=Cm(14.3));d.paragraphs[-1].paragraph_format.keep_with_next=True
 d.paragraphs[-1].paragraph_format.line_spacing=1.0;d.paragraphs[-1].paragraph_format.space_before=Pt(6)
 cp=para(d,caption,'Caption');cp.paragraph_format.keep_with_next=False;cp.paragraph_format.keep_together=True
def test_report():
 src=POWER/'SDTwin_Power_Test_Report_CN.docx';d=Document(src);body=d.element.body;kids=list(body)
 form=deepcopy(d.tables[6]._tbl);portrait=deepcopy(d.sections[3]._sectPr);landscape=deepcopy(d.sections[-1]._sectPr)
 for e in kids[16:]:body.remove(e)
 body.append(portrait)
 replace(Paragraph(kids[5],d),'SDTwin BATTERY MODULE');replace(Paragraph(kids[6],d),'电池模型与状态估计软件测试报告');replace(Paragraph(kids[7],d),'一阶 ECM  RLS 参数辨识  EKF 状态估计  Simulink 对照')
 cell(d.tables[0].cell(3,1),'2026年10月8日')
 abstract='一阶 ECM 与 EKF RLS 已实现并逐步执行。9 组 Python 与 Simulink 动态记录一致。留出脉冲电压通过预设门槛；3 mV 噪声下的 R₁ 与时间常数辨识失败，不能据此宣称联合算法全面通过。'
 cell(d.tables[1].cell(0,1),abstract);cell(d.tables[1].cell(1,1),'一阶等效电路；递推最小二乘；扩展卡尔曼滤波；实测验证')
 for ri in range(4,8):
  vals=['2.0','电池专项重构','2026-10-08','', '替换电池方法，保留未通过判据与历史数据'] if ri==4 else ['']*5
  for j,v in enumerate(vals):cell(d.tables[1].cell(ri,j),v)
 for st in d.styles:
  if st.type==1 and st.name.startswith(('Heading','Title')):st.font.color.rgb=RGBColor(0,0,0)
 d.add_heading('1 概述',1);d.add_heading('1.1 标识',2)
 para(d,'报告编号 SDTwin BATTERY 2.0。被测实现为 Power/battery.py，模型为一阶 Thevenin ECM，状态估计器为 EKF，在线参数辨识器为带遗忘因子的 RLS。报告沿用 Orbit 的七章结构和逐项用例记录表。')
 d.add_heading('1.2 软件与文档概述',2);para(d,abstract)
 para(d,'本次验证对象是电池核心与估计器。旧 Thermal 联合适配器及旧 Power 测试保留为历史证据。新核心尚未接入整星联合求解，因此本报告不覆盖日食、配电器或完整 Thermal 闭环。')
 d.add_heading('1.3 缩略语与约定',2)
 table(d,['术语','含义'],[['ECM','开路电压源、欧姆电阻和一个 RC 支路组成的等效电路'],['RLS','从电流与端电压的递推回归辨识 R₀、R₁、τ，C₁＝τ/R₁'],['EKF','利用电压观测修正 SOC 和极化状态，保留校正前预测'],['SOC 参考','实测工况采用电流积分参考；合成工况有已知真值，两者不混用'],['正负号','电流为正表示放电，负数表示充电；原始 BAK 电流符号相反']])
 d.add_heading('2 测试内容',1);d.add_heading('2.1 实验配置与数据划分',2)
 para(d,'参考 PDF 采用脉冲放电、静置回稳、参数估计及独立工况验证。本轮选用其方法，未声称完全复刻其示意实验的时长。BAK 实测数据来自 MathWorks 配套示例，电芯为 N18650CL-29，电流、电压与时间来自实验记录。标称容量 2.9 Ah；本次电量换算按官方示例使用 2.84 Ah。')
 table(d,['项目','实际配置'],[['环境','25°C 恒温，单体电芯。温度是试验设定，文件未提供逐点实测温度。'],['试验序列','约 30 s 放电、40 s 静置、10 s 充电、40 s 静置；25°C 电流约为 6.19 A 放电、−4.64 A 充电，以保存的实测序列为准。'],['训练','放电脉冲 1、2、5、7、9，仅 0 至 30 s 加载段拟合 RC；训练脉冲前静置电压构建 OCV。'],['留出验证','脉冲 4、6、8，完整 0 至 119.9 s，包含未参与拟合的静置与充电。'],['异常诊断','脉冲 3 的起始静置电压问题由官方示例指出，保留单独测试，不修补后加入训练。'],['其他温度','保存 0、10、35、45°C 原始数据，但容量和初始 SOC 未分别校准，不作为本轮通过项目。']])
 d.add_heading('2.2 对比测试框架',2)
 image(d,'FLOW','图 1 物理电池与估计器的职责。上排物理模型只接收电流与状态；下排估计器读取电压测量，输出用于评估的预测和估计。')
 para(d,'Python 和独立 MATLAB 方程在 Simulink 固定步长引擎中分别运行。两侧共享参数与输入文件，不读取另一侧输出。EKF 为 0.1 s，RLS 为 1 s。原始电流按零阶保持；电压仅在相同电流方向段内线性重采样，切换处不跨段插值。CSV 保存完整序列，图中残差直接由上图曲线逐点相减。')
 para(d,'三层证据分别是解析解与已知真值、留出实测电压、Python 与 Simulink 一致性。合成数据可判断 SOC 与参数误差；实测数据没有独立真实 SOC 和在线电阻测量，不能据电压贴合宣称这些量已被真实验证。')
 d.add_heading('3 详细测试项目',1)
 cases=[]
 def case(cid,title,purpose,inputs,expected,actual,passed,figs,anomaly='无额外异常。'):
  number=len(cases)+1;cases.append((cid,title,'通过' if passed else '不通过'))
  hp=d.add_heading(f'3.{number} 测试项目 {number} {title}',2)
  if number>1:hp.paragraph_format.page_break_before=True
  el=deepcopy(form);body.insert(len(body)-1,el);t=Table(el,d)
  vals={1:purpose,2:'25°C，单体，参数与判据见第 5 章。物理状态和估计状态分别初始化。',3:inputs,4:'载入冻结输入，逐步推进状态，保存所有输出；逐点对比参考，并按预设门槛判定。',5:expected,6:'使用第 5 章对应门槛。未执行内容不得作为通过项。',8:actual,10:anomaly}
  cell(t.cell(0,1),cid);cell(t.cell(0,3),title)
  for row,text in vals.items():cell(t.cell(row,1),text)
  cell(t.cell(7,1),'自动执行');cell(t.cell(7,3),'2026年10月8日');cell(t.cell(9,1),'自动结果归档');cell(t.cell(9,3),'通过' if passed else '不通过')
  for ri,row in enumerate(t.rows):
   for c in row.cells:
    for pp in c.paragraphs:pp.paragraph_format.keep_with_next=ri<len(t.rows)-1;pp.paragraph_format.keep_together=True
  for f,caption in figs:image(d,f,caption)
 case('ECM-01','脉冲与停流回稳','验证 RC 支路的动态记忆及电流阶跃后的电压变化。','SOC 0.65；0 至 10 s 放电 1 A，随后静置至 100 s；步长 0.1 s。','极化符合解析指数响应，停流后仍逐步恢复。',f"RC 电压与解析解最大差 {S['checks']['analytic_rc_max_error_V']:.3g} V。",True,[('ANALYTIC','图 2 上图为端电压，下图为 RC 极化状态及解析解。10 s 停流时欧姆压降立即消失，RC 压降随后逐渐衰减。')])
 case('ECM-02','发热储能与边界','检查电压损耗中的热耗散和暂存能量，检查满电及空电方向保护。','沿用 1 A 脉冲；常参数精确能量积分；另在支持 SOC 上下端点请求继续充放电，并检查 10 串 2 并换算。','OCV 功减端口功等于热量与 RC 储能变化之和；端点禁止越界电流。',f"能量残差 {S['checks']['energy_max_residual_J']:.3g} J；上下端点向外电流均为 0 A；非法温度与 SOC 越界被拒绝；串并联换算误差为 0。",True,[('ENERGY','图 3 OCV 功与端口功之差分成电阻发热和 RC 储能。停流后剩余 RC 能量继续耗散，因此仅比较端口损失与热量会漏掉储能。')])
 titles={'EXACT':'无噪声参数恢复','NOISE':'测量噪声与初始 SOC 偏差','FIXED_EKF':'固定参数 EKF 对照','DRIFT':'参数变化跟踪','REST':'静置时冻结参数辨识'}
 for m in S['cases']:
  cid=m['id'];real=cid.startswith('MEASURED');title=f"实测脉冲 {m['pulse']} 与充放电回稳" if real else titles[cid]
  if real:
   inputs=f"BAK 25°C 第 {m['pulse']} 个脉冲，初始电量参考 {json.loads((HERE/'fixtures'/f'{cid}.json').read_text())['initial_soc']:.6f}；0 至 119.9 s，1200 个样本。电流输入为该段实际记录，不用规则矩形替代。"
   expected='开环电压 RMSE 不超过 30 mV、最大误差不超过 100 mV；EKF 校正前预测 RMSE 不超过 30 mV。'
   actual=f"开环 RMSE {m['open_loop']['rmse_V']*1000:.3f} mV，最大误差 {m['open_loop']['max_V']*1000:.3f} mV；预测 RMSE {m['prior']['rmse_V']*1000:.3f} mV。最终估计 SOC 与电量参考差 {100*m['final_soc_error']:.3f} 个百分点。"
   anomaly='电压判据通过不等于真实 SOC 或参数已验证。SOC 缺少独立真值。' + ('本项为异常诊断，不能计入三组独立留出验证。' if m['anomaly_case'] else '')
  else:
   inputs='SOC 真值 0.65；每 10 s 切换一次电流，序列为 0、1.5、0、−1、2、0、−1.5、0、1、−0.5 A，100 s 周期重复至 600 s。'
   inputs+=('全程电流为 0 A，替代上述序列。' if cid=='REST' else '')
   inputs+=('电压加入标准差 3 mV 的高斯噪声，随机种子 20261008；EKF 初始 SOC 0.75。' if cid in ('NOISE','DRIFT','FIXED_EKF') else '不添加测量噪声；初始 SOC 正确。')
   inputs+=('300 s 时真实 R₀、R₁、τ 同时增大 30%，C₁ 不变。' if cid=='DRIFT' else '')
   inputs+=('初始 R₀ 与 τ 偏高 20%，R₁ 偏低 20%。' if cid!='REST' else '')
   expected='最终 SOC 误差不超过 2 个百分点，R₀、R₁、τ 相对误差不超过 20%，且至少接受一次参数更新。'
   if cid=='REST':expected='全程冻结参数，无接受或拒绝的候选更新。'
   if cid=='FIXED_EKF':expected='关闭 RLS，仅检验最终 SOC 误差不超过 2 个百分点。'
   actual=f"预测电压 RMSE {m['prior']['rmse_V']*1000:.3f} mV；最终 SOC 误差 {m['final_soc_error']*100:.3f} 个百分点；接受 {m['accepted']} 次、拒绝 {m['rejected']} 次参数更新。"
   if 'parameter_relative_error'in m:actual+='R₀、R₁、τ 最终相对误差分别为 '+ '、'.join(f'{100*m["parameter_relative_error"][x]:.2f}%'for x in ('r0','r1','tau'))+'。'
   anomaly='含噪回归存在偏差，R₁ 与 τ 超过门槛，不能以较小电压误差掩盖参数辨识失败。' if not m['pass'] else '本项使用合成真值检验算法，不能代替真实电芯参数验证。'
  figs=[(cid,f'图 {len(cases)+2} {title}。上图同时给出参考、开环与校正前预测；中图直接计算各预测减参考的误差；下图给出实际电流时序。')]
  if cid in ('EXACT','NOISE','DRIFT'):figs.append((cid+'_parameters','参数与状态误差。虚线为已知真值，实线为 RLS 估计。右下图显示估计 SOC 减真实 SOC，单位为百分点。'))
  case(cid,title,'区分动态预测、状态估计与参数辨识各自是否达标。',inputs,expected,actual,m['pass'],figs,anomaly)
 case('ECM-03','容量窗口内三次充放电','给出按电流积分得到的容量与电压曲线，检查循环状态继承。','0.1 C 即 0.284 A；SOC 在 0.12 至 0.96 间充放电三次；每步不超过 1 s。容量坐标每半循环重新累计，物理状态不重置。','各分支电量等于 SOC 变化乘容量；无老化模型时稳定循环应重合。',f"每个分支容量 {S['cycles']['capacity_per_branch_Ah']:.4f} Ah；总物理时间 {S['cycles']['duration_s']:.1f} s。",True,[('CYCLES','容量特性图。横轴为单体电流积分的 Ah，不是没有质量依据的 mAh/g。仅覆盖已标定 SOC 窗口；没有老化或循环衰退状态，因此不能伪造逐圈差异。')],'该图是标定模型的特性检查，没有整段实测容量曲线对照，不构成全容量电芯验证。')
 case('SL-01','独立 Simulink 逐步对照','检验两种实现对同一输入是否产生一致的状态、预测和参数。','9 个输入文件；共 '+str(sum(m['rows']for m in S['cases']))+' 个样本；Simulink FixedStepDiscrete，0.1 s；独立 MATLAB 实现，无 Python 回调。','所有同点连续输出最大差小于 10⁻⁶，离散标志相同。','9 组全部通过，逐点差值保存为 difference.csv。'+f"最大电压差 {max(m['max_errors']['prior_voltage_V']for m in S['simulink_comparison']):.3g} V。",S['simulink_all_pass'],[],'数值一致性只证明两侧实现相符，不是对实际电芯的独立物理证明。')
 table(d,['用例','最大预测电压差 V','最大 SOC 差'],[[m['id'],f"{m['max_errors']['prior_voltage_V']:.3g}",f"{m['max_errors']['estimated_soc']:.3g}"]for m in S['simulink_comparison']])
 d.add_heading('4 测试内容充分性分析',1)
 para(d,'物理核算覆盖脉冲回稳、常参数能量、SOC 边界和同质电池组换算。估计器覆盖无噪声、有噪声、初始 SOC 偏差、参数阶跃及无激励；固定参数 EKF 提供对照。另有 8 项程序回归验证独立 ODE 一致性、因果性、协方差半正定、非法输入、状态连续和参数冻结。')
 para(d,'实测独立验证仅覆盖 25°C 的三个留出片段。OCV 与电量参考仍依赖容量及初始 SOC 假设。真实 SOC、在线参数真值、其他温度、长时衰退、传感器偏置及目标航天电芯尚无独立验证。当前 RLS 含噪参数恢复失败，属于明确未通过项目。')
 d.add_heading('5 测试条件与要求',1);d.add_heading('5.1 参数与验收门槛',2)
 table(d,['参数','数值与含义'],[['电芯与温度','BAK N18650CL-29，25°C'],['Q','2.84 Ah，官方示例的电量换算容量'],['R₀ / R₁',f"{P['r0']:.8f} / {P['r1']:.8f} Ω"],['τ / C₁',f"{P['tau']:.6f} s / {P['tau']/P['r1']:.3f} F"],['采样与辨识','EKF 0.1 s，RLS 1 s，预滤波时间常数 2 s'],['噪声与协方差','电压标准差 0.003 V；过程方差率为 10⁻⁸ SOC²/s 与 10⁻⁶ V²/s；初始状态方差 0.01 与 0.0001。'],['RLS','遗忘因子 0.9995；初始协方差对角为 10⁶、10³、10³、10³。'],['实测电压','开环 RMSE ≤30 mV 且最大差 ≤100 mV；预测 RMSE ≤30 mV。'],['合成参数与 SOC','最终各参数误差 ≤20%，最终 SOC 误差 ≤2 个百分点；参数辨识须有实际接受更新。'],['跨实现','连续输出最大差 ≤10⁻⁶，离散标志逐点相同。']])
 image(d,'CALIBRATION','离线标定图。左图为训练静置电压形成的 OCV 曲线，灰色表示未验证的端点延伸。右图为五段加载数据分别拟合的电阻，在线参数初值取中位数。')
 para(d,'工程门槛在拟合前写入 contract.json，后续修正回归公式未放宽门槛。历史版本显示仅看电压误差会掩盖参数错误；修订后要求参数辨识确实发生更新，静置冻结单独验收。')
 table(d,['训练脉冲','SOC','加载拟合 RMSE mV','τ s'],[[f['pulse'],f"{f['soc']:.4f}",f"{f['fit']['rmse_V']*1000:.3f}",f"{f['tau']:.3f}"]for f in P['fits']])
 d.add_heading('5.2 执行环境',2)
 para(d,'Python 使用 Thermal/.venv 环境中的 NumPy 与 SciPy。独立对照使用 MATLAB R2026b 和 Simulink。保存的 slx 由自定义 MATLAB S 函数实现模型及估计器，不是 Simscape Battery 工厂电芯模块。S 函数使用离散状态保存所有历史，Outputs 不改变状态，Update 每个采样点提交一次状态。')
 d.add_heading('6 测试数据',1)
 table(d,['位置','内容'],[['data','五个温度的原始导出 CSV 及 BSD 许可'],['fixtures','各用例电流、电压、初值与参数；参考 SOC 仅供评分'],['results','完整 Python、Simulink、逐点差值和 summary.json'],['history','早期错误回归方案及其失败结果，不覆盖原始判据'],['matlab','可执行独立 MATLAB 公式、Simulink 模型和批处理入口'],['design_before','修订前双语设计报告']])
 para(d,'重现顺序为 run.py、matlab/run_simulink.m、analyze.py、test_model.py、build_reports.py。README 提供完整命令。原始时间没有压缩；容量曲线仍需按实际电量走完全过程。')
 para(d,'数据来源为 MathWorks 示例，作者 Anandaroop Bhattacharya 与 Subhasish Basu Majumder，IIT Kharagpur。原始许可保存在 data/LICENSE.txt。本文数据不是 SpaceDC 卫星电芯的实测，也不是用户提供图片的化学体系复现。')
 d.add_heading('7 测试总结',1)
 para(d,'重构已提供一阶 ECM、因果 EKF 与 RLS 的实际代码和独立 Simulink 对照。实测留出片段的开环电压达标，无噪声参数恢复达标。含 3 mV 噪声时，R₁ 误差 21.40%、τ 误差 30.46%，超过 20% 门槛，故联合辨识不能判为全面通过。下一步需要改善噪声下的辨识偏差并取得独立真实 SOC 数据；本版不应作为已经全面验证的航天电池模型。')
 table(d,['编号','测试项目','结论'],cases)
 para(d,'参考资料与输入 SHA256 见 results/summary.json。')
 for e in d.element.body.iter(qn('w:trHeight')):e.set(qn('w:hRule'),'atLeast')
 d.save(POWER/'SDTwin_Battery_EKF_RLS_Test_Report_CN.docx')

if __name__=='__main__':
 design('CN');design('EN');test_report()
 print('Updated CN/EN design reports and battery test report')
