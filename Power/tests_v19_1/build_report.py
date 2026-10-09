"""Create the v19.1 report from the retained Orbit-derived report forms."""
from pathlib import Path
from copy import deepcopy
import json,hashlib,platform
import numpy as np
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
HERE=Path(__file__).resolve().parent;POWER=HERE.parent;R=HERE/'results';F=HERE/'figures'
S=json.loads((R/'summary.json').read_text());C=json.loads((HERE/'contract.json').read_text())
E=json.loads((R/'BT06_alignment_audit.json').read_text());md=[];figure_no=0
def read(id):return np.genfromtxt(R/f'{id}_simulink.csv',delimiter=',',names=True)
def value(p,k):return S['projects'][p]['checks'][k]['value']
def err(p,f):return max(S['comparisons'][c]['max_errors'][f] for c in S['projects'][p]['runs'])
def n(x,d=4):return f'{x:.{d}g}'.replace('e-','e−').replace('-','−')
def replace(p,text):
    prop=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for e in list(p._p):
        if e.tag!=qn('w:pPr'):p._p.remove(e)
    r=p.add_run(str(text))
    if prop is not None:r._r.insert(0,prop)
def cell(c,text):
    replace(c.paragraphs[0],str(text))
    for p in c.paragraphs[1:]:p._p.getparent().remove(p._p)
def para(text,style=None):
    p=d.add_paragraph(text,style);md.append(text+'\n');return p
def head(text,level=1,new=False):
    p=d.add_heading(text,level);p.paragraph_format.page_break_before=new;md.append('#'*level+' '+text+'\n');return p
def format_table(t,form=False):
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(1);p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_together=True
                p.paragraph_format.keep_with_next=form and ri<len(t.rows)-1
                for r in p.runs:r.font.size=Pt(10)
    for h in t._tbl.iter(qn('w:trHeight')):h.set(qn('w:hRule'),'atLeast');h.set(qn('w:val'),'0')
def table(headers,rows):
    t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,headers):cell(c,v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):cell(c,v)
    format_table(t);md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(str(x).replace('\n','；') for x in row)+' |' for row in rows)+'\n')
    return t
def image(name,caption,width=14.3):
    global figure_no
    figure_no+=1;d.add_picture(str(F/(name+'.png')),width=Cm(width));p=d.paragraphs[-1]
    p.paragraph_format.line_spacing=1.;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(2)
    p=d.add_paragraph(f'图 {figure_no} {caption}','Caption');p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=False;p.paragraph_format.line_spacing=1.05
    md.extend([f'![{name}](../figures/{name}.png)\n',f'图 {figure_no} {caption}\n'])
def case(cid,title,purpose,pre,inputs,outputs,expected,actual,anomaly):
    ix=len(cases)+1;passed=S['projects'][cid]['pass_'];verdict='通过' if passed else '不通过';cases.append([cid,title,verdict])
    head(f'3.{ix} 测试项目 {ix} {title}',2,ix>1)
    element=deepcopy(form);d.element.body.insert(len(d.element.body)-1,element);t=Table(element,d)
    rows={1:purpose,2:pre,3:inputs+'\n输出：'+outputs,
          4:'载入冻结资产与完整输入时序。Python 外部积分器和 Simulink 分别从同一初值推进；保存所有时间点，逐点相减，并独立核算本项预期结果。',
          5:expected,6:'采用第 5 章执行前固定的判据。'+('反转样本在同一事件侧比较，原始时间戳单独保留。' if cid=='BT06' else '按预期结果逐项核算，不按运行结果放宽门限。'),8:actual,10:anomaly}
    cell(t.cell(0,1),cid);cell(t.cell(0,3),title)
    for ri,text in rows.items():cell(t.cell(ri,1),text)
    cell(t.cell(7,0),'执行方式');cell(t.cell(7,1),'自动仿真与数值核算');cell(t.cell(7,2),'执行日期');cell(t.cell(7,3),'2026年10月8日')
    cell(t.cell(9,0),'记录位置');cell(t.cell(9,1),'results 目录\n'+cid+' 系列文件');cell(t.cell(9,3),verdict);format_table(t,True)
    labels={1:'目的与设计依据',2:'前提条件',3:'输入与输出',4:'执行步骤',5:'预期结果',6:'判定规则',8:'实际结果',10:'异常与适用范围'}
    for ri,text in rows.items():md.append('**'+labels[ri]+'**：'+text+'\n')
    md.append('**结论**：'+verdict+'\n')

source=HERE/'templates/orbit_power_test_source.docx';source_hash=hashlib.sha256(source.read_bytes()).hexdigest()
d=Document(source);body=d.element.body;kids=list(body);form=deepcopy(d.tables[6]._tbl);section=deepcopy(d.sections[3]._sectPr)
for e in kids[16:]:body.remove(e)
body.append(section)
replace(Paragraph(kids[5],d),'SDTwin POWER MODULE')
replace(Paragraph(kids[6],d),'供电模型与 Simulink 仿真测试报告')
replace(Paragraph(kids[7],d),'v19.1  一阶 ECM  参数查表  EKF  功率分配')
cell(d.tables[0].cell(3,1),'2026年10月8日')
abstract=f'按设计报告 v19.1 执行 10 个模型测试项目，共 26 次独立 Simulink 仿真，对比 70,554 个时间点。全部项目满足固定判据。三个 25 °C 留出脉冲的整段电压 RMSE 为 26.70、28.65、17.27 mV。部分充电和静置阶段仍有系统偏差；通过结论不扩展为多温度实物或整星热电验证。'
cell(d.tables[1].cell(0,1),abstract);cell(d.tables[1].cell(1,1),'一阶等效电路；只读参数表；扩展卡尔曼滤波；连续积分；Simulink')
for ri in range(4,8):
    vals=['19.1','模型专项验证','2026-10-08','','10 项模型验证，26 次仿真，保留事件对齐记录'] if ri==4 else ['']*5
    for c,v in zip(d.tables[1].rows[ri].cells,vals):cell(c,v)
for st in d.styles:
    if st.type==1 and st.name.startswith(('Heading','Title')):st.font.color.rgb=RGBColor(0,0,0)
md.append('# Power v19.1 模型与 Simulink 仿真测试报告\n')
head('1 概述');head('1.1 标识',2)
para('报告依据为用户修订的 SDTwin Power 设计报告 v19.1 和执行前确定的 10 项模型验证方案。被测核心为 Power/battery.py，模型采用一阶 ECM、SOC 与温度二维只读资产表及 EKF。运行时不做 RLS 参数辨识。')
head('1.2 软件与文档概述',2);para(abstract)
para('太阳、电池、配电和外部功率求解器组成独立试验台。物理状态为 SOC 与极化电压 u1，由外部求解器推进；估计器只读取电流、端电压和温度，估计状态不写回电池，也不参与功率分配。温度作为规定输入，本轮不积分 Thermal 温度状态。')
head('1.3 术语与方向约定',2)
table(['量','含义与单位'],[['z 与 SOC','剩余电量比例，无量纲。0.6 表示 60%。误差 0.02 等于 2 个百分点。'],['u1','RC 支路的极化电压，单位 V，表示加载与静置过程中的动态记忆。'],['OCV','无负载时的开路电压，由当前 SOC 和温度查表。'],['i 与电池功率','单体电流单位 A。放电为正，充电为负。电池端口功率为正表示向外供电。'],['太阳上限与实际取电','上限由辐照度和方向决定；充电受限时实际取电可低于上限。'],['RMSE 与最大误差','在同一时间点计算模型减参考。RMSE 为差值平方的均值再开平方；最大误差为绝对值最大值。']])
head('2 测试内容');head('2.1 十项测试与设计能力',2)
titles=['日照与朝向下的太阳发电','电池静态响应与整组换算','放电脉冲与静置恢复','充电脉冲与充放电反转','SOC 温度查表及响应','电阻发热与极化储能','完整充放电与容量曲线','三组实测脉冲留出验证','初始偏差与噪声下的 SOC 估计','供需变化下的功率分配']
refs=['4.2，P2','4.4.1 至 4.4.5，P4 至 P9','4.4.2、4.4.3，P5 至 P7','4.4.2、4.4.3，P5 至 P7','4.4.4，P8','4.4.5，P9','4.4.2、4.4.3，P5、P7','4.4、9.4','4.4.6，P11、P12','4.1、4.3、4.5，P1、P3、P13']
table(['编号','验证内容','设计章节与公式','运行数'],[[cid,title,ref,str(len(S['projects'][cid]['runs']))] for cid,title,ref in zip(C['projects'],titles,refs)])
head('2.2 输入资产与实验环境',2)
table(['资产','数值与用途'],[['A 解析核算','人工 2 Ah 单体，25 °C，OCV=3.2+z；R0=0.05 Ω，R1=0.02 Ω，τ=10 s，C1=500 F。SOC 范围 0.1 至 0.9。'],['B 温度插值','人工 2 Ah 单体，SOC 网格 0.1、0.5、0.9，温度网格 0、25、50 °C。由第 3.5 节显式函数生成，不代表真实电芯温度性能。'],['C 实测验证','BAK N18650CL-29，按来源示例采用 2.84 Ah，恒温设定 25 °C。R0=0.04448418743 Ω，R1=0.01120275491 Ω，τ=10.35558481 s。'],['默认初值','除 BT01、BT06、BT07、EK01 特别注明外，物理 z=0.6，u1=0 V。单体项目均为 1 串 1 并。'],['温度含义','输入接口使用 K。25 °C 输入 298.15 K；0、50 °C 分别输入 273.15、323.15 K。恒温是受控设定，不是已完成热平衡预测。']])
para('资产 C 复用已完成的离线标定数值，在本轮运行前冻结。训练使用脉冲 1、2、5、7、9 的放电加载段和相应静置电压；脉冲 4、6、8 留出。异常脉冲 3 不参与本轮标准验证。实测支持的 SOC 区间约 0.103 至 0.987，旧表的端点延伸不进入本轮。留出输入沿用已有的分段重采样文件，本轮重新运行模型，没有沿用旧版通过结论。')
head('2.3 连续时间仿真与对照方法',2)
para('Python 使用 SciPy DOP853 进行外部积分。Simulink 使用原生连续 Integrator 块保存 SOC 和 u1，并使用 ode45 求解。两侧相对容差 1×10⁻¹⁰、绝对容差 1×10⁻¹²。短时项目最大步长与保存间隔均为 0.1 s；BT02 另算 0.05、0.025 s；BT06 为 1 s。静态项目明确冻结 SOC 和 u1，只核算瞬时响应。')
para('输入电流按零阶保持，切换时间作为求解边界。完整容量试验在 Simulink 中由 SOC 触发 Relay 反转电流；Python 侧使用常电流电量公式确定对应时刻。估计器每 0.1 s 更新，先保存校正前预测电压，再读本次测量完成校正。两侧只共享参数和输入，不调用对方实现，也不以对方输出为输入。')
image('simulink_ecm','实际 Simulink 电池子系统。左侧输入单体电流和温度，中间两个连续积分器保存 SOC 与极化状态；电压、发热和串并联换算由显式方程计算。参数表仅读取，求根与 EKF 均在子系统之外。')
para('图下残差由相应上图所用数组直接相减，不使用另一次仿真曲线。BT06 保留固定时间网格原始 CSV，同时保存事件前后及其真实时间戳；仅反转样本按同一事件侧配对。未配对时的失败摘要和原始差值也保留。所有其他时间点保持原样。')
head('3 详细测试项目');cases=[]
case('PV01',titles[0],'验证 P2 的方向余弦、有效面积和日照输入对最大可发功率的作用。','面积 2 m²，效率 0.25，电池不带负载，温度 25 °C。','角度取 0、60、90、180 度，G=1000 W/m²。时间试验正对太阳，0 至 10 s 的 G=1000，10 至 20 s 为 500，20 至 40 s 为 0，40 至 60 s 为 1000 W/m²；再把面积改为 4 m²。','最大太阳功率；输入余弦与板面有效辐照度另按公式核算。','角度功率依次为 500、250、0、0 W；时间序列为 500、250、0、500 W；面积加倍为 1000 W。','角度与时序符合预期，逐点时间序列误差为 0 W，面积加倍输出 1000 W。','G 已含日食遮挡，G=0 就表示本测试的全遮挡环境，不另乘第二个日食因子。没有轨道位置或温度反馈。')
image('PV01','PV01 日照和朝向试验。上图横轴为秒，20 至 40 s 辐照度为零，因此最大可发功率为零；下图在固定辐照度下转动板面，90 度掠入射与 180 度背向都不发电。曲线表示太阳可发上限，不代表负载已取用这些功率。')
table(['入射角 度','输入余弦','有效辐照度 W/m²','Simulink W','理论 W'],[[str(a),n(np.cos(np.deg2rad(a))),n(1000*max(0,np.cos(np.deg2rad(a)))),n(value('PV01',f'angle_{a}_W')),str(p)]for a,p in [(0,500),(60,250),(90,0),(180,0)]])
case('BT01',titles[1],'核算 P4 至 P9 的单体端电压、功率、导数、热源和电池组倍率。','资产 A，恒温 25 °C，z=0.6，u1=0.02 V。冻结两项状态，只比较瞬时方程。','单体电流依次为 0、2、−2 A；分别配置 1 串 1 并和 4 串 2 并，共 6 组。','单体及整组电压、电流、功率、热源、RC 储能、SOC 与 u1 导数。','2 A 单体 V=3.68 V、P=7.36 W、热源 0.22 W、储能 0.1 J；dz/dt=−2/7200，du1/dt=0.002 V/s。4 串 2 并电压乘 4、电流乘 2、功率热源储能乘 8。','6 组全部通过手算及跨实现判据。单体状态导数不随组数重复放大。','0 A 仍有 0.02 W 极化耗散，因为初始 u1 不为零。冻结状态仅用于本项瞬时公式核算，动态项目均实际积分。')
image('BT01','BT01 瞬时核算表。Cell V 为单体端电压，Pack W 为整组端口功率，Heat W 为整组耗散。充电功率为负而热源为正。两组拓扑的状态导数相同，证明未把串并联倍率重复计入单体电量。')
static=json.loads((R/'BT01_hand_calculations.json').read_text())
case('BT02',titles[2],'验证 P5 至 P7：欧姆瞬态、极化建立、静置恢复和 SOC 电量积分。','资产 A，25 °C，z=0.6，u1=0 V，1 串 1 并。','0 至 10 s 为 0 A；10 至 30 s 为 2 A；30 至 60 s 为 0 A。最大步长依次 0.1、0.05、0.025 s。','电流、SOC、u1、OCV、欧姆项、端电压、解析与跨实现差值。','10 s 加载下降 0.1 V，30 s 撤流升高 0.1 V。20 s 放电后 SOC=0.5944444444，u1=0.03458658867 V。静置 SOC 不变而 u1 指数衰减。',f"Python 与 Simulink 最大电压差 {n(err('BT02','voltage_V'))} V；最细两级电压差 {n(value('BT02','refinement_005_vs_0025')['voltage_V'])} V。解析轨迹、状态和跳变量全部达标。",'残差在浮点舍入量级，不是物理噪声；实际测量误差由 BT07 另行评价。')
image('BT02','BT02 放电脉冲全过程。四行依次为实际电流、两侧端电压与 OCV、SOC 与极化电压、Python 电压减 Simulink 电压。25 °C 与初值见本项表。加载时立即损失 0.1 V，之后极化继续增加；撤流后欧姆压降消失，极化用 10 s 时间常数恢复。残差坐标的科学计数倍率必须一并读取。')
image('BT02_voltage_terms','BT02 电压分解。每条彩线分别为相对初始 OCV 的变化、−iR0 与 −u1；黑虚线为三项之和，即端电压相对 3.8 V 的变化。加载瞬态来自欧姆项，后续弯曲来自极化项，OCV 的缓慢下降来自 SOC 消耗。')
edges=json.loads((R/'pulse_edge_limits.json').read_text())
table(['时刻 s','电流前后 A','电压左极限 V','电压右极限 V','跳变量 V'],[[n(e['time_s']),n(e['current_left_A'])+' → '+n(e['current_right_A']),n(e['voltage_left_V'],10),n(e['voltage_right_V'],10),n(e['jump_V'])] for e in edges if e['case']=='BT02'])
para('表中左右极限在同一已接受边界状态上分别代入切换前后电流计算，不把 29.9 s 的电压当作 30 s 左极限。完整解析解独立检查边界状态和区间轨迹。')
case('BT03',titles[3],'验证 P5 至 P7 的充电方向、反向电流和极化记忆。','资产 A，25 °C，z=0.6，u1=0 V。','0 至 10 s 为 0 A；10 至 30 s 为 −2 A；30 至 50 s 为 2 A；50 至 80 s 为 0 A。','端电压、SOC、u1、OCV、电压三项贡献和端口功率。','初始充电电压升高 0.1 V，SOC 增加；30 s 反转时电压下降 0.2 V，极化不清零；充放电净电量为零，50 s 后 SOC 回到 0.6。',f"解析与跨实现判据均通过。跨实现最大电压差 {n(err('BT03','voltage_V'))} V；反转瞬态为 −0.2 V。",'撤去充电电流时电压回落并不表示电量突然减少；电压由 OCV、欧姆项和保留的极化状态共同决定。')
image('BT03','BT03 80 s 充放电与静置全过程。10 s 开始充电，端电压由 3.8 V 升至 3.9 V，随后继续上升；30 s 转为放电时立即下降 0.2 V。SOC 与极化连续，极化由负值逐渐转正。最下图直接计算两条模型电压之差。')
image('BT03_voltage_terms','BT03 电压分解。30 s 电流反转只立即改变 −iR0，原有负极化保留并逐渐消退。此分解解释了反向初期的电压行为，而不是按充电或放电标签强行规定整条曲线形状。')
table(['时刻 s','电流前后 A','电压左极限 V','电压右极限 V','跳变量 V'],[[n(e['time_s']),n(e['current_left_A'])+' → '+n(e['current_right_A']),n(e['voltage_left_V'],10),n(e['voltage_right_V'],10),n(e['jump_V'])] for e in edges if e['case']=='BT03'])
case('BT04',titles[4],'验证 P8 的二维双线性插值、单位换算和参数影响。','人工资产 B，容量 2 Ah，初始 z=0.6、u1=0 V。各次温度保持不变。','25 个 SOC 温度查询覆盖所有节点、网格中心及边上中点。另分别在 0、25、50 °C 执行与 BT02 相同的 2 A 脉冲。','OCV、R0、R1、τ、C1、端电压与极化。','z=0.3、T=12.5 °C 时 OCV=3.505 V、R0=0.0545 Ω、R1=0.021425 Ω、τ=10.65 s；C1=τ/R1。全部查询按显式函数独立核对。',f"25 次查询及三温度动态轨迹全部通过。三温度最大跨实现电压差 {n(err('BT04','voltage_V'))} V；四个紧邻支持域外的输入均被拒绝。",'温度影响由人为函数定义。此项证明查表和动态耦合实现正确，不能说明真实 BAK 电芯在 0 或 50 °C 的精度。')
para('生成 B 表的函数如下，T 单位为 °C，表行对应温度、列对应 SOC。OCV=3.2+z−0.001×ΔT+0.002×z×ΔT；R0=0.05+0.01×0.5−0.01×z−0.0002×ΔT；R1=0.02+0.004×0.5−0.004×z−0.00005×ΔT；τ=10+2×0.5−2×z−0.02×ΔT。其中 ΔT=T−25，C1 始终由 τ/R1 得到。')
image('BT04','BT04 三个恒温环境的电压和极化响应。输入脉冲相同，0 °C 人工表的电阻较大，加载压降和极化较大；50 °C 相反。OCV 本身也随温度改变，因此应结合下表分辨基线与内阻效应。这里只比较人工表机制。')
table(['初始温度 °C','OCV V','R0 Ω','R1 Ω','τ s','C1 F'],[[str(T)]+[n(read(f'BT04_T{T}')[f][0],7)for f in ['ocv_V','r0_ohm','r1_ohm','tau_s','c1_F']]for T in [0,25,50]])
case('BT05',titles[5],'验证 P9 的两项耗散与 RC 暂存能量，避免把内部电压损失全部计为热。','常参数资产 A，25 °C，z=0.6、u1=0 V，C1=500 F。','采用 BT03 的 80 s 时序，含充电、反转和最后 30 s 静置。','i²R0、u1²/R1、总热功率、0.5C1u1²、累计热量和能量收支残差。','常 C1 时，积分 i×OCV 与积分端口功之差，等于累计热量加 RC 储能变化；热源始终非负，停流后极化未消失时仍发热。',f"Simulink 最大能量残差 {n(value('BT05','simulink_energy_residual_J'))} J；Python 为 {n(value('BT05','python_energy_residual_J'))} J。两项热源及跨实现轨迹通过。",'没有实测热量。本项为方程能量核算；常 C1 恒等式不直接应用于变参数 B，也未覆盖可逆熵热、老化或副反应热。')
image('BT05','BT05 热源与储能核算。第一行区分即时欧姆耗散和极化电阻耗散；第二行比较累计内部电功差、累计热量和当前 RC 储能；第三行为第一种累计能量减去后两者。初始 RC 储能为零。50 s 停流后，欧姆耗散为零而极化耗散仍为正。')
rr=read('BT05');table(['时刻 s','欧姆热 W','极化热 W','RC 储能 J','累计热量 J'],[[n(t)]+[n(rr[f][int(t*10)],8)for f in ['ohmic_heat_W','polarization_heat_W','polarization_energy_J','heat_energy_J']]for t in [10,30,50,80]])
case('BT06',titles[6],'验证 P5、P7 的长时电量积分、SOC 边界和容量曲线。','人工资产 A，25 °C，初始 z=0.12、u1=0 V；最大步长 1 s。','以 −0.2 A 充至 SOC=0.88，再以 0.2 A 放回 0.12；即 0.1C，一次往返，反转时不重置极化。','电压、SOC、累计绝对 Ah、两段电量及事件时间和事件两侧状态。','每个方向 1.52 Ah，每段 27360 s，总计 54720 s；电量误差不超过 10⁻⁶ Ah，反转事件时间误差不超过 10⁻⁴ s。',f"充入 {value('BT06','charge_Ah'):.12f} Ah，放出 {value('BT06','discharge_Ah'):.12f} Ah；反转时间偏差 {n(E['time_error_s'])} s。同侧配对后电压最大差 {n(err('BT06','voltage_V'))} V，全部通过。",'固定网格在 27360 s 原始比较出现 0.02 V 差，原因是两侧相差约 17 ns 而分别落在反转前后。保留原始失败记录，仅按既定规则配对相同事件侧。人工线性 OCV 决定了容量曲线近似直线，不能用此曲线代表真实材料平台或老化。')
image('BT06','BT06 容量曲线与完整电量历程。上图横轴是各充、放电分支分别从零累计的 Ah，因此两条分支可能相交；没有活性物质量数据，不使用 mAh/g。下图给出真实物理时间，完整 0.1C 往返确需 54720 s。资产 A 的 OCV 对 SOC 为线性，因此曲线形状是核算参数的结果。')
image('BT06_comparison','BT06 完整时域对比。上图显示电流反转，中图为两套实现的端电压，下图直接相减。反转样本使用 Simulink 的事件右侧值与 Python 右侧值配对，其真实来源时间保存在事件表中，其他样本未作调整。')
table(['反转侧','Python 时间 s','Simulink 来源时间 s','SOC','u1 V','端电压 V'],[['左侧','27360',f"{E['source_left_s']:.10f}",'0.88','−0.004','4.094'],['右侧','27360',f"{E['source_right_s']:.10f}",'0.88','−0.004','4.074']])
para('事件两侧电流变化 0.4 A，欧姆压降变化为 0.4×0.05=0.02 V；状态连续。仅用两个不同事件侧的电压相减会把正确的物理跳变当成数值误差。对齐审计同时保存原始时刻、原始误差、配对误差及选取规则，没有修改原始 CSV 或 MAT。')
case('BT07',titles[7],'验证第 4.4、9.4 节：固定参数的一阶 ECM 对未参与标定的实际端电压是否有效。','资产 C，25 °C 恒温设定。每段初始 u1=0 V，初始 SOC 由完整记录的电流积分获得；不按该段电压重新拟合初值。','留出脉冲 4、6、8，每组 0 至 119.9 s，间隔 0.1 s。原始电流反号后输入；约 30 s 放电、40 s 静置、10 s 充电、40 s 静置，实际幅值采用记录。','实测与开环模型电压、直接残差、整段及分阶段 RMSE 和最大误差。','每组整段 RMSE≤30 mV，最大绝对误差≤100 mV；物理状态只由电流推进，不利用测量电压纠正。','三组整段 RMSE 为 26.70、28.65、17.27 mV；最大绝对误差为 60.25、59.47、50.51 mV，均通过。','部分充电和静置阶段 RMSE 约 35 至 41 mV，高于整段 30 mV 的量级。整段通过不等于每阶段都低于 30 mV。SOC 由电量推算，缺少独立真值；0.1 s 是重采样间隔，不是原传感器采样率。')
table(['留出脉冲','初始 SOC','放电峰值 A','充电峰值 A','整段 RMSE mV','最大误差 mV'],[[str(j),n(json.loads((HERE/f'fixtures/BT07_P{j:02d}.json').read_text())['initial'][0],8),n(max(read(f'BT07_P{j:02d}')['current_A']),6),n(min(read(f'BT07_P{j:02d}')['current_A']),6),n(1000*value('BT07',f'BT07_P{j:02d}')['rmse'],6),n(1000*value('BT07',f'BT07_P{j:02d}')['max_abs'],6)]for j in [4,6,8]])
for j in [4,6,8]:
    image(f'BT07_P{j:02d}',f'BT07 留出脉冲 {j}，25 °C。第一行为输入实测电流；第二行蓝色和橙色为两套开环模型，黑色为测量电压；第三行为 Simulink 电压减测量电压，单位 mV；第四行为 Python 减 Simulink，单位 V。黑色测量曲线未送入物理模型，因此第三行是真正的开环预测误差。')
    m=value('BT07',f'BT07_P{j:02d}')['segments']
    table(['阶段','区间 s','RMSE mV','最大误差 mV'],[[lab,interval,n(1000*m[key]['rmse'],5),n(1000*m[key]['max_abs'],5)]for key,lab,interval in [('discharge','放电','0 至 30'),('rest_after_discharge','放电后静置','30 至 70'),('charge','充电','70 至 80'),('rest_after_charge','充电后静置','80 至 120')]])
para('观察到的偏差主要集中于训练中未覆盖的充电与长静置阶段。固定 RC 参数、OCV 与初始电量误差都可能造成这些偏差，但当前证据不能唯一归因。后续若改参数，必须使用训练数据重新标定并另存版本，不能用这三段留出曲线调到通过。')
case('EK01',titles[8],'验证 P11、P12 的因果 SOC 状态估计，保留一个核心估计项目。','资产 A，25 °C；物理 z=0.6、u1=0，估计初值 z=0.7、u1=0。物理与估计状态分离。','每 10 s 依次施加 0、1.5、0、−1、2、0、−1.5、0、1、−0.5 A，重复三次，共 300 s。电压噪声标准差 3 mV，种子 20261008，EKF 周期 0.1 s。','SOC 真值与估计、估计极化、SOC 误差、校正前预测电压、协方差。','最后 30 s SOC RMSE 和最终绝对误差≤0.02；协方差有限、对称、半正定；估计轨迹与独立 MATLAB 实现一致。',f"最后 30 s SOC RMSE 为 {n(100*value('EK01','soc_accuracy')['last_30s_rmse'],6)} 个百分点，最终误差 {n(100*value('EK01','soc_accuracy')['final_abs'],6)} 个百分点。首次测量校正后即进入 2 个百分点范围，协方差检查通过。",'人工 OCV 斜率为 1 V 每单位 SOC，初始 SOC 方差较大，因此首次电压观测即可显著纠偏。不能据此声称真实电池在任意 SOC 都立即收敛。首点先验电压仍含约 100 mV 初始偏差，未从指标中删除。')
image('EK01','EK01 300 s 噪声试验。第二行红叉明确标出校正前初始 SOC=0.70，连续曲线为每次测量后的估计与物理真值；第三行为后验 SOC 减真值，单位为百分点；第四行为校正前电压预测减本次测量，保留首点约 100 mV 误差。不能将后验电压当作先验预测精度。')
table(['配置或结果','数值'],[['P0 对角元','0.01、0.0001'],['每步过程协方差对角元','1×10⁻⁹、1×10⁻⁷'],['电压测量方差','9×10⁻⁶ V²'],['先验电压 RMSE，包含首点',n(1000*value('EK01','soc_accuracy')['prior_voltage']['rmse'],6)+' mV'],['协方差最小特征值',n(value('EK01','covariance')['min_eigenvalue'])],['跨实现最大 SOC 估计差',n(value('EK01','independent_estimator')['estimated_soc'])]])
case('PA01',titles[9],'验证 P1、P3、P13 的供需平衡、可行电流根、限充和供电不足断开。','资产 A，4 串 2 并，25 °C，z=0.6、u1=0，负载初始接通；配电效率 0.9，单体电流界限 −3 至 5 A，电压界限 2.8 至 4.3 V。','每段 20 s。太阳上限依次为 500、500、500、500、0 W；请求负载依次为 450、540、405、90、180 W，总计 100 s。','太阳上限与实际输出、电池电流与功率、请求与实际负载、配电热、SOC、电压、连接状态和残差。','五段分别为电池零输出、补足 100 W、吸收 50 W、按 −3 A 限充并弃光、供电不足后实际负载降至 0 W。每点按实际太阳功率核算守恒。',f"五种模式及电压电流约束均通过；最大端口功率残差 {n(value('PA01','max_balance_residual_W'))} W。80 s 断开前最大可交付负载约 {n(value('PA01','maximum_load_before_disconnect_W'),7)} W，确实低于 180 W。",'断开后的零供电是本测试保护结果，不等于满足了请求负载。此项仅含一条必要模式序列，未扩展为整星启停与故障矩阵。')
image('PA01','PA01 100 s 功率分配。上两行区分太阳上限、实际取电、电池端口、请求负载、交付负载与配电热。60 s 开始限充，太阳实际取电下降；80 s 无日照且电池能力不足，交付负载归零。第三行为单体电流；第四行为端口守恒残差以及两套实现的电池功率差。')
pa=read('PA01');table(['时刻 s','太阳实际 W','电池 W','交付负载 W','配电热 W','单体电流 A'],[[str(t)]+[n(pa[f][t*10],7)for f in ['actual_solar_W','power_W','load_W','distribution_heat_W','current_A']]for t in [0,20,40,60,80]])
para('配电入口需求等于请求负载除以 0.9。例如请求 540 W 时入口需要 600 W，太阳提供 500 W，电池补 100 W，配电器耗散 60 W。电池内部热已体现在端电压与功率中，不在 P1 中重复扣除。求解器在当前真实状态下求可行根，EKF 不参与。')
head('4 测试内容充分性分析',new=True)
para('10 项覆盖太阳输入、静态电池、秒级充放电、SOC 温度查表、耗散与储能、完整容量窗口、留出实测电压、SOC 估计和基本功率分配。26 次运行属于这 10 个项目内部的参数配置及步长核算，不新增正式测试项目。')
para('解析解、手算和能量恒等式用于发现两套实现可能共同存在的公式错误；独立 Simulink 用于检查跨实现一致性；留出实测电压用于判断 25 °C 电芯预测精度。这三类证据分别有不同作用，数值一致不能替代物理准确性。')
para('覆盖边界明确为：多温度只验证人工表，真实温度性能未测；EKF 只验证合成已知 SOC；真实 SOC、实际发热、老化、单体不一致、可逆热和完整 Thermal 闭环未验证。接口装配、工程异常矩阵和完整轨道联调不计入本轮 10 项。')
head('5 测试条件与要求');head('5.1 固定判据',2)
table(['对象','执行前固定的门限'],[['静态与插值','绝对误差≤1×10⁻¹⁰ 加参考量绝对值的 1×10⁻¹⁰。'],['动态对照','SOC≤1×10⁻⁷；u1≤1×10⁻⁶ V；单体电压≤1×10⁻⁵ V；单体电流≤1×10⁻⁵ A；功率与热源≤1×10⁻⁴ W。'],['步长核算','0.05 与 0.025 s 的差异低于上述 SOC、u1、电压门限的四分之一。'],['端口守恒','残差≤1×10⁻⁶ W 加当时最大端口绝对功率的 1×10⁻⁹。'],['实测电压','每组完整留出段 RMSE≤30 mV，最大绝对误差≤100 mV。'],['EKF SOC','最后 30 s 的 RMSE 和最终绝对误差均≤0.02。'],['BT05 能量','残差≤1×10⁻⁵ J 加最大累计内部电功差绝对值的 1×10⁻⁶。'],['BT06 电量与时刻','单方向电量误差≤1×10⁻⁶ Ah；反转时刻误差≤1×10⁻⁴ s；同事件侧比较电压。']])
head('5.2 执行环境与可复现设置',2)
meta=json.loads((R/'simulink_execution.json').read_text())
table(['环境','配置'],[['MATLAB 与 Simulink',meta['matlab']+'；Simulink '+meta['simulink']['Version']],['Python','Thermal/.venv，SciPy DOP853；纯电池响应不拥有积分时间。'],['Simulink 状态','5 个原生连续积分器：SOC、u1、累计热量、累计内部电功差、累计绝对 Ah。后二项状态不反馈物理电池。'],['离散模块','参数查询及外部功率分配为独立 MATLAB 方程；EKF 使用独立离散状态，每次已接受观测更新一次。'],['可执行模型','matlab 下每个项目对应同名 slx，另有通用模型 sdtwin_power_v191.slx。装载与启动回调自动加载相应配置。'],['源码范围','Power/battery.py 为本版纯 ECM 与 EKF；外部 Python 积分和求根在 tests_v19_1/run_python.py。历史 RLS 实现独立保留。']])
para('这是以设计方程构建的独立 Simulink 模型，使用原生连续积分器，不是 Simscape Battery 厂商预设电芯块。模型验证的输入和参数来自本报告定义的资产。')
head('6 测试数据');head('6.1 文件与重现入口',2)
table(['目录或文件','保存内容'],[['fixtures','26 份时间序列及配置，A、B、C 冻结参数，BAK 数据许可，执行前方案副本。'],['results','Python 与 Simulink CSV、原始 Simulink MAT、逐点差值、手算、能量与事件审计、summary.json。'],['figures','每项图表的 PNG 和 PDF；原生 Simulink 电池子系统与整体图。'],['matlab','10 个项目模型、通用模型、模型构建器、独立公式、运行入口。'],['reports','本报告 Markdown，以及报告布局与溯源记录。'],['README.md','完整重现命令、字段定义、结果摘要和已知适用边界。']])
para('从仓库根目录运行 prepare.py 冻结输入，再运行 run_python.py；MATLAB 中执行 addpath 后运行 run_suite；最后执行 analyze.py。build_report.py 只读取结果并生成报告。用户查看模型时可直接打开同名 slx 后运行，不需要 Python 回调。完整命令见本试验目录 README。')
head('6.2 原始数据与版本证据',2)
para('BAK 数据来自 MathWorks 的 Estimate Battery Model Parameters from HPPC Data 示例，原始记录作者为 Anandaroop Bhattacharya 与 Subhasish Basu Majumder，IIT Kharagpur。许可随导出输入保存。原始电流以放电负号记录，进入当前模型前反号。电流保持及电压同工作段内重采样的旧输入文件按哈希固定，本轮没有再拟合留出电压。')
table(['证据','SHA256 前 16 位'],[['执行前方案',C['plan_sha256'][:16]],['用户设计报告 v19.1',C['design_sha256'][:16]],['25 °C 原始数据',C['raw_data_sha256'][:16]],['离线标定记录',C['calibration_record_sha256'][:16]],['保留的 Orbit 格式来源',source_hash[:16]]])
para('完整哈希在 contract.json 与 evidence_manifest.json。BT06_alignment_audit.json 记录唯一事件配对。summary_before_event_alignment.json 保存最初按固定网格比较得到的 9 项通过、1 项失败；该失败原因及后续配对规则已在第 3.7 节说明。判据和原始轨迹均未覆盖。')
head('6.3 参考资料',2)
for text in ['用户修订的 SDTwin_Power_Design_Report_CN_v19.1.docx；本轮设计依据。','Power/TEST_PLAN.md 的执行前冻结版本；本轮范围与判据。','用户提供的《用 MATLAB Simulink 做锂电池建模》；脉冲、静置回稳与独立工况验证方法。','MathWorks. Estimate Battery Model Parameters from HPPC Data. https://www.mathworks.com/help/simscape-battery/ug/estimate-battery-model-parameters-from-hppc-data.html','MathWorks. Integrator. https://www.mathworks.com/help/simulink/slref/integrator.html','MathWorks. From Workspace. https://www.mathworks.com/help/simulink/slref/fromworkspace.html']:para(text)
head('7 测试总结',new=True)
table(['编号','测试项目','结论'],cases)
para('本轮 10 项满足预设模型判据。所有动态项目已沿完整物理时间积分，未用少量采样点替代动态过程。三组留出实测电压整段达标，同时保留充电与静置阶段的系统偏差。')
para('可支持的结论是：v19.1 一阶 ECM、只读参数表、EKF 与基本功率分配在本报告列出的输入域中完成数值核算和独立 Simulink 对照；BAK 25 °C 三个留出片段完成电压验证。结论不扩展到实际航天电芯、多温度实物、真实 SOC、发热实测、老化或整星热电联调。')
table(['项目','最大电压差 V','最大 SOC 差','比较点数'],[[cid,n(err(cid,'voltage_V')),n(err(cid,'soc')),str(sum(S['comparisons'][c]['rows']for c in S['projects'][cid]['runs']))]for cid in C['projects']])
for e in d.element.body.iter(qn('w:trHeight')):e.set(qn('w:hRule'),'atLeast')
output=POWER/'SDTwin_Power_Model_Test_Report_CN_v19.1.docx';d.save(output)
(HERE/'reports/TEST_REPORT_CN.md').write_text('\n'.join(md),encoding='utf8')
assert hashlib.sha256(source.read_bytes()).hexdigest()==source_hash
(HERE/'reports/artifact.md').write_text('''# Retained report template and fill contract

Source: Power/tests_v19_1/templates/orbit_power_test_source.docx. This is the retained Orbit-derived
report, with cover, revision record, contents, seven numbered chapters and an
eleven-row case form. Original source remains unchanged.

Editable slots: cover subtitle, scope line, date, abstract, revision row, body
chapters, ten cloned case forms, result tables and figure captions.
Preserved: cover organization, typography, section geometry, header/footer,
TOC field, merged-cell case form and original approval blanks. No reviewer is invented.

The new report is a separate v19.1 artifact. The user design report is immutable.
Numerical content is read from current results. Raw event-time mismatch and its
same-side correction are disclosed; threshold values remain those frozen before runs.
All final pages must be rendered and visually inspected before delivery.
Source SHA256: '''+source_hash+'\n',encoding='utf8')
print(output)
