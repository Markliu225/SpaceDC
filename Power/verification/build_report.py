"""Bilingual evidence-led Power versus Simulink reports and scientific figures."""
import json,sys,math,hashlib
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from prepare import ROOT,OUT,FIELDS
from report_content import TEXT,PLOT,LABELS

HERE=Path(__file__).resolve().parent
FIG=HERE/'figures';FIG.mkdir(exist_ok=True)
SUITE=json.loads((OUT/'suite.json').read_text(encoding='utf-8'))
SUMMARY=json.loads((OUT/'summary.json').read_text(encoding='utf-8'))
RESULT={x['id']:x for x in SUMMARY['cases']}
plt.rcParams.update({'font.family':'Microsoft YaHei','axes.unicode_minus':True,'font.size':9,
                     'axes.spines.top':False,'axes.spines.right':False,'figure.dpi':150})
def pick(cn,en,lang):return cn if lang=='cn' else en
def fmt(x):
    if isinstance(x,(bool,np.bool_)):return str(int(x))
    if isinstance(x,(int,float,np.floating)):
        if not np.isfinite(x):return '—'
        if x==0:return '0'
        s=f'{x:.6g}' if abs(x)>=1e-3 else f'{x:.2e}'
        return s.replace('-', '−')
    return str(x)
def csv(cid,side):return np.genfromtxt(OUT/f'{cid}_{side}.csv',delimiter=',',names=True)
def label(n,lang):return LABELS[n][lang=='en']

def figure(c,lang):
    cid=c['id'];a=csv(cid,'python');b=csv(cid,'simulink_aligned');d=csv(cid,'difference');signals=PLOT[cid]
    fig,axes=plt.subplots(2,len(signals),figsize=(7.15,3.85),squeeze=False,gridspec_kw={'height_ratios':[2.4,1.25]},layout='constrained')
    x=a['time_s']/60 if c['kind']=='dynamic' else np.arange(1,len(a)+1)
    for j,n in enumerate(signals):
        ax=axes[0,j];adx=axes[1,j]
        ax.plot(x,a[n],color='#1674bc',lw=1.9,label=pick('本项目模型','Project model',lang))
        ax.plot(x,b[n],color='#e66b25',lw=1.25,ls='--',marker='.' if c['kind']=='static' else None,label='Simulink')
        ax.set_ylabel(label(n,lang));ax.legend(fontsize=7.5,loc='best');ax.grid(alpha=.22)
        ax.ticklabel_format(axis='y',style='plain',useOffset=False)
        values=np.r_[a[n],b[n]];values=values[np.isfinite(values)]
        span=.05 if n in ['SOC','x_n','x_p'] else .1 if n.endswith('_A') else .1 if n.endswith('_V') else 1.
        if len(values) and np.ptp(values)<1e-6*max(1.,abs(float(np.mean(values)))):
            middle=float(np.mean(values));ax.set_ylim(middle-span/2,middle+span/2)
        if cid=='CT-003' and n=='i_A':ax.set_ylim(min(values)-.05,.05)
        if n=='SOC':ax.axhline(1.,color='#b52f38',ls=':',lw=1)
        if cid=='CT-003' and n=='i_A':ax.axhline(0.,color='#b52f38',ls=':',lw=1)
        actual=a[n]-b[n]
        assert np.allclose(actual,d[n],equal_nan=True,rtol=0,atol=1e-13)
        adx.plot(x,actual,color='#7446a2',lw=1.2,marker='.' if c['kind']=='static' else None)
        adx.axhline(0,color='#666',lw=.6);adx.grid(alpha=.22)
        adx.set_ylabel(pick('模型减 Simulink','Model minus Simulink',lang),fontsize=8)
        adx.set_xlabel(pick('仿真时间 min','Simulation time min',lang) if c['kind']=='dynamic' else pick('输入组合编号','Input combination index',lang))
        adx.ticklabel_format(axis='y',style='sci',scilimits=(-2,3),useOffset=False)
        if c['kind']=='static' and len(a)<=12:adx.set_xticks(x)
    fig.savefig(FIG/f'{cid}_{lang}.png',dpi=200);plt.close(fig)

def build(lang):
    en=lang=='en';L=lambda a,b:pick(a,b,lang);doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8)
    sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1);sec.header_distance=Cm(.85);sec.footer_distance=Cm(.9)
    for name in ['Normal','Body Text','Caption','Heading 1','Heading 2','Heading 3','Title']:
        s=doc.styles[name];s.font.name='Times New Roman';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        s.font.size=Pt(10.5);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.12
    for s in doc.styles:
        for border in list(s._element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
    for name,size in [('Title',23),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].paragraph_format.space_after=Pt(5)
    head=sec.header.paragraphs[0];head.text=L('SDTwin 供电模块测试报告','SDTwin Power Module Test Report');head.style=doc.styles['Caption']
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
    def p(t,style=None):return doc.add_paragraph(t,style)
    def h(t,level=1):return doc.add_heading(t,level)
    def page():doc.add_page_break()
    def table(headers,rows,widths=None,size=9):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
        if widths:
            for col,width in zip(t.columns,widths):col.width=Cm(width)
        for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
        rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
        for row in rows:
            for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
        for i,row in enumerate(t.rows):
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
            for j,cell in enumerate(row.cells):
                if widths:cell.width=Cm(widths[j])
                for par in cell.paragraphs:
                    par.paragraph_format.space_after=Pt(2);par.paragraph_format.space_before=Pt(2);par.paragraph_format.line_spacing=1.02
                    for run in par.runs:run.font.size=Pt(size);run.bold=i==0
                if i==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');cell._tc.get_or_add_tcPr().append(sh)
        p('');return t
    def caption(t):p(t,'Caption')
    def verdict(res):return L('通过','Pass') if res['status']=='pass' else L('不通过','Fail')
    def observed(c):
        a=csv(c['id'],'python');r=RESULT[c['id']]
        if c['id']=='CT-003':return L(f'三个满电点均继续充电，单体电流最小值为 {fmt(np.min(a["i_A"]))} A。两边数值一致，但不满足满电停止充电要求。',f'All three full-charge samples continue charging, with minimum cell current {fmt(np.min(a["i_A"]))} A. Numerical agreement passes, but the full-charge requirement fails.')
        if c['id'] in ['DY-002','DY-007']:
            return L(f'SOC 峰值为 {100*np.max(a["SOC"]):.4f}%，超过配置中定义的 100% 满电点。两边均出现这一现象，故总体验收不通过。',f'Peak SOC is {100*np.max(a["SOC"]):.4f}%, above the configured 100% point. Both implementations exhibit this behavior, so overall acceptance fails.')
        if c['id']=='BT-003':return L(f'整组热源范围为 {fmt(np.min(a["Q_B_W"]))} 至 {fmt(np.max(a["Q_B_W"]))} W，负热源被保留。',f'Pack heat ranges from {fmt(np.min(a["Q_B_W"]))} to {fmt(np.max(a["Q_B_W"]))} W; negative heat is retained.')
        if c['id']=='DY-005':return L('失电发生在 100 s；日照于 200 s 恢复后负载仍断开；300 s 显式启动、400 s 停机和 500 s 再启动均与预期一致。','Supply trips at 100 s; the load remains disconnected after sunlight returns at 200 s. Explicit start at 300 s, stop at 400 s and restart at 500 s agree with expectations.')
        if c['kind']=='dynamic':return L(f'初始 SOC 为 {100*a["SOC"][0]:.3f}%，结束时为 {100*a["SOC"][-1]:.3f}%；电池温度由 {a["T_B_K"][0]:.3f} K 变化至 {a["T_B_K"][-1]:.3f} K。',f'SOC changes from {100*a["SOC"][0]:.3f}% to {100*a["SOC"][-1]:.3f}%; battery temperature changes from {a["T_B_K"][0]:.3f} K to {a["T_B_K"][-1]:.3f} K.')
        return L(f'共完成 {r["samples"]} 个输入组合，输出可用性、离散标志和适用的数值误差检查全部通过。',f'Completed {r["samples"]} input combinations. Output availability, discrete status and applicable numerical difference checks pass.')
    # Cover and an explicit scope statement.
    p('NTU SPACE DATA CENTER','Subtitle');p('');p('');p(L('SDTwin 供电模块','SDTwin Power Module'),'Title')
    p(L('模型与 Simulink 仿真对比测试报告','Model Comparison with Simulink Simulation'),'Title')
    p(L('完整测试框架 用例 输入 逐项结果与缺陷','Test framework, cases, inputs, individual results and findings'))
    p('');p(L('版本 1.0','Version 1.0'));p(L('2026年10月7日','7 October 2026'))
    p(L('20 项已执行    17 项验收通过    3 项验收不通过','20 cases executed    17 accepted    3 failed acceptance'))
    p(L('20 项数值对比均通过，不能据此认定全部设计行为正确。','All 20 numerical comparisons pass; this does not establish that every design behavior is correct.'))
    page();h(L('1 概述与测试对象','1 Overview and systems under test'))
    p(L('本报告比较本项目 Power 模型与 Simulink 仿真在相同参数、状态和输入下的结果。测试覆盖太阳阵发电、电池电压与热源、串并联换算、功率分配、保护指令、动态轨迹、能量守恒及求解精度。','This report compares the project Power model with Simulink under identical parameters, states and inputs. Tests cover array power, cell voltage and heat, pack scaling, allocation, protection commands, dynamic trajectories, conservation and numerical precision.'))
    table([L('对象','Object'),L('实现与边界','Implementation and boundary')],[
      [L('本项目模型','Project model'),'Thermal/sdtwin_sim/power_stand_in.py'],
      [L('实现状态','Implementation status'),L('按 Power 设计报告实现的测试用模型，尚不是正式生产供电软件；原文件保持不变。','A test implementation of the Power design report, not production EPS software. The source under test is unchanged.')],
      ['Simulink',L('独立 MATLAB 方程，通过 Level 2 S 函数放入真实 .slx 模型，连续状态由 Simulink ode45 积分。','Independent MATLAB equations in a Level 2 S-function in a real .slx model, with continuous states integrated by Simulink ode45.')],
      [L('独立性','Independence'),L('双方共享输入参数，不共享计算结果；MATLAB 不调用 Python，Python 不读取 Simulink 输出作为参考状态。','Only input parameters are shared. MATLAB does not call Python; Python does not take Simulink outputs as its state trajectory.')],
      [L('参数性质','Parameter provenance'),L('太阳阵和电池参数沿用现有示例资产，电極材料曲线为示例多项式，未经目标实物电芯标定。','Existing illustrative array and battery assets, including illustrative electrode polynomials, not fitted to a selected physical cell.')],
      [L('可支持的结论','Supported conclusion'),L('同一物理假设下的实现一致性与选定设计行为验证。不能等同于实测验证，也不是与 Simscape 内置电池块的对比。','Cross-implementation verification under common physics assumptions and selected design behavior checks. This is neither measured validation nor a comparison against a factory Simscape battery block.')]
    ],[3.2,13.5])
    h(L('实际结论','Observed conclusion'),2)
    p(L('全部 20 个项目完成执行，数值对比均通过。满电边界 CT-003、单周期 DY-002 和三周期 DY-007 不满足满电相关要求。根因证据是两侧均只限制电流、电压与材料占比，没有将可用电量的 100% 端点纳入充电约束。','All 20 projects executed and passed numerical comparison. CT-003, DY-002 and DY-007 fail full-charge behavior requirements. Both implementations constrain current, voltage and material fractions but omit the configured 100% usable-charge endpoint from the charging constraint.'))
    p(L('失败项保留原始输入与结果。本轮没有修改受测模型，也没有放宽误差门限或裁剪 SOC 来消除失败。','Failed cases retain their inputs and raw results. The model under test was not changed, tolerances were not relaxed, and SOC was not clipped to hide the failures.'))
    page();h(L('2 测试框架','2 Test framework'))
    steps=[
      ('共同输入','Shared inputs','suite.json 保存器件参数、初值、输入组合、事件时间、积分设置和验收门限；SHA256 锁定输入版本。','suite.json stores device parameters, initial conditions, input combinations, event times, numerical settings and thresholds, bound by SHA256.'),
      ('分别求解','Independent execution','Python 运行现有接口并采用 SciPy DOP853 积分；MATLAB 独立求解电池与分配方程，由 Simulink ode45 积分。','Python calls the existing API with SciPy DOP853 integration; MATLAB independently evaluates cell and allocation equations while Simulink ode45 integrates states.'),
      ('事件与时间','Events and time','所有预设输入突变点精确分段，分别保存同一时间的左侧与右侧。保护只在已接受的边界处理。','Scheduled discontinuities split integration intervals exactly. Both left and right values are retained at each boundary. Protection is applied only at accepted boundaries.'),
      ('逐项比较','Signal comparison','同一时间、同一侧、同一单位的值直接相减；不跨事件插值。无效结果比较有效性和事件标志，不把缺失端口替换成零。','Subtract values at the same time, side and unit. No interpolation crosses an event. Invalid results compare validity and event flags without replacing missing ports by zero.'),
      ('独立判据','Independent acceptance','数值一致性之外检查功率与锂守恒、太阳可用功率、电量范围、满电停止充电和预期启停行为。','In addition to numerical agreement, check power and lithium conservation, solar feasibility, usable SOC range, full-charge stopping and expected protection behavior.'),
      ('证据与报告','Evidence and report','保存两侧 CSV、直接差值 CSV、事件 JSON、每项判据和源文件哈希；每个项目均给出输入表、误差表与对比图。','Retain both CSV trajectories, direct difference CSVs, event JSON, check verdicts and source hashes. Every case includes an input table, error table and comparison figure.')]
    table([L('环节','Stage'),L('执行约定','Execution contract')],[[L(a,b),L(c,d)] for a,b,c,d in steps],[3.2,13.5])
    h(L('物理关系与状态','Physical relations and states'),2)
    p('Ppv + PB = Pload + QD     PB = Ns Np i v     Pload = ηD Pbus')
    p('dxn/dt = −i/Qn     dxp/dt = i/Qp     SOC = (xn − xn,0)/(xn,100 − xn,0)')
    p(L('正电流与正电池功率均表示放电；负值表示充电。QB 是电池内部热源，不在使用电池端口功率的守恒式中重复扣除。电池温度在热域只积分一次。','Positive cell current and battery power denote discharge; negative values denote charging. QB is internal battery heat and is not subtracted again from terminal-power balance. Battery temperature is integrated once in the thermal domain.'))
    p(L('动态电热用例使用 CB dTB/dt = QB − (TB − Tsink)/R。其余动态用例指定恒温，不另外运行温度模型。','Electrothermal cases use CB dTB/dt = QB − (TB − Tsink)/R. Other dynamic cases prescribe a constant temperature and do not run a second temperature model.'))
    page();h(L('3 基准场景与输入含义','3 Baseline scenario and input meaning'))
    a=SUITE['parameters'];b=a['battery'];ini=SUITE['scene']['initial_state']
    table([L('输入','Input'),L('数值','Value'),L('含义','Meaning')],[
      ['A','1.8 m²',L('有效太阳能电池片面积','Active solar cell area')],['ηpv','0.30',L('固定最大功率点效率','Fixed maximum-power-point efficiency')],
      ['G','1361 W/m²',L('正常日照基准，地影用零表示；已包含日食影响，不再乘一次日食因子。','Nominal illumination; zero denotes eclipse. Eclipse is already included and is not multiplied twice.')],
      ['cos θ','1',L('太阳正对板面；为零是侧对太阳，负值是背面朝阳。','Normal incidence; zero is edge-on and negative is rear-face illumination.')],
      ['Ns × Np','8 × 6',L('每条支路串联八节，六条支路并联，共四十八个相同单体。','Eight series cells per string and six parallel strings, totaling 48 identical cells.')],
      ['ηD','0.95',L('配电器及其到计算负载的线路效率','Efficiency of distribution and its load-side harness')],
      ['RΩ','0.025 Ω',L('单体欧姆电阻','Cell ohmic resistance')],
      ['TB','293.15 K / 20 ℃',L('电池初始或规定温度，不是空间空气温度。','Initial or prescribed battery temperature, not space ambient air.')],
      ['xn / xp',f'{ini["x_n"]:.6f} / {ini["x_p"]:.9f}',L('两极平均锂占比，无量纲','Mean electrode lithium fractions, dimensionless')],
      ['SOC',f'{100*(ini["x_n"]-.03)/(.90-.03):.6f}%',L('由负极占比换算的可用剩余电量','Usable state of charge derived from negative fraction')],
      ['xn,0 / xn,100 / xp,100','0.03 / 0.90 / 0.27',L('同一参数包规定的空电与满电端点','Empty and full-charge endpoints from the same parameter set')],
      [L('单体电流范围','Cell current range'),'−3 to 6 A',L('负值为充电，正值为放电','Negative for charging, positive for discharge')],
      [L('单体电压范围','Cell voltage range'),'3.0 to 4.2 V',L('电芯端口电压约束','Cell terminal voltage constraint')],
      [L('材料有效范围','Material ranges'),'xn: 0.01 to 0.95; xp: 0.20 to 0.95',L('平均、中心与表面锂占比约束；不等同于 SOC 的 0 至 1 范围。','Constraints on mean, center and surface fractions; not equivalent to the usable SOC interval 0 to 1.')]
    ],[3.3,4.9,8.5],8.8)
    p(L('除明确覆盖值外，所有用例采用本表及附录材料参数。太阳常数和理想化周期是受控输入，不引用真实 ISS 工况；低温和高温均指表中给定的电池温度。','Unless explicitly overridden, cases use this table and the material parameters in the appendix. Illumination and idealized cycles are controlled inputs, not actual ISS conditions. Low and high temperature refer to specified battery temperatures.'))
    p(L('输入表中的 imin 与 imax 是单体充电与放电电流边界，单位 A；vmin 与 vmax 是单体电压下限与上限，单位 V；eta 是配电效率。所有逻辑标志中，1 表示是，0 表示否。精确输入保存在 suite.json，表内显示值可能四舍五入。','In input tables, imin and imax are cell charge and discharge current bounds in A; vmin and vmax are cell voltage limits in V; eta is distribution efficiency. Logical flags use 1 for true and 0 for false. Exact inputs are retained in suite.json; displayed table values may be rounded.'))
    page();h(L('4 判据与测试环境','4 Criteria and execution environment'))
    table([L('比较对象','Quantity'),L('最大允许绝对差','Maximum absolute difference')],[
      [L('静态功率与热源','Static power and heat'),'0.000001 W'],[L('动态功率与热源','Dynamic power and heat'),'0.001 W'],
      [L('电压','Voltage'),'0.0001 V'],[L('电流','Current'),'0.0001 A'],[L('SOC 与锂占比','SOC and lithium fractions'),'0.000001'],
      [L('电池温度','Battery temperature'),'0.001 K'],[L('锂占比变化率','Lithium fraction derivative'),'0.00000001 s⁻¹'],
      [L('离散标志与原因编码','Discrete flags and reason codes'),L('完全相同','Exact agreement')],
      [L('预设事件时刻','Scheduled event time'),'0.001 s'],[L('累计能量相对差','Relative integrated energy difference'),'0.00001'],
      [L('功率守恒残差','Power conservation residual'),'0.000002 W'],[L('总锂相对漂移','Relative lithium inventory drift'),'0.0000000001'],
      [L('可用 SOC 范围','Usable SOC range'),L('0 至 1，允许 0.000001 的舍入误差','0 to 1, with a 0.000001 rounding allowance')],
      [L('满电充电边界','Full-charge boundary'),L('SOC 为 1 时单体电流不得小于 −0.00000001 A','At SOC 1, cell current must be at least minus 0.00000001 A')]
    ],[8.4,8.3],9)
    p(L('数值门限用于软件一致性验证，并非来自真实电芯精度指标。平均误差同时报告均方根值；总体数值通过要求每个适用信号均满足最大差值门限。精度复核要求加密前后的差异小于上述门限的十分之一。','Numerical thresholds assess software consistency, not measured cell accuracy. RMSE is reported alongside maximum error; every applicable signal must meet its maximum-error limit. Refinement changes must be below one tenth of the corresponding limit.'))
    p(L('输入版本 1.0 在首次求解前冻结。覆盖性复核发现满电要求漏测后，版本 1.1 增加 CT-003 和 SOC 行为判据。原数值误差门限保持不变，旧版输入与判据保存在 results/history。','Input version 1.0 was frozen before initial execution. Coverage review added CT-003 and SOC behavior criteria in version 1.1 after identifying missing full-charge coverage. Numerical thresholds remain unchanged; prior inputs and criteria are retained in results/history.'))
    p('MATLAB R2026b 26.2.0.3386108; Simulink 26.2; Windows; Python Thermal/.venv; SciPy DOP853.')
    p(L('基准积分相对容差 10⁻⁸，绝对容差 10⁻¹⁰，最大步长 10 s。复核相对容差 10⁻¹⁰，绝对容差 10⁻¹²，最大步长 2 s。','Baseline integration uses relative tolerance 10⁻⁸, absolute tolerance 10⁻¹⁰ and maximum step 10 s. Refinement uses 10⁻¹⁰, 10⁻¹² and 2 s.'))
    page();h(L('5 测试项目与覆盖关系','5 Test projects and coverage'))
    table([L('编号','ID'),L('测试项目','Project'),L('对比','Agreement'),L('验收','Acceptance')],
      [[c['id'],c['name_en' if en else 'name_cn'],L('通过','Pass'),verdict(RESULT[c['id']])] for c in SUITE['cases']],[2.1,10,2.2,2.4],9)
    p(L('太阳阵对应设计 P2；电池与热源对应 P4 至 P8 和 P10；功率分配对应 P1、P3 与 P11；保护与接口对应设计第 5 章。动态项目组合这些组件，NU-001 对求解精度进行独立复核。','Array tests map to P2; cell and heat tests to P4–P8 and P10; allocation to P1, P3 and P11; protection and interfaces to design chapter 5. Dynamic cases combine these components; NU-001 independently reviews numerical convergence.'))
    # Two deliberate pages per test: specification and results.
    for num,c in enumerate(SUITE['cases'],1):
        cid=c['id'];res=RESULT[cid];txt=TEXT[cid];page();h(f'6.{num} {cid} '+c['name_en' if en else 'name_cn'])
        h(L('测试目的与场景','Purpose and physical scenario'),2);p(txt[en]);p(txt[2+en])
        h(L('输入与初始条件','Inputs and initial conditions'),2)
        if c['kind']=='dynamic':
            start=0;rows=[]
            for seg in c['segments']:
                rows.append([f'{fmt(start)} to {fmt(seg["end"])}',fmt(seg['G']),fmt(seg['request']),seg['command']]);start=seg['end']
            table([L('时间 s','Time s'),'G W/m²',L('请求 W','Demand W'),L('起点指令','Start command')],rows,[4.2,3,3,6.5],9)
            p(L(f'初始 xn = 0.60，xp = {ini["x_p"]:.9f}；初始温度 {c["T0"]:.2f} K；板面入射余弦恒为 1；输出间隔 {fmt(c["step"])} s。',f'Initial xn = 0.60, xp = {ini["x_p"]:.9f}; initial temperature {c["T0"]:.2f} K; incidence cosine 1; output interval {fmt(c["step"])} s.'))
            if c['thermal']:p(f'CB = {fmt(c["C"])} J/K; R = {fmt(c["R"])} K/W; Tsink = {fmt(c["Tsink"])} K.')
            if c['overrides']:p(L('参数覆盖值 ','Overrides ')+', '.join(f'{k} = {fmt(v)}' for k,v in c['overrides'].items()))
        else:
            mode=c['rows'][0]['mode'];rows=[]
            for i,r in enumerate(c['rows'],1):
                if mode=='solar':rows.append([i,fmt(r['G']),fmt(r['cosine'])])
                elif mode=='battery':rows.append([i,fmt(r['xn']),fmt(r['xp']),fmt(r['T']),fmt(r['current']),f'{r["overrides"].get("Ns",8)} × {r["overrides"].get("Np",6)}'])
                elif cid=='PR-001':rows.append([i,fmt(r['G']),fmt(r['request']),f'{r["connected"]} / {r["latched"]}',r['command']])
                elif cid=='IV-001':rows.append([i,r['bad'],L('边界事件','Boundary event') if i==6 else L('无效输入','Invalid input')])
                else:rows.append([i,fmt(r['G']),fmt(r['request']),fmt(r['T']),', '.join(f'{k}={fmt(v)}' for k,v in r['overrides'].items()) or L('基准','Baseline')])
            if mode=='solar':headers=[L('编号','Index'),'G W/m²','cos θ'];widths=[2,7,7.7]
            elif mode=='battery':headers=[L('编号','Index'),'xn','xp','TB K','i A','Ns × Np'];widths=[1.5,2.6,3.2,3,3,3.4]
            elif cid=='PR-001':headers=[L('编号','Index'),'G W/m²',L('请求 W','Demand W'),L('连接 / 锁存','Connected / latched'),L('指令','Command')];widths=[1.3,2.6,2.7,4,6.1]
            elif cid=='IV-001':headers=[L('编号','Index'),L('输入变更','Input change'),L('预期分类','Expected class')];widths=[2,8,6.7]
            else:headers=[L('编号','Index'),'G W/m²',L('请求 W','Demand W'),'TB K',L('参数变更','Override')];widths=[1.3,2.7,2.7,2.7,7.3]
            table(headers,rows,widths,8.5)
            if cid=='CT-003':p('xn = 0.90; xp = 0.27; SOC = 1; cos θ = 1.')
        h(L('执行步骤与预期结果','Procedure and expected outcome'),2)
        p(L('读取同一输入，分别运行受测模型与 Simulink，按本章输入顺序或规定时间保存结果。检查输出是否有效，再逐个比较功率、电压、电流、SOC、热源和适用状态标志。依据第 4 章门限判定数值一致性，另外检查本用例的物理与保护要求。','Load the same inputs and independently run the project model and Simulink. Save outputs in the listed input order or at prescribed times. Check output validity, then compare applicable power, voltage, current, SOC, heat and flags. Apply section 4 numerical limits and separately check the physical and protection requirements.'))
        p(L('证据文件 ','Evidence files ')+f'results/{cid}_python.csv; {cid}_simulink.csv; {cid}_difference.csv; {cid}.json.','Caption')
        page();h(f'{cid} '+L('对比结果','Comparison results'))
        figure(c,lang);doc.add_picture(str(FIG/f'{cid}_{lang}.png'),width=Cm(16.7))
        names=L('与',' and ').join(label(n,lang) for n in PLOT[cid])
        caption(L(f'图 {num}  {cid} 的 {names}。上排为双方同一采样点的输出，下排为本项目模型减 Simulink，单位与对应上图相同。',f'Figure {num}  {cid}: {names}. Top panels show both outputs at identical samples; lower panels show project model minus Simulink in the same units.'))
        if c['kind']=='static':p(L('横轴编号严格对应前页输入表，连线仅帮助观察变化，不表示这些独立工况之间的动态过程。','Horizontal indices match the preceding input table. Connecting lines aid reading and do not imply a dynamic trajectory between independent samples.'),'Caption')
        else:p(L('横轴为分钟。输入变化时保留同刻左右两侧值，竖直跳变表示输入或连接状态突变；未在突变两侧之间插值。','The horizontal axis is minutes. Vertical changes retain both sides of a discontinuity at the same time; no interpolation spans that discontinuity.'),'Caption')
        wanted=list(dict.fromkeys(PLOT[cid]+['P_pv_W','P_load_W','V_B_V','I_B_A','Q_B_W','SOC','T_B_K']))[:8]
        metrics=[m for n in wanted for m in res['metrics'] if m['signal']==n]
        table([L('比较量','Signal'),L('最大绝对差','Max absolute error'),'RMSE',L('门限','Limit')],
          [[label(m['signal'],lang),fmt(m['max_abs']),fmt(m['rmse']),fmt(m['threshold'])] for m in metrics],[7.6,3.1,3,3],8.5)
        p(L(f'数值对比通过；总体验收{verdict(res)}。已记录 {len(res["checks"])} 项检查。',f'Numerical comparison passes; overall acceptance: {verdict(res)}. Recorded {len(res["checks"])} checks.'))
        p(observed(c))
    page();h(L('7 结果汇总与缺陷','7 Summary and findings'))
    p(L('本轮完成 20 个项目，数值对比 20 项通过，设计验收 17 项通过、3 项不通过。三个失败项是同一满电约束缺失在静态与动态工况下的表现。','All 20 projects completed. Numerical agreement passes in 20 cases; design acceptance passes in 17 and fails in 3. The three failures expose the same omitted full-charge constraint in static and dynamic conditions.'))
    table([L('项目','Case'),L('本项目结果','Project result'),L('判定原因','Reason')],[
      ['CT-003',L('SOC 为 1 时仍以约 −0.545 A 单体电流充电','Charging at about minus 0.545 A per cell at SOC 1'),L('未按满电端点停止充电','No stop at the configured full-charge endpoint')],
      ['DY-002',L('SOC 峰值 100.6153%','Peak SOC 100.6153%'),L('单周期越过满电点','Full charge exceeded in one cycle')],
      ['DY-007',L('SOC 峰值 104.2070%','Peak SOC 104.2070%'),L('重复周期进一步暴露超充状态','Repeated cycles further expose the overcharge state')]
    ],[2.1,7.5,7.1])
    p(L('代码依据：受测模型的 _linear_limit 仅考虑电流、颗粒表面和中心占比，_search 再检查端电压；SOC 换算的 xn,100 未被加入充电约束。独立 Simulink 方程按相同数学模型求解后复现该问题。因此证据支持约束集合与满电文字要求不一致，并非两种求解器的数值误差。','Code evidence: _linear_limit checks current and particle surface or center fractions, and _search additionally checks terminal voltage. The xn,100 endpoint used to compute SOC is absent from charging limits. The independently implemented Simulink equations reproduce this behavior. Evidence supports a mismatch between the constraint set and the full-charge requirement, rather than an integration error.'))
    p(L('建议修正：明确可用 SOC 端点是否属于充放电保护约束，在到达边界的已接受时刻重新分配电力，满电时将多余太阳功率限发。修正后应重跑 CT-003、DY-002、DY-007 及全部回归项目。不能仅把显示 SOC 裁剪到 1。','Recommended correction: explicitly decide that usable SOC endpoints are charge and discharge constraints, reallocate at accepted boundary events, and curtail surplus solar at full charge. Rerun CT-003, DY-002, DY-007 and the regression suite after correction. Clipping displayed SOC to 1 is insufficient.'))
    h(L('验证范围与未覆盖内容','Validation scope and uncovered behavior'),2)
    p(L('本报告验证同一方程体系的实现一致性及列出的设计行为。未覆盖目标电芯实测精度、快速扩散松弛、高倍率电解液效应、开关纹波、MPPT 跟踪瞬态、连续耗尽所触发的自主事件定位、完整六节点 Thermal 联调和全部 SimReady 资产错误。预设输入边界的事件处理已经覆盖，不能据此宣称所有连续边界定位均已验证。','This report verifies implementations of the same equations and the listed design behaviors. It does not cover measured-cell accuracy, rapid diffusion relaxation, high-rate electrolyte effects, switching ripple, MPPT tracking transients, autonomous event localization during continuous depletion, full six-node Thermal integration, or every SimReady asset error. Scheduled-boundary protection is covered; this is not evidence of all continuous-boundary localization behavior.'))
    page();h(L('附录 A 电池材料参数','Appendix A Battery material parameters'))
    n=b['electrodes']['n'];pp=b['electrodes']['p']
    desc=[('radius_m','颗粒半径 m','Particle radius m'),('active_fraction','活性材料体积分数','Active material volume fraction'),
      ('electrode_area_m2','有效涂覆面积 m²','Effective electrode area m²'),('thickness_m','电极厚度 m','Electrode thickness m'),
      ('c_max_mol_m3','最大锂浓度 mol/m³','Maximum lithium concentration mol/m³'),('D_ref_m2_s','参考扩散系数 m²/s','Reference diffusion coefficient m²/s'),
      ('I0_ref_A','参考交换电流 A','Reference exchange current A'),('E_D_J_mol','扩散活化能 J/mol','Diffusion activation energy J/mol'),
      ('E_I_J_mol','反应活化能 J/mol','Reaction activation energy J/mol'),('b_V','平衡电位多项式系数 V','Equilibrium potential polynomial coefficients V'),
      ('d_V_K','温度系数多项式 V/K','Temperature polynomial coefficients V/K')]
    def val(v):return ', '.join(fmt(x) for x in v) if isinstance(v,list) else fmt(v)
    table([L('参数','Parameter'),L('负极','Negative electrode'),L('正极','Positive electrode')],[[L(cn,ee),val(n[k]),val(pp[k])] for k,cn,ee in desc],[6.1,5.3,5.3],9)
    p(L('多项式系数按常数项到高次项排列，参考温度为 298.15 K。电极电荷量 Qk = F εk Ael Lk ck,max；Qk 是锂占比从 0 到 1 的电荷量，不是直接填写的铭牌 Ah。所有数值为现有测试示例，不代表已选型或标定电芯。','Polynomial coefficients run from constant term to highest power, with reference temperature 298.15 K. Electrode charge is Qk = F εk Ael Lk ck,max, the charge for a fraction change from 0 to 1 rather than a directly specified rated Ah capacity. These are existing illustrative test values, not a selected or calibrated cell.'))
    h(L('参数来源与约束','Sources and constraints'),2)
    p('Power/SDTwin_Power_Design_Report_CN.docx; Power/SDTwin_Power_Design_Report_EN.docx; Thermal/sdtwin_sim/power_stand_in.py, example_power_records.')
    p(L('材料温度有效范围为 253.15 至 333.15 K。电池欧姆损耗已含在电池热源中；配电效率覆盖 PDU 至 Compute 线路，避免对同一线路损耗重复计入。','The material temperature range is 253.15 to 333.15 K. Cell ohmic loss is included in battery heat; distribution efficiency covers the PDU-to-Compute harness to avoid duplicate loss accounting.'))
    page();h(L('附录 B 复现与证据索引','Appendix B Reproduction and evidence'))
    p(L('工作目录为 C:/Workspace/SpaceDC。按以下顺序运行即可重新生成输入、双方结果、逐项判定与报告。','Use C:/Workspace/SpaceDC as the working directory. Run the following sequence to regenerate fixtures, independent results, verdicts and reports.'))
    for line in ['Thermal/.venv/Scripts/python.exe Power/verification/prepare.py',
       'Thermal/.venv/Scripts/python.exe Power/verification/run_python.py',
       "matlab -batch \"addpath('Power/verification/matlab'); run_suite\"",
       'Thermal/.venv/Scripts/python.exe Power/verification/compare.py',
       'python Power/verification/build_report.py']:
        p(line)
    p(L('最后一条使用安装了 python-docx 和 matplotlib 的文档运行环境；matlab 命令可替换为本机 R2026b 的完整可执行路径。','Use a document runtime with python-docx and matplotlib for the last command. Replace matlab with the full local R2026b executable path when needed.'))
    table([L('文件','File'),L('内容','Content')],[
      ['results/suite.json',L('完整参数、用例与判据','Complete parameters, cases and criteria')],
      ['matlab/sdtwin_power_reference.slx',L('实际运行的 Simulink 模型','Executed Simulink model')],
      ['matlab/power_reference.m',L('独立参考方程','Independent reference equations')],
      ['results/<ID>_python.csv / <ID>_simulink.csv',L('双方原始结果','Raw results from both implementations')],
      ['results/<ID>_difference.csv',L('同一采样点的直接差值','Direct differences at shared samples')],
      ['results/<ID>.json / summary.json',L('全部检查、误差与判定','All checks, errors and verdicts')],
      ['results/<ID>_*_events.json',L('实际事件时刻与连接状态','Actual event times and connectivity')],
      ['results/NU-001_*_fine.csv',L('独立精度加密结果','Independent refinement runs')],
      ['results/history',L('覆盖性补充前的原输入与判据','Original inputs and criteria before coverage expansion')]
    ],[9,7.7],9)
    p('Input SHA256: '+SUMMARY['suite_sha256'],'Caption')
    p(L('Simulink 实现依据 MathWorks 的 Level 2 MATLAB S 函数接口，连续状态由引擎调用 Derivatives 和 InitializeConditions 方法处理。','The implementation uses MathWorks Level 2 MATLAB S-function interfaces. The Simulink engine manages continuous states through Derivatives and InitializeConditions callbacks.'))
    p('https://www.mathworks.com/help/simulink/sfg/writing-level-2-matlab-s-functions.html','Caption')
    page();h(L('附录 C 累计能量对比','Appendix C Integrated energy comparison'))
    p(L('对输出时间序列分段使用梯形积分。事件左右两侧具有相同时间，因此时间宽度为零，不引入跨边界能量。这里比较的是双方用同一输出间隔计算的累计能量，并非额外实测能量。','Integrate each output series with the trapezoidal rule. Left and right event values have identical time, so their interval contributes zero energy. This compares energy at the common output cadence, not additional measured energy.'))
    for cid in ['DY-002','DY-007']:
        h(cid,2)
        table([L('能量来源','Energy source'),L('本项目 J','Project J'),'Simulink J',L('相对差','Relative difference')],
          [[label(n,lang).replace(' W',''),fmt(e['python_J']),fmt(e['simulink_J']),fmt(e['relative_error'])] for n,e in RESULT[cid]['energy'].items()],
          [6.5,3.4,3.4,3.4],9)
    p(L('电池端口能量带符号：正值表示净放电，负值表示净充电。电池热源能量与配电损耗分别积分，电池内部发热不再加入端口功率平衡。能量比较通过不改变这两个用例的 SOC 验收失败结论。','Battery terminal energy is signed: positive for net discharge and negative for net charge. Battery heat and distribution loss are integrated separately; internal heat is not added to terminal-power balance. Passing energy comparison does not change the failed SOC acceptance of these cases.'))
    page();h(L('附录 D 数值收敛复核','Appendix D Numerical convergence review'))
    p(L('NU-001 分别加密 Python 和 Simulink 求解器。表中差值为同一求解器的基准结果减加密结果，取最大绝对值，不是两个软件之间的差值。','NU-001 refines Python and Simulink separately. Each entry is the maximum absolute difference between baseline and refined runs of the same solver, not the difference between the two programs.'))
    table([L('求解器','Solver'),L('物理量','Quantity'),L('最大变化','Maximum change'),L('允许变化','Allowed change')],
      [[m['solver'],label(m['signal'],lang),fmt(m['max_abs']),fmt(m['threshold'])] for m in RESULT['NU-001']['convergence']],
      [3,7,3.3,3.4],9)
    p(L('两个求解器均满足十分之一验收门限的收敛要求。该复核针对本用例的平滑电热过程，不代表已验证所有刚性工况或自主保护边界的事件定位精度。','Both solvers meet the refinement criterion of one tenth of the acceptance threshold. This check covers this smooth electrothermal scenario, not all stiff regimes or autonomous protection-event localization accuracy.'))
    from orbit_format import standardize, record_template_contract
    record_template_contract()
    doc=standardize(doc,lang,SUITE,SUMMARY,TEXT)
    dest=ROOT/'Power'/f'SDTwin_Power_Test_Report_{lang.upper()}.docx';doc.save(dest);print(dest)
    return dest
if __name__=='__main__':
    for lang in (sys.argv[1:] or ['cn','en']):build(lang)
