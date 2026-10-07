"""Validate final evidence, DOCX structure, and frozen test fixture hashes."""
import json,hashlib,re,sys
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as E
from docx import Document
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[2];HERE=Path(__file__).resolve().parent;OUT=HERE/'results';QA=HERE/'qa';QA.mkdir(exist_ok=True)
summary=json.loads((OUT/'summary.json').read_text(encoding='utf-8'));suite=json.loads((OUT/'suite.json').read_text(encoding='utf-8'))
assert summary['execution_complete'] and len(summary['cases'])==20
assert summary['counts']=={'pass':17,'fail':3,'not_completed':0}
assert all(c['comparison_status']=='pass' for c in summary['cases'])
assert {c['id'] for c in summary['cases'] if c['status']=='fail'}=={'CT-003','DY-002','DY-007'}
assert hashlib.sha256((OUT/'suite.json').read_bytes()).hexdigest()==summary['suite_sha256']
for who in ['python','simulink']:
    meta=json.loads((OUT/f'{who}_execution.json').read_text(encoding='utf-8'))
    assert len(meta['cases'])==20 and all(c['status']=='executed' for c in meta['cases'])
docs=[]
for lang in ['CN','EN']:
    path=ROOT/'Power'/f'SDTwin_Power_Test_Report_{lang}.docx'
    doc=Document(path)
    expected_chapters=(['1 概述','2 测试内容','3 详细测试项目','4 测试内容充分性分析','5 测试条件与要求','6 测试数据','7 测试总结'] if lang=='CN' else ['1 Overview','2 Test Content','3 Detailed Test Projects','4 Test Content Sufficiency Analysis','5 Test Conditions and Requirements','6 Test Data','7 Test Summary'])
    assert [p.text for p in doc.paragraphs if p.style.name=='Heading 1']==expected_chapters
    forms=[t for t in doc.tables if t.cell(0,0).text==('用例编号' if lang=='CN' else 'Case ID')]
    assert len(forms)==20 and all(len(t.rows)==11 and len(t.columns)==4 for t in forms)
    assert [t.cell(0,1).text for t in forms]==[c['id'] for c in suite['cases']]
    assert doc.sections[-1].page_width>doc.sections[-1].page_height
    assert len(doc.inline_shapes)==20
    assert len(doc.sections)==5
    pdf=PdfReader(path.with_suffix('.pdf'))
    assert len(pdf.pages)==59
    assert all(len(p.extract_text().strip())>80 for p in pdf.pages)
    assert pdf.pages[3].extract_text().strip().startswith('1')
    with ZipFile(path) as z:
        xml=E.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        text=' '.join(n.text or '' for n in xml.findall('.//w:t',ns));tables=len(xml.findall('.//w:tbl',ns));drawings=len(xml.findall('.//w:drawing',ns))
        for c in suite['cases']:assert c['id'] in text and c['name_cn' if lang=='CN' else 'name_en'] in text
        assert drawings==20 and tables>=45
        assert '104.2070%' in text and '100.6153%' in text
    docs.append(dict(language=lang,format='Orbit seven-chapter test report',standard_case_forms=len(forms),figures=drawings,tables=tables,pages=len(pdf.pages),sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
files=[p for p in HERE.rglob('*') if p.is_file() and p.suffix in ['.m','.slx','.py','.json','.csv','.md','.ps1'] and '__pycache__' not in str(p) and 'qa' not in p.parts and 'cache' not in p.parts and p.name!='delivery_manifest.json']
files += [ROOT/'Power'/'IMPLEMENTATION.md',ROOT/'Thermal'/'sdtwin_sim'/'power_stand_in.py',*list((ROOT/'Power').glob('*Test_Report*'))]
manifest=dict(status='verified',case_count=20,numerical_comparison_pass=20,acceptance_counts=summary['counts'],
              check_count=sum(len(c['checks']) for c in summary['cases']),sample_count=sum(c['samples'] for c in summary['cases']),documents=docs,
              files={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in files})
(OUT/'delivery_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in manifest.items() if k!='files'},ensure_ascii=False,indent=2))
