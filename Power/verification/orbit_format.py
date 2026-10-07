"""Apply the retained Orbit report's document and test-record format to Power evidence.

The numerical report builder supplies content only. This module clones the actual
Orbit DOCX, its front matter, styles, eleven-row case form, and landscape summary.
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import json
import re
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
REFERENCE = next((ROOT / 'Orbit').glob('2*.docx'))
GROUPS = [
    ('太阳阵与电池组件', 'Solar Array and Battery Components', ['PV-001','BT-001','BT-002','BT-003','BT-004']),
    ('功率分配与保护接口', 'Power Allocation and Protection Interfaces', ['PD-001','PD-002','CT-001','CT-002','CT-003','PR-001','IV-001']),
    ('动态供电与电热过程', 'Dynamic Supply and Electrothermal Processes', ['DY-001','DY-002','DY-003','DY-004','DY-005','DY-006','DY-007']),
    ('数值精度复核', 'Numerical Accuracy Review', ['NU-001']),
]
EXPECTED = {
 'PV-001': ('背面照射与地影时发电为零；其余输出等于面积、效率及正向辐照度的乘积。','Rear illumination and eclipse yield zero power; other outputs equal active area times efficiency times front-face irradiance.'),
 'BT-001': ('零电流时发热与锂占比导数为零，端电压等于两极平衡电位之差。','Zero current yields zero heat and fraction derivatives; terminal voltage equals the electrode equilibrium-potential difference.'),
 'BT-002': ('双方电压、极化和边界裕量一致；越界电流点保留诊断，不误称为可供电状态。','Voltage, polarization and margins agree; infeasible current samples retain diagnostics and are not labeled feasible supply states.'),
 'BT-003': ('反应热、欧姆热及可逆热求和一致；负的可逆热与总热源不被截断。','Reaction, ohmic and reversible heat sum consistently; negative reversible heat and total heat are retained.'),
 'BT-004': ('整组电压随串联数变化，电流随并联数变化，功率与热源随总单体数变化。','Pack voltage scales with series count, current with parallel count, and power and heat with total cell count.'),
 'PD-001': ('太阳富余时充电，不足时放电；请求不可满足时明确报告供电不足。','Surplus solar charges the battery; a deficit discharges it; infeasible requests explicitly report shortfall.'),
 'PD-002': ('负载与配电发热之和等于配电入口功率；效率为一时配电发热为零。','Load plus distribution heat equals distribution input; unit efficiency produces zero distribution heat.'),
 'CT-001': ('遵守充电限制并降低太阳实际输出，仍满足可供电负载与端口守恒。','Honor charge limits and curtail solar while satisfying feasible load and terminal-power conservation.'),
 'CT-002': ('不可满足的负载返回供电不足及事件请求，和数值求解错误明确区分。','Infeasible demand returns shortfall and an event request, distinguished from numerical failure.'),
 'CT-003': ('SOC 为一时停止继续充电，多余太阳功率限发。','At SOC one, stop further charging and curtail surplus solar.'),
 'PR-001': ('启停、锁存、指令优先级及试算拒绝均符合预设事件序列。','Start, stop, latching, command priority and trial-event rejection match the prescribed event sequence.'),
 'IV-001': ('前五项输入无效，第六项为有效的器件边界事件；无效结果不提供伪造端口值。','The first five inputs are invalid; the sixth is a valid device-boundary event. Invalid results do not fabricate port values.'),
 'DY-001': ('SOC 连续上升且处于零至一，功率守恒，双方动态输出在误差范围内一致。','SOC rises continuously within zero to one, power is conserved, and dynamic outputs agree within tolerance.'),
 'DY-002': ('入影后由电池供电，出影后恢复太阳供电；SOC 全程处于零至一。','Battery supplies the eclipse interval; array supply resumes in sunlight; SOC stays within zero to one.'),
 'DY-003': ('功率与电流响应负载阶跃，锂状态在事件时刻保持连续。','Power and current respond to demand steps while lithium states remain continuous at each event.'),
 'DY-004': ('充电电流满足限制，太阳限发量随充电能力变化，双方轨迹一致。','Charging respects limits, curtailment follows charge capability, and both trajectories agree.'),
 'DY-005': ('100 s 失电，200 s 日照恢复但不自动启动，300 s 启动，400 s 停机，500 s 再启动。','Trip at 100 s; sunlight at 200 s causes no restart; start at 300 s, stop at 400 s, and restart at 500 s.'),
 'DY-006': ('电池温度仅积分一次，热源与散热共同决定温度变化，并反馈到电池电压与电流。','Temperature is integrated once; heat generation and rejection determine temperature and feed back into voltage and current.'),
 'DY-007': ('三个周期内守恒成立，双方累计能量一致，SOC 保持在零至一。','Conservation holds over three cycles, integrated energies agree, and SOC stays within zero to one.'),
 'NU-001': ('两侧各自加密后的变化均小于对应数值对比门限的十分之一。','Each solver changes by less than one tenth of the corresponding comparison limit under independent refinement.'),
}

def text(el):
    return ''.join(n.text or '' for n in el.iter(qn('w:t')))

def function(cid):
    if cid.startswith('PV'): return 'solar_power'
    if cid.startswith('BT'): return 'battery_response'
    if cid == 'PR-001': return 'apply_power_event'
    if cid == 'IV-001': return 'PowerInputs / PowerState'
    if cid.startswith(('PD','CT')): return 'solve_power_allocation'
    if cid == 'NU-001': return 'DOP853 / ode45'
    return 'solve_power_allocation / state integration'

def requirement(cid):
    if cid.startswith('PV'): return 'P2'
    if cid.startswith('BT'): return 'P4–P8, P10'
    if cid.startswith(('PD','CT')): return 'P1, P3, P11'
    if cid.startswith(('PR','IV')): return '5'
    return 'P1–P11'

def standardize(source, lang, suite, summary, prose):
    en = lang == 'en'
    L = lambda cn, ee: ee if en else cn
    results = {c['id']: c for c in summary['cases']}
    cases = {c['id']: c for c in suite['cases']}
    doc = Document(REFERENCE)
    body = doc.element.body
    kids = list(body)
    form = deepcopy(doc.tables[4]._tbl)
    portrait = deepcopy(doc.sections[3]._sectPr)
    # Body page numbering must start on its own page after the contents. Keeping
    # the source's continuous break lets Word merge the shortened TOC with it.
    typ=portrait.find(qn('w:type'))
    if typ is not None:typ.set(qn('w:val'),'nextPage')
    landscape = deepcopy(doc.sections[-1]._sectPr)
    chunks = {}
    current = None
    for el in source.element.body:
        if el.tag == qn('w:sectPr'): continue
        st = el.find('./' + qn('w:pPr') + '/' + qn('w:pStyle'))
        if st is not None and st.get(qn('w:val')) == source.styles['Heading 1'].style_id:
            current = text(el); chunks[current] = []
        elif current:
            chunks[current].append(el)
    def get(prefix): return next(v for k,v in chunks.items() if k.startswith(prefix))
    for el in kids[16:]: body.remove(el)
    body.append(deepcopy(landscape))
    # Source front matter and section furniture remain authoritative.
    def replace_par(el, value):
        p = Paragraph(el, doc)
        old = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        for child in list(el):
            if child.tag != qn('w:pPr'): el.remove(child)
        r = p.add_run(value)
        if old is not None: r._r.insert(0, old)
        r.font.color.rgb = RGBColor(0,0,0)
        return p
    replace_par(kids[5], 'SDTwin POWER MODULE')
    replace_par(kids[6], L('供电模块软件测试报告', 'Power Module Software Test Report'))
    replace_par(kids[7], L('太阳阵发电  电池响应  功率分配  保护控制  Simulink 对比',
                         'Solar Array  Battery Response  Power Allocation  Protection  Simulink Comparison'))
    replace_par(kids[11], L('版本记录','Version Control'))
    cover = Table(kids[8], doc)
    for i, lab in enumerate([L('单位','Institute'), L('编制','Author'), L('审核','Reviewer'), L('日期','Issue date')]):
        replace_par(cover.cell(i,0).paragraphs[0]._p, lab)
    replace_par(cover.cell(1,1).paragraphs[0]._p, L('自动测试与报告生成','Automated verification and report generation'))
    replace_par(cover.cell(2,1).paragraphs[0]._p, '')
    replace_par(cover.cell(3,1).paragraphs[0]._p, L('2026年10月7日','7 October 2026'))
    abstract = L('本报告记录供电模型与独立 Simulink 仿真的测试框架、输入、判据和逐项结果。20 项数值对比全部通过，设计验收 17 项通过、3 项不通过，异常集中于满电充电约束。',
        'This report records the framework, inputs, criteria and individual comparisons of the Power model with independent Simulink simulation. All 20 numerical comparisons pass; 17 cases pass design acceptance and 3 fail due to the full-charge constraint.')
    vt = Table(kids[12], doc)
    front = [(0,0,L('摘要','Abstract')),(0,1,abstract),(1,0,L('关键词','Keywords')),(1,1,L('供电模型；Simulink；测试用例；数值对比；验收记录','Power model; Simulink; test cases; numerical comparison; acceptance records')),(2,0,L('修改记录','Change Record'))]
    for r,c,s in front: replace_par(vt.cell(r,c).paragraphs[0]._p,s)
    for c,s in enumerate([L('版本','Revision'),L('修改项','Change Request'),L('日期','Date'),L('编制','Author'),L('说明','Description')]): replace_par(vt.cell(3,c).paragraphs[0]._p,s)
    records = [
        ['1.0',L('测试执行','Test execution'),'2026-10-07','',L('完成 20 项模型与 Simulink 对比，保留全部结果。','Executed 20 model comparisons with Simulink and retained all results.')],
        ['1.1',L('标准格式修订','Standard format'),'2026-10-07','',L('按 Orbit 测试报告重排七章、用例记录表和横向汇总；实验数据与判据不变。','Applied the Orbit seven-chapter layout, case records and landscape summary; data and criteria unchanged.')],
    ]
    for row in range(4,len(vt.rows)):
        for j,cell in enumerate(vt.rows[row].cells): replace_par(cell.paragraphs[0]._p,records[row-4][j] if row<6 else '')
    # Preserve the TOC content control and field; Word rebuilds its cached entries.
    toc = kids[14]
    tp = next(toc.iter(qn('w:p')))
    replace_par(tp,L('目录','TABLE OF CONTENTS'))
    for el in toc.iter(qn('w:instrText')):
        if 'TOC ' in (el.text or ''): el.text = ' TOC \\o "1-3" \\h \\z \\u '
    for s in doc.styles:
        if s.type == 1:
            rp = s.element.get_or_add_rPr()
            fonts = rp.find(qn('w:rFonts'))
            if fonts is None: fonts=OxmlElement('w:rFonts');rp.insert(0,fonts)
            fonts.set(qn('w:eastAsia'),'SimSun')
        if s.name.startswith('Heading'):
            s.font.color.rgb=RGBColor(0,0,0)
        if s.style_id in ('TOC1','TOC2','TOC3'):
            s.font.size=Pt(10.5)
            s.paragraph_format.line_spacing=Pt(13);s.paragraph_format.space_after=Pt(0);s.paragraph_format.space_before=Pt(0)
            snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');s.element.get_or_add_pPr().append(snap)
    doc.core_properties.title=L('SDTwin 供电模块软件测试报告','Test Report for Power Module of SDTwin')
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    tabno=0
    def p(t,style=None): return doc.add_paragraph(t,style)
    def h(t,lev=1):
        par=doc.add_heading(t,lev)
        if lev==1:par.paragraph_format.page_break_before=True
        return par
    def page(): doc.add_page_break()
    def sublabel(value):
        par=p(value);par.paragraph_format.keep_with_next=True
        par.paragraph_format.space_before=Pt(9)
        for run in par.runs:run.bold=True
        return par
    def caption(t):
        nonlocal tabno
        tabno+=1
        p(L(f'表 {tabno}  ',f'Table {tabno}. ')+t,'Caption')
    def tidy_table(tbl,width=14.65):
        old=sum(int(x.get(qn('w:w'))) for x in tbl._tbl.tblGrid)
        target=int(Cm(width).twips);ratio=target/old
        for x in tbl._tbl.tblGrid:x.set(qn('w:w'),str(round(int(x.get(qn('w:w')))*ratio)))
        for x in tbl._tbl.iter(qn('w:tcW')):x.set(qn('w:w'),str(round(int(x.get(qn('w:w')))*ratio)))
        tw=tbl._tbl.tblPr.find(qn('w:tblW'))
        if tw is not None:tw.set(qn('w:w'),str(target));tw.set(qn('w:type'),'dxa')
        tbl.alignment=1;tbl.autofit=False
        for row in tbl.rows:
            tr=row._tr.get_or_add_trPr()
            if tr.find(qn('w:cantSplit')) is None:tr.append(OxmlElement('w:cantSplit'))
            seen=set()
            for cell in row.cells:
                if id(cell._tc) in seen:continue
                seen.add(id(cell._tc))
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next=False
                    par.paragraph_format.space_before=Pt(1);par.paragraph_format.space_after=Pt(1)
                    par.paragraph_format.line_spacing=1.05
                    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');par._p.get_or_add_pPr().append(snap)
                    for run in par.runs:
                        run.font.name='Times New Roman';run.font.size=Pt(9 if en else 9.5)
                        run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun')
        return tbl
    def table(headers,rows,widths,title,width=14.65):
        caption(title)
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
        for col,w in zip(t.columns,widths):col.width=Cm(w)
        for cell,s in zip(t.rows[0].cells,headers):cell.text=str(s)
        for vals in rows:
            for cell,s in zip(t.add_row().cells,vals):cell.text=str(s)
        for row in t.rows:
            for j,cell in enumerate(row.cells):cell.width=Cm(widths[j])
        for cell in t.rows[0].cells:
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'D9D9D9');cell._tc.get_or_add_tcPr().append(sh)
            for run in cell.paragraphs[0].runs:run.bold=True
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        tidy_table(t,width);return t
    def append(elements,table_title='',subheads=True):
        for original in elements:
            if original.tag==qn('w:p') and not text(original) and original.find('.//'+qn('w:drawing')) is None:continue
            el=deepcopy(original)
            for st in el.iter(qn('w:pStyle')):
                old=st.get(qn('w:val'))
                try:name=source.styles[old].name
                except KeyError:name='Normal'
                if name.startswith('Heading'):
                    st.set(qn('w:val'),doc.styles['Normal'].style_id)
                    for r in el.iter(qn('w:r')): 
                        pr=r.find(qn('w:rPr'))
                        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
                        pr.append(OxmlElement('w:b'))
                else: st.set(qn('w:val'),doc.styles[name if name in doc.styles else 'Normal'].style_id)
            for st in el.iter(qn('w:tblStyle')):st.set(qn('w:val'),doc.styles['Table Grid'].style_id)
            for blip in el.iter(qn('a:blip')):
                rid=blip.get(qn('r:embed'))
                if rid:
                    nr,_=doc.part.get_or_add_image(BytesIO(source.part.related_parts[rid].blob));blip.set(qn('r:embed'),nr)
            for ext in el.iter(qn('wp:extent')):
                old=int(ext.get('cx'));ratio=int(Cm(14.65))/old
                ext.set('cx',str(int(Cm(14.65))));ext.set('cy',str(round(int(ext.get('cy'))*ratio)))
            for t in el.iter(qn('w:t')):
                if t.text:t.text=t.text.replace('第 4 章门限','第 5.1 节门限').replace('section 4 numerical limits','section 5.1 numerical limits').replace('前页输入表','本用例输入表').replace('preceding input table','input table for this case').replace('附录材料参数','第 6.1 节材料参数').replace('material parameters in the appendix','material parameters in section 6.1')
            if el.tag==qn('w:tbl'):caption(table_title)
            body.insert(len(body)-1,el)
            if el.tag==qn('w:tbl'):tidy_table(Table(el,doc))
    # 1 Overview, with the same subsection hierarchy as Orbit.
    h(L('1 概述','1 Overview'));h(L('1.1 标识','1.1 Identification'),2)
    p(L('平台名称：SDTwin。模块名称：供电模块。文档版本：1.1。测试日期：2026年10月7日。','Platform: SDTwin. Module: Power. Document revision: 1.1. Test date: 7 October 2026.'))
    p(L('文档用途：记录本项目模型与 Simulink 的可重复对比测试及实际验收结论。','Purpose: record repeatable comparison tests of the project model against Simulink and the actual acceptance verdicts.'))
    h(L('1.2 软件与文档概述','1.2 Software and Document Overview'),2)
    append(get('1 '),L('受测实现与参考实现','System Under Test and Reference Implementation'))
    h(L('1.3 缩略语与约定','1.3 Abbreviations and Conventions'),2)
    table([L('术语','Term'),L('含义','Definition')],[['SOC',L('可用剩余电量比例，零为空电，一为满电。','Usable charge fraction: zero is empty and one is full.')],['PDU',L('配电器，其损耗由配电效率决定。','Power distribution unit; losses follow the specified efficiency.')],['RMSE',L('双方在相同采样点的均方根误差。','Root mean square error at identical samples.')],['G / cos θ',L('已含地影的太阳辐照度与板面入射余弦。','Irradiance including eclipse and panel incidence cosine.')],['i / PB',L('正值为放电，负值为充电。','Positive for discharge, negative for charge.')]], [3,11.65],L('缩略语与符号','Abbreviations and Symbols'))
    page();h(L('2 测试内容','2 Test Content'))
    p(L('20 个测试项目分为四组，每个项目对应一个详细用例记录及一条最终汇总。测试同时检查数值一致性与设计行为，二者分别记录。','Twenty projects form four capability groups. Each maps to one detailed case record and one final summary record. Numerical agreement and design behavior are evaluated and recorded separately.'))
    h(L('2.1 功能需求与用例对应关系','2.1 Functional Requirement and Test Case Traceability'),2)
    table([L('序号','No.'),L('用例编号','Case ID'),L('被测函数或对象','Function'),L('测试名称','Test Name'),L('设计依据','Basis'),L('状态','Status')],[[i,c['id'],function(c['id']),c['name_en' if en else 'name_cn'],requirement(c['id']),L('通过','Pass') if results[c['id']]['status']=='pass' else L('不通过','Fail')] for i,c in enumerate(suite['cases'],1)],[.9,1.8,3.7,4.95,1.75,1.55],L('功能需求与测试用例对应关系','Functional Requirement and Test Case Traceability'))
    h(L('2.2 对比测试框架','2.2 Comparison Framework'),2)
    append(get('2 '),L('测试框架与执行约定','Test Framework and Execution Contract'))
    # 3 Source-derived formal case sheets plus the original inputs and figures.
    page();h(L('3 详细测试项目','3 Detailed Test Projects'))
    p(L('每项采用 Orbit 测试报告的统一记录字段。输入表给出具体场景，结果页提供双方输出、直接差值和误差表。编制及审核签名留待实际人员填写。','Every project uses the Orbit test report record fields. Input tables specify the scenario; result pages retain both outputs, direct differences and error tables. Human signature fields are intentionally blank.'))
    num=0
    for gi,(cn,ee,ids) in enumerate(GROUPS,1):
        if gi>1:page()
        h(f'3.{gi} '+L(cn,ee),2)
        for si,cid in enumerate(ids,1):
            num+=1
            if si>1:page()
            c=cases[cid];res=results[cid];name=c['name_en' if en else 'name_cn']
            h(f'3.{gi}.{si} '+L(f'测试项目 {num} ',f'Test Project {num} ')+name,3)
            specs=get(f'6.{list(cases).index(cid)+1} ')
            outs=get(cid+' ')
            actual=[text(e) for e in outs if e.tag==qn('w:p') and text(e)][-1]
            caption(cid+' '+L('测试用例','Test Case'))
            el=deepcopy(form);body.insert(len(body)-1,el);t=Table(el,doc)
            def cell(r,col,value):
                tc=t.cell(r,col)
                for extra in list(tc._tc)[1:]:tc._tc.remove(extra)
                pa=tc.add_paragraph(value);pa.style=doc.styles['Normal']
            labels=L(['用例编号','用例名称','测试目的','前置条件','输入数据','测试步骤','预期结果','验收判据','用例设计','执行人与日期','实际结果','记录人','测试结论','异常记录'],['Case ID','Case Name','Test Objective','Preconditions','Input Data','Procedure','Expected Results','Acceptance Criteria','Case Designer','Executor / Date','Actual Results','Recorded By','Test Conclusion','Anomalies'])
            cell(0,0,labels[0]);cell(0,1,cid);cell(0,2,labels[1]);cell(0,3,name)
            pre=L('采用第 5 章冻结参数与环境，双方独立执行。设计依据为 ', 'Use the frozen section 5 inputs and environment; execute both implementations independently. Design basis: ')+requirement(cid)+'.'
            inp=prose[cid][2+en]+L(' 完整数值见本用例输入表。',' Exact values appear in the input table for this case.')
            steps=L('1．载入本用例参数与初始状态，分别运行模型与 Simulink。\n2．按输入表顺序或规定时间保存双方输出与事件。\n3．核对有效性、同点差值及本用例行为判据，保留检查记录。','1. Load this case and run the project model and Simulink independently.\n2. Save outputs and events in input order or at the prescribed times.\n3. Check validity, same-sample differences and case behavior; retain check records.')
            accept=L('所有适用信号满足第 5.1 节门限，离散标志与预期事件一致，并满足本表预期结果。','All applicable signals meet section 5.1 limits, flags and events match expectations, and the expected results in this form hold.')
            values=[prose[cid][en],pre,inp,steps,EXPECTED[cid][en],accept]
            for r,v in enumerate(values,1):cell(r,0,labels[r+1]);cell(r,1,v)
            cell(7,0,labels[8]);cell(7,1,'');cell(7,2,labels[9]);cell(7,3,L('自动测试  2026年10月7日','Automated test  2026-10-07'))
            cell(8,0,labels[10]);cell(8,1,actual+L(f' 共记录 {len(res["checks"])} 项检查；证据为 ',f' {len(res["checks"])} checks recorded; evidence: ')+f'results/{cid}.json.')
            cell(9,0,labels[11]);cell(9,1,'');cell(9,2,labels[12]);cell(9,3,L('☒ 通过  ☐ 不通过  ☐ 未执行','[X] Pass   [ ] Fail   [ ] Not Executed') if res['status']=='pass' else L('☐ 通过  ☒ 不通过  ☐ 未执行','[ ] Pass   [X] Fail   [ ] Not Executed'))
            cell(10,0,labels[13]);cell(10,1,L('未发现不符合本用例判据的异常。','No anomaly against this case criteria.') if res['status']=='pass' else actual)
            tidy_table(t)
            for r in [0,1,2,3,4,5,6,7,8,9,10]:
                for col in ([0,2] if r in [0,7,9] else [0]):
                    for run in t.cell(r,col).paragraphs[0].runs:run.bold=True
            # The prior report's explicit input table and all associated initial conditions.
            start=next(i for i,e in enumerate(specs) if text(e)==L('输入与初始条件','Inputs and initial conditions'))
            end=next(i for i,e in enumerate(specs) if text(e)==L('执行步骤与预期结果','Procedure and expected outcome'))
            if c['kind']=='static' and len(c['rows'])>12:page()
            sublabel(cid+' '+L('场景输入与初始条件','Scenario Inputs and Initial Conditions'))
            append(specs[start+1:end],cid+' '+L('场景输入','Scenario Inputs'))
            p(L('未覆盖的参数沿用第 5.2 节基准及第 6.1 节材料参数。','Parameters not overridden here use sections 5.2 and 6.1.'))
            page();sublabel(cid+' '+L('对比结果','Comparison Results'))
            append(outs,cid+' '+L('数值误差与门限','Numerical Errors and Limits'))
    # 4 coverage and limitations, separated from final findings as in Orbit.
    page();h(L('4 测试内容充分性分析','4 Test Content Sufficiency Analysis'))
    table([L('能力组','Capability'),L('用例','Cases'),L('已覆盖的验证重点','Verified Focus')],[[L(a,b),', '.join(ids),L('组件响应、边界输入与独立对比','Component response, boundary inputs and independent comparison') if i<2 else L('动态状态、事件、能量或收敛检查','Dynamic states, events, energy or convergence checks')] for i,(a,b,ids) in enumerate(GROUPS)],[3.4,4.5,6.75],L('测试覆盖性','Test Coverage'))
    summ=get('7 ');scope=next(i for i,e in enumerate(summ) if text(e)==L('验证范围与未覆盖内容','Validation scope and uncovered behavior'))
    append(summ[scope+1:])
    p(L('全部 20 项已执行，覆盖不等于全部验收通过。CT-003、DY-002 和 DY-007 仍需满电约束修正后闭环；更广的实物精度与系统联调需新增测试。','All 20 cases executed, but coverage does not mean all acceptance criteria passed. CT-003, DY-002 and DY-007 require correction of full-charge constraints and retesting. Physical accuracy and broader system integration need additional tests.'))
    page();h(L('5 测试条件与要求','5 Test Conditions and Requirements'))
    h(L('5.1 测试环境与验收判据','5.1 Execution Environment and Acceptance Criteria'),2)
    append(get('4 '),L('数值及行为验收判据','Numerical and Behavioral Acceptance Criteria'))
    h(L('5.2 基准场景与输入要求','5.2 Baseline Scenario and Input Requirements'),2)
    append(get('3 '),L('基准参数与物理含义','Baseline Parameters and Physical Meaning'))
    page();h(L('6 测试数据','6 Test Data'))
    for i,(prefix,cn,ee) in enumerate([('附录 A' if not en else 'Appendix A','电池材料参数','Battery Material Parameters'),('附录 B' if not en else 'Appendix B','数据清单与复现','Data Inventory and Reproduction'),('附录 C' if not en else 'Appendix C','累计能量对比数据','Integrated Energy Comparison Data'),('附录 D' if not en else 'Appendix D','数值收敛数据','Numerical Convergence Data')],1):
        if i>1:page()
        h(f'6.{i} '+L(cn,ee),2);append(get(prefix),L(cn,ee))
    # Last portrait section closes before the landscape summary, matching Orbit.
    sp=doc.add_paragraph()._p.get_or_add_pPr();sp.append(portrait)
    h(L('7 测试总结','7 Test Summary'))
    p(abstract)
    rows=[]
    for i,c in enumerate(suite['cases'],1):
        cid=c['id'];out=get(cid+' ');actual=[text(e) for e in out if e.tag==qn('w:p') and text(e)][-1]
        rows.append([i,cid,function(cid),c['name_en' if en else 'name_cn'],L('通过','Pass') if results[cid]['status']=='pass' else L('不通过','Fail'),actual+' '+f'results/{cid}.json'])
    table([L('序号','No.'),L('用例编号','Case ID'),L('被测函数或对象','Function'),L('验证重点','Verification Focus'),L('状态','Status'),L('结果与证据','Result / Evidence')],rows,[.8,1.7,4,5,1.6,11.5],L('供电模块测试汇总','Power Module Test Summary'),24.6)
    page();h(L('7.1 异常记录与结论','7.1 Anomalies and Conclusion'),2)
    for el in summ[:scope]:
        if el.tag==qn('w:tbl'):
            caption(L('异常记录与判定依据','Anomalies and Verdicts'))
            new=deepcopy(el);body.insert(len(body)-1,new)
            nt=Table(new,doc);nt.style=doc.styles['Table Grid'];tidy_table(nt,24.6)
        else:append([el])
    p(L('总体结论：☐ 通过  ☒ 不通过  ☐ 未执行。20 项数值一致性通过；3 项设计行为验收未通过。','Overall Conclusion: [ ] Pass   [X] Fail   [ ] Not Executed. All 20 numerical comparisons pass; three design-behavior cases fail acceptance.'))
    p(L('现场测试专家签名：________________________    日期：________________','On-site Test Expert Signature: ________________________    Date: ________________'))
    for par in doc.paragraphs:
        if par.style.name=='Caption':par.paragraph_format.keep_with_next=not par.text.startswith(('图 ','Figure '))
    # Make chapter starts explicit without a redundant empty break paragraph.
    # paragraph that can move onto an otherwise empty page after a full table.
    for par in list(doc.paragraphs):
        if par.style.name=='Heading 1':
            prev=par._p.getprevious()
            if prev is not None and not text(prev) and prev.find('.//'+qn('w:br')) is not None and prev.find('.//'+qn('w:sectPr')) is None:
                prev.getparent().remove(prev)
    # No stale source-module text survives outside preserved styles or opaque metadata.
    final_text=' '.join(text(e) for e in body if e is not toc)
    assert 'OD-001' not in final_text and 'Orbit Dynamics Module' not in final_text
    doc.settings.element.append(OxmlElement('w:updateFields'))
    doc.settings.element[-1].set(qn('w:val'),'true')
    return doc

def record_template_contract():
    qa=HERE/'qa'/'orbit_template';qa.mkdir(parents=True,exist_ok=True)
    reference_hash=hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    (qa/'artifact.md').write_text(f'''# Power report template contract

Reference: {REFERENCE.as_posix()}
SHA256: {reference_hash}
Retained reference is read only. Its cover, version table, TOC content control,
styles, eleven-row case form and landscape summary are the formatting authority.

Page geometry: A4 portrait, 25.4 mm top and bottom, 31.75 mm left and right;
landscape summary, 19.05 mm top and bottom, 25.4 mm left and right.
Normal: Times New Roman 11 pt, 1.2 line spacing, 4 pt after. Heading 1 15 pt,
Heading 2 12.5 pt, black. Chinese glyphs use SimSun. Table grid is black,
labels gray, portrait case grid 1512/1584/1656/3528 twips in the source.

Slots: replace source module identity, abstract, dates, version entries and all
body content. Leave human author/reviewer/signature slots blank. Preserve actual
TOC control and refresh it in Word. New figures are Power numerical evidence.
Seven body chapters and 3.group.project hierarchy follow Orbit. Common inputs,
material/energy/convergence data move into chapters 5 and 6. Every case retains
its full input table and evidence figure beyond the standard form.

Repeat the portrait body section as needed; the redundant source section before
the original requirement matrix is omitted. Keep front-matter section properties
and final landscape properties. Source-specific bookmarks and generated TOC
entries are regenerated. Preserve header/footer/numbering/theme parts; styles
permit Chinese glyph font and compact TOC spacing. Table text is 9 pt English or
9.5 pt Chinese for readability. Source prose, signers and results are not copied.

Fidelity gates: same seven chapters, exact case-form fields, 20 distinct case
IDs, 20 figures, landscape summary, actual 17 pass and 3 fail, updated TOC/page
fields; render all final pages and inspect. No input, result or threshold changes.
''',encoding='utf-8')
